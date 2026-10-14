"""Hour-boundary semantics, immutable snapshots and existing algorithm reuse."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, timedelta, timezone
import uuid

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.hourly_brief.columns import COLUMNS, cell
from app.hourly_brief.inputs import Inputs, RAW_ASH, COAL_AMOUNT
from app.hourly_brief.models import HourlyBriefRow, HourlyBriefSchedule
from app.hourly_brief.service import tick, calculate, build_snapshot
from app.models import ManualMeasurement, CoarseModel, AutoState, Setting
from app.services.density import calc_total_ash, compute_density_guidance
from app.services.modeling import MLR_FEATURES
from app.telemetry.config import CollectorConfig, PointConfig
from app.telemetry.derived import DENSITY_AVERAGE_MEMBERS
from app.telemetry.domain import normalize, iso_time
from app.telemetry.models import TelemetryReading, TelemetryPoint
from app.telemetry.repository import TelemetryRepository

HOUR = datetime(2026, 10, 8, 2, tzinfo=timezone.utc)  # 10:00 Beijing


@pytest.fixture
def sessions(tmp_path):
    engine = create_engine("sqlite:///" + str(tmp_path / "hourly.db"),
                           connect_args={"check_same_thread": False, "timeout": 15})
    Base.metadata.create_all(engine)
    yield sessionmaker(engine, autoflush=False)
    engine.dispose()


def send(repo, cfg, when, values, statuses=None):
    samples = [normalize(p, value, code, when, when, when)
               for p, value, code in zip(cfg.points, values, statuses or [0] * len(values))]
    batch = dict(source_id=cfg.source_id, batch_id=uuid.uuid4().hex, samples=samples)
    repo.persist(cfg, "owner", batch)
    repo.state(cfg, "owner", "connected")
    return batch


def config(items):
    return CollectorConfig(enabled=True, endpoint="opc.tcp://localhost:49320",
        security="None", allow_insecure=True, points=[
            PointConfig(key="point" + str(i), name=item, node_id="ns=2;s=" + item, item_id=item)
            for i, item in enumerate(items)])


def setup_clock(monkeypatch, moment):
    monkeypatch.setattr("app.telemetry.repository.utc_now", lambda: moment)


def test_next_hour_only_idempotent_and_missing_rows_survive_restart(sessions):
    tick(HOUR - timedelta(minutes=36), sessions)
    with sessions() as db:
        assert db.get(HourlyBriefSchedule, "main").next_hour == iso_time(HOUR)
        assert list(db.scalars(select(HourlyBriefRow))) == []
    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(lambda _: tick(HOUR, sessions), range(2)))
    with sessions() as db:
        first = deepcopy(db.get(HourlyBriefRow, iso_time(HOUR)).snapshot)
        assert len(first["cells"]) == 17
        assert first["cells"]["coarse_amount"]["value"] is None
        assert first["cells"]["heavy_ash"]["status"] == "default"
    tick(HOUR + timedelta(hours=3, minutes=12), sessions)
    with sessions() as db:
        rows = list(db.scalars(select(HourlyBriefRow).order_by(HourlyBriefRow.hour)))
        assert len(rows) == 4
        assert rows[0].snapshot == first
        assert all(c["value"] is None for c in rows[-1].snapshot["cells"].values())
        assert "未回算" in rows[-1].snapshot["reason"]


def test_nearest_pre_hour_is_preserved_when_latest_has_moved(sessions, monkeypatch):
    setup_clock(monkeypatch, HOUR + timedelta(seconds=3))
    cfg = config([RAW_ASH])
    repo = TelemetryRepository(sessions)
    repo.acquire(cfg, "owner")
    send(repo, cfg, HOUR - timedelta(seconds=50), [40])
    last = send(repo, cfg, HOUR - timedelta(seconds=2), [42])
    send(repo, cfg, HOUR + timedelta(seconds=3), [99])
    with sessions() as db:
        value = Inputs(db, HOUR).telemetry(RAW_ASH)
        assert value["value"] == 42
        assert value["observed_at"] == iso_time(HOUR - timedelta(seconds=2))
        assert db.scalar(select(TelemetryReading).where(TelemetryReading.batch_id == last["batch_id"]))


def test_zero_fallback_excludes_future_nonzero_and_sets_stop(sessions, monkeypatch):
    setup_clock(monkeypatch, HOUR + timedelta(seconds=3))
    cfg = config([RAW_ASH, COAL_AMOUNT])
    repo = TelemetryRepository(sessions)
    repo.acquire(cfg, "owner")
    send(repo, cfg, HOUR - timedelta(minutes=2), [42, 800])
    send(repo, cfg, HOUR - timedelta(seconds=2), [0, 0])
    send(repo, cfg, HOUR + timedelta(seconds=3), [99, 999])
    with sessions() as db:
        result = build_snapshot(db, HOUR)
    for key, value in (("raw_ash", 42), ("coal_amount", 800)):
        assert result["cells"][key]["value"] == value
        assert result["cells"][key]["raw_value"] == 0
        assert result["cells"][key]["status"] == "held"
    assert result["model_inputs"]["is_stoppage"] == 1


def test_density_one_uses_own_history_and_excludes_future(sessions, monkeypatch):
    setup_clock(monkeypatch, HOUR + timedelta(seconds=3))
    cfg = config([item for _, item in DENSITY_AVERAGE_MEMBERS])
    repo = TelemetryRepository(sessions)
    repo.acquire(cfg, "owner")
    send(repo, cfg, HOUR - timedelta(seconds=50), [1.4, 1.6])
    send(repo, cfg, HOUR - timedelta(seconds=4), [1.46, 1.66])
    send(repo, cfg, HOUR - timedelta(seconds=2), [1, 1.68])
    with sessions() as db:
        value = Inputs(db, HOUR).density()
        assert value["value"] == pytest.approx(1.57) and value["status"] == "held"
        assert [c["value"] for c in value["components"]] == [1.46, 1.68]
        assert value["components"][0]["raw_value"] == 1
    send(repo, cfg, HOUR + timedelta(seconds=3), [1.9, 1.9])
    with sessions() as db:
        assert Inputs(db, HOUR).density()["value"] == pytest.approx(1.57)
    repo.state(cfg, "owner", "stopped")
    with sessions() as db:
        assert Inputs(db, HOUR).density()["value"] == pytest.approx(1.57), "Later samples/disconnect cannot change the earlier hour"
        held = Inputs(db, HOUR + timedelta(seconds=40)).density()
        assert held["value"] == 1.9 and held["status"] == "held"


def test_hour_boundary_keeps_old_held_density_with_legacy_history_gap(sessions, monkeypatch):
    setup_clock(monkeypatch, HOUR + timedelta(seconds=3))
    cfg = config([item for _, item in DENSITY_AVERAGE_MEMBERS])
    repo = TelemetryRepository(sessions)
    repo.acquire(cfg, "owner")
    send(repo, cfg, HOUR - timedelta(days=10), [1.44, 1.64])
    send(repo, cfg, HOUR - timedelta(seconds=2), [1, 1])
    send(repo, cfg, HOUR + timedelta(seconds=3), [1.9, 1.9])
    # Simulate data removed by the former policy, not the current collector.
    with sessions.begin() as db:
        db.execute(delete(TelemetryReading).where(
            TelemetryReading.received_at < iso_time(HOUR - timedelta(days=7))))
    with sessions() as db:
        result = Inputs(db, HOUR).density()
        assert result["value"] == pytest.approx(1.54)
        assert result["status"] == "held"


def add_manual(db, value, observed, created, metric="coarse_slime_amount"):
    db.add(ManualMeasurement(metric=metric, name="test", value=value, unit="t/h",
        observed_at=iso_time(observed), created_at=iso_time(created), request_hash=uuid.uuid4().hex))


def test_manual_nearest_observation_and_received_gate(sessions):
    with sessions() as db:
        add_manual(db, 41, HOUR - timedelta(minutes=1), HOUR - timedelta(seconds=20))
        add_manual(db, 39, HOUR - timedelta(minutes=10), HOUR - timedelta(seconds=10))
        add_manual(db, 99, HOUR - timedelta(seconds=1), HOUR + timedelta(seconds=1))
        add_manual(db, 98, HOUR + timedelta(minutes=1), HOUR - timedelta(seconds=1))
        db.commit()
        assert Inputs(db, HOUR).manual("coarse_slime_amount")["value"] == 41


def test_auto_manual_subvalues_have_individual_cutoffs(sessions):
    with sessions() as db:
        db.add(AutoState(key="amountInputs", updated_at=(HOUR + timedelta(hours=8, seconds=1)).replace(tzinfo=None),
            value=dict(coarseAmount=dict(manual=41, manualAt=(HOUR-timedelta(minutes=2)).timestamp()*1000),
                       floatAmount=dict(manual=99, manualAt=(HOUR+timedelta(seconds=1)).timestamp()*1000))))
        db.commit()
        data = Inputs(db, HOUR)
        assert data.manual(auto_paths=[("amountInputs", "coarseAmount")])["value"] == 41
        assert data.manual(auto_paths=[("amountInputs", "floatAmount")]) is None
        assert "amountInputs" not in data.auto


def test_saved_legacy_model_without_feature_names_is_compatible(sessions):
    from app.hourly_brief.service import settings_and_model
    with sessions() as db:
        db.add(CoarseModel(method="mlr", is_current=True, trained_at="2026-09-22 09:24:38",
                          feature_names=None, intercept=10, coefs=[.1]+[0]*9))
        db.commit()
        _, metadata = settings_and_model(db, HOUR, Inputs(db, HOUR))
        cells = complete_cells()
        calculate(cells, {}, metadata, HOUR)
        assert cells["coarse_ash"]["value"] == pytest.approx(14.2)
        assert db.scalar(select(CoarseModel)).feature_names is None


def test_status_group_rejects_mixed_batches(sessions, monkeypatch):
    setup_clock(monkeypatch, HOUR)
    cfg = config(["system-a", "system-b"])
    for p in cfg.points:
        p.value_type = "boolean"
    repo = TelemetryRepository(sessions)
    repo.acquire(cfg, "owner")
    send(repo, cfg, HOUR - timedelta(seconds=2), [True, False])
    with sessions() as db:
        point = db.get(TelemetryPoint, (cfg.source_id, "point1"))
        point.latest = {**point.latest, "batch_id": "another-batch"}
        db.commit()
        result = Inputs(db, HOUR).group([("A", "system-a"), ("B", "system-b")])
        assert result["value"] == {"A":None, "B":None}
        assert "批次" in result["reason"]


def complete_cells():
    cells = {key: cell(reason="test missing") for key, _, _ in COLUMNS}
    for key, value in dict(raw_ash=42, coal_amount=800, level=55, coarse_amount=40,
        heavy_ash=7.9, heavy_amount=235.2, float_ash=9.85, float_amount=20, density=1.45).items():
        cells[key] = cell(value, "test", iso_time(HOUR), "good")
    cells["systems"] = cell(dict(A=True, B=False, **{"401":True, "402":False}), status="good")
    cells["desliming"] = cell({"473":True, "474":False}, status="good")
    return cells


def model():
    return dict(model=dict(type="mlr", featureNames=MLR_FEATURES, intercept=10, coefs=[.1]+[0]*9))


@pytest.mark.parametrize("scheme", ["total", "heavy"])
def test_model_formula_and_original_density_algorithm(scheme):
    cells, meta = complete_cells(), model()
    before = deepcopy(meta)
    _, guide = calculate(cells, dict(guideScheme=scheme), meta, HOUR)
    total = calc_total_ash(7.9, 235.2, 9.85, 20, 14.2, 40)
    expected = compute_density_guidance(dict(rho_cur=1.45, heavy_ash=7.9, actual_total=total,
        scheme=scheme, heavy_amt=235.2, float_amt=20, coarse_amt=40, float_ash=9.85, coarse_ash=14.2), 8.5)
    assert cells["coarse_ash"]["value"] == pytest.approx(14.2)
    assert cells["total_ash"]["value"] == pytest.approx(total)
    assert cells["suggested_density"]["value"] == expected["rhoNew"]
    assert guide["maxStep"] == .02 and meta == before
    assert cells["reagent"]["value"] is None, "Unrelated missing display fields do not block calculations"


def test_missing_factor_no_imputation_and_zero_total_not_fabricated():
    cells = complete_cells()
    cells["systems"]["value"]["A"] = None
    calculate(cells, {}, model(), HOUR)
    assert cells["coarse_ash"]["value"] is None
    assert cells["total_ash"]["value"] is None
    cells = complete_cells()
    for key in ("heavy_amount", "float_amount", "coarse_amount"):
        cells[key]["value"] = 0
    calculate(cells, {}, model(), HOUR)
    assert cells["total_ash"]["value"] is None
    assert cells["suggested_density"]["value"] is None


@pytest.mark.parametrize("rho", [1.25, 1.75])
def test_measured_density_outside_old_bounds_keeps_step_limit(rho):
    cells = complete_cells()
    cells["density"]["value"] = rho
    calculate(cells, {}, model(), HOUR)
    assert cells["density"]["value"] == rho
    assert cells["suggested_density"]["value"] == pytest.approx(rho - .02)


def test_density_guard_uses_actual_driver_time_and_updated_settings_block():
    from app.services.guidance import density_guard_state
    from app.hourly_brief.inputs import LOCAL
    cells = complete_cells()
    when = HOUR - timedelta(seconds=3)
    for key in ("raw_ash", "coal_amount", "level"):
        cells[key]["observed_at"] = iso_time(when)
    total = calc_total_ash(7.9, 235.2, 9.85, 20, 14.2, 40)
    guard = density_guard_state(dict(coarseCoal=[{"timestamp":when.astimezone(LOCAL).strftime("%Y-%m-%d %H:%M:%S")}]),
                                total, 7.9, "total", HOUR.timestamp()*1000)
    calculate(cells, {"densityActionLatch":{"key":guard["driveKey"]}}, model(), HOUR)
    assert cells["suggested_density"]["value"] == 1.45
    assert "调整过" in cells["suggested_density"]["reason"]
    calculate(cells, {"_unavailable_settings":["ash_target"]}, model(), HOUR)
    assert cells["suggested_density"]["value"] is None


def test_get_is_read_only_and_day_filter_uses_beijing_midnight(sessions):
    from app.main import app
    def get_test_db():
        with sessions() as db:
            yield db
    app.dependency_overrides[get_db] = get_test_db
    try:
        client = TestClient(app)
        response = client.get("/api/v1/hourly-brief?day=2026-10-08").json()
        assert response["rows"] == [] and response["next_hour"] is None
        tick(HOUR - timedelta(minutes=1), sessions)
        tick(HOUR, sessions)
        response = client.get("/api/v1/hourly-brief?day=2026-10-08").json()
        assert response["total"] == 1 and len(response["columns"]) == 17
        assert client.get("/api/v1/hourly-brief?day=2026-10-07").json()["total"] == 0
        assert client.get("/api/v1/hourly-brief?limit=99999").status_code == 422
        assert client.post("/api/v1/hourly-brief").status_code == 405
        from openpyxl import load_workbook
        from io import BytesIO
        export = client.get("/api/v1/hourly-brief/export", params={"hour":iso_time(HOUR)})
        assert export.status_code == 200 and export.content.startswith(b"PK")
        workbook = load_workbook(BytesIO(export.content))
        assert workbook.sheetnames == ["小时推测简报", "数据来源与说明", "人工经验反馈"]
        assert workbook.worksheets[0].max_column == 19
        assert workbook.worksheets[0].max_row == 2
        assert workbook.worksheets[1].max_row == 18
        assert workbook.worksheets[0]["A2"].value == "2026-10-08 10:00:00"
    finally:
        app.dependency_overrides.pop(get_db, None)

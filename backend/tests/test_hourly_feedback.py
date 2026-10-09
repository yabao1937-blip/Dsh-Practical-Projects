"""Feedback durability and hour linkage; all writes use a disposable SQLite DB."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import auth
from app.database import Base, get_db
from app.hourly_brief.columns import COLUMNS, cell
from app.hourly_brief.models import HourlyBriefFeedback, HourlyBriefRow
from app.routers import hourly_brief, state
from app.services import migrate

URL = "/api/v1/hourly-brief"
HOUR = "2026-10-09T01:00:00.000+00:00"
OTHER = "2026-10-09T02:00:00.000+00:00"


@pytest.fixture
def setup(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{(tmp_path / 'feedback.db').as_posix()}",
                           connect_args={"check_same_thread": False, "timeout": 15})
    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine, autoflush=False)
    snapshot = {"cells": {key: cell(1.48 if 'density' in key else 8.5, status="good") for key, _, _ in COLUMNS},
                "model": {"method": "gpt", "trained_at": "fixture"}, "guidance": {"maxStep": .03}}
    with sessions() as db:
        for hour in [HOUR, OTHER]:
            db.add(HourlyBriefRow(hour=hour, generated_at=hour, snapshot=deepcopy(snapshot)))
        db.commit()
    app = FastAPI()
    app.include_router(hourly_brief.router, prefix="/api/v1")
    app.include_router(state.router, prefix="/api/v1")
    def database():
        with sessions() as db:
            yield db
    app.dependency_overrides[get_db] = database
    monkeypatch.setattr(migrate, "SessionLocal", sessions)
    monkeypatch.setattr(migrate, "backup_database", lambda *_: "isolated-test-backup")
    with TestClient(app) as client:
        yield client, sessions, snapshot
    engine.dispose()


def body(content="建议维持密度，观察下一小时灰分。", hour=HOUR):
    return {"hour": hour, "content": content, "client_id": str(uuid4())}


def test_persist_normalize_append_summary_export_and_mirror(setup):
    client, sessions, snapshot = setup
    first = body("  预测偏高\n建议观察  ", "2026-10-09T09:00:00+08:00")
    before = datetime.now(timezone.utc)
    saved = client.post(URL + "/feedback", json=first)
    assert saved.status_code == 201, saved.text
    record = saved.json()
    assert record["hour"] == HOUR and record["content"] == "预测偏高\n建议观察"
    assert datetime.fromisoformat(record["created_at"]) >= before
    assert client.post(URL + "/feedback", json={**first, "hour": HOUR}).json() == record
    second = client.post(URL + "/feedback", json=body('=1+1\n补充：<img src=x onerror=alert(1)>')).json()
    client.post(URL + "/feedback", json=body("另一个小时", OTHER))
    report = client.get(URL).json()
    assert len(report["columns"]) == 17
    by_hour = {r["hour"]: r for r in report["rows"]}
    assert by_hour[HOUR]["feedback"] == {"total": 2, "latest": second}
    assert by_hour[OTHER]["feedback"]["total"] == 1
    # Mirror replacement may touch business records, never report/feedback tables.
    for force in (False, True):
        store = client.get("/api/v1/state").json()
        result = client.put("/api/v1/state", params={"force": force}, json=store)
        assert result.status_code == 200, result.text
    with sessions() as db:
        assert db.query(HourlyBriefFeedback).count() == 3
        assert db.get(HourlyBriefFeedback, record["id"]).content == record["content"]
        assert db.get(HourlyBriefRow, HOUR).snapshot == snapshot
    export = client.get(URL + "/export", params={"hour": HOUR})
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content))
    assert workbook.active["S1"].value == "反馈（最新）"
    assert workbook.active["S2"].value.startswith(second["content"])
    history = workbook["人工经验反馈"]
    assert history.max_row == 3
    assert history["D2"].value == second["content"] and history["D2"].data_type == "s"
    assert history["D3"].value == record["content"]
    assert history["B2"].value == "2026-10-09 09:00:00"


def test_idempotent_concurrent_retry_and_conflicts(setup):
    client, sessions, _ = setup
    payload = body()
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(lambda _: client.post(URL + "/feedback", json=payload), range(4)))
    assert sorted(r.status_code for r in responses) == [200, 200, 200, 201]
    assert len({r.json()["id"] for r in responses}) == 1
    for changes in [{"hour": OTHER}, {"content": "不同内容"}]:
        assert client.post(URL + "/feedback", json={**payload, **changes}).status_code == 409
    with sessions() as db:
        assert db.query(HourlyBriefFeedback).count() == 1


def test_history_cursor_stays_stable_during_new_submissions(setup):
    client, _, _ = setup
    for n in range(3):
        client.post(URL + "/feedback", json=body(f"现场经验 {n}"))
    first = client.get(URL + "/feedback", params={"hour": HOUR, "limit": 2}).json()
    assert [r["content"] for r in first["items"]] == ["现场经验 2", "现场经验 1"]
    client.post(URL + "/feedback", json=body("新增经验"))
    last = client.get(URL + "/feedback", params={"hour": HOUR, "before_id": first["next_before_id"], "limit": 2}).json()
    assert [r["content"] for r in last["items"]] == ["现场经验 0"]
    assert last["total"] == 4 and last["next_before_id"] is None


@pytest.mark.parametrize("changes", [
    {"content": "  \n "}, {"content": "长" * 2001}, {"content": "a\x00b"},
    {"content": "a\ufffeb"}, {"content": 1}, {"client_id": "invalid"},
    {"hour": "2026-10-09T09:00:00"}, {"hour": "2026-10-09T09:01:00+08:00"},
    {"hour": "0001-01-01T00:00:00+08:00"}, {"hour": "9999-12-31T23:00:00-01:00"},
])
def test_invalid_feedback_never_writes(setup, changes):
    client, sessions, _ = setup
    assert client.post(URL + "/feedback", json={**body(), **changes}).status_code == 422
    if "hour" in changes:
        assert client.get(URL + "/feedback", params={"hour": changes["hour"]}).status_code == 422
    with sessions() as db:
        assert db.query(HourlyBriefFeedback).count() == 0


def test_missing_hour_and_write_auth(setup, monkeypatch):
    client, sessions, _ = setup
    missing = "2026-10-09T03:00:00Z"
    assert client.post(URL + "/feedback", json=body(hour=missing)).status_code == 404
    assert client.get(URL + "/feedback", params={"hour": missing}).status_code == 404
    monkeypatch.setattr(auth, "_LOCALHOST_WRITES", False)
    monkeypatch.setattr(auth, "WRITE_TOKEN", "feedback-test")
    assert client.post(URL + "/feedback", json=body()).status_code == 401
    assert client.get(URL + "/feedback", params={"hour": HOUR}).json()["total"] == 0
    assert client.post(URL + "/feedback", json=body(), headers={"X-DMCS-Token": "feedback-test"}).status_code == 201

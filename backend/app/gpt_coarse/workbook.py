"""Read the supplied 12-column sheet without editing it or the shared DS records."""
from collections import Counter
from datetime import date, datetime, time
import hashlib
from pathlib import Path

from openpyxl import load_workbook

from .predict import FEATURES, number


def read_workbook(path, year):
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook['Sheet1']
    all_rows = list(sheet.values)
    if not (str(all_rows[1][2]).startswith('原煤灰分') and '315' in str(all_rows[1][10])
            and list(all_rows[2][4:10]) == ['A', 'B', 401, 402, 473, 474]):
        raise ValueError('表头与指定的9因子/315灰分模板不一致')
    valid_rows, stopped, rejected, fixes = [], [], [], []
    previous_day = date(year, 1, 1)
    for row_number, row in enumerate(all_rows[3:], 4):
        if not any(v is not None for v in row):
            continue
        try:
            raw_day, raw_time = row[:2]
            if isinstance(raw_day, (datetime, date)):
                day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
            else:
                parts = str(raw_day).strip().split('.')
                month, day_number = int(parts[0]), int(parts[1])
                candidates = []
                for d in [day_number]+([day_number*10] if len(parts[1]) == 1 else []):
                    try:
                        candidate = date(year, month, d)
                        if candidate >= previous_day:
                            candidates.append(candidate)
                    except ValueError:
                        pass
                if not candidates:
                    raise ValueError('日期倒序或无法解析，请人工核实')
                day = min(candidates)
                if day.day != day_number:
                    fixes.append(dict(row=row_number, original=raw_day, resolved=day.isoformat(),
                                      reason='Excel数字省略末尾0；按日期段顺序消歧'))
            previous_day = day
            if not isinstance(raw_time, time):
                raise ValueError('采样时刻不是Excel时间')
            values = list(row[2:10])+[row[11]]
            for i in range(2, 8):
                values[i] = {'开':1, '停':0, '关':0}.get(values[i], values[i])
            record = dict(zip(FEATURES, values))
            record.update(ash_content=row[10], day=day.isoformat(), source_row=row_number,
                          timestamp=f'{day.isoformat()} {raw_time.isoformat()}')
            if not all(number(v) for v in [*values, row[10]]):
                raise ValueError('缺失/非数值因子或目标')
            if not (0 < record['raw_ash'] <= 100 and 0 < row[10] <= 100 and
                    0 <= record['level'] <= 100 and record['coal_amount'] >= 0 and
                    all(v in (0, 1) for v in values[2:8])):
                raise ValueError('因子或目标超出物理/开关范围')
            if record['coal_amount'] <= 10:
                stopped.append(record)
            else:
                valid_rows.append(record)
        except (ValueError, TypeError, IndexError) as e:
            rejected.append(dict(row=row_number, reason=str(e)))
    workbook.close()
    # Ambiguous duplicates are not silently collapsed, relabelled or averaged.
    times = Counter(r['timestamp'] for r in valid_rows)
    duplicates = {t for t,n in times.items() if n > 1}
    if duplicates:
        rejected.extend(dict(row=r['source_row'], reason='同采样时刻存在多条记录，待核实')
                        for r in valid_rows if r['timestamp'] in duplicates)
        valid_rows = [r for r in valid_rows if r['timestamp'] not in duplicates]
    valid_rows.sort(key=lambda r:(r['day'], r['timestamp'], r['source_row']))
    audit = dict(source=str(path.resolve()), sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                 sheet='Sheet1', range=f'A4:L{sheet.max_row}', factor_range='C:J,L:L', target_column='K',
                 assumed_year=year, original_rows=sum(any(v is not None for v in r) for r in all_rows[3:]),
                 usable_rows=len(valid_rows), stopped_rows=stopped, rejected_rows=rejected, date_repairs=fixes,
                 dates=sorted({r['day'] for r in valid_rows}),
                 chronology_note='沿用表中生产日期；同日夜班时刻不推断跨日，所有时间验证按完整日期分组。',
                 ranges={k:[min(r[k] for r in valid_rows),max(r[k] for r in valid_rows)] for k in FEATURES+['ash_content']})
    return valid_rows, audit

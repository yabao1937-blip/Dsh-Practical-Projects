# -*- coding: utf-8 -*-
"""生成 Excel：前后端分离（FastAPI）改造规划与建议"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\前后端分离-FastAPI改造规划.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def make_sheet(wb, title, headers, rows, widths, first_col_fill=True):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER
            if first_col_fill and cell.column == 1:
                cell.fill = SEC_FILL
                cell.font = SEC_FONT
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws


wb = Workbook()
wb.remove(wb.active)

# ---------------- Sheet 1 总体架构与目标 ----------------
make_sheet(wb, "1总体架构与目标",
    ["板块", "内容"],
    [
        ["目标", "把当前『原生JS + localStorage』的纯前端系统，拆成 FastAPI 后端 + 前端 分离架构，数据和计算下沉到后端"],
        ["前端定位", "保留现有 6 个页面（密度推荐总览/数据采集补录/粗精煤泥分析/浮精影响分析/批量导入/历史查询）与 Canvas 图表，改为通过 HTTP API 读写"],
        ["后端定位", "FastAPI + SQLAlchemy 承担：持久化、三表导入解析、模型训练(MLR/PLS)、建议密度计算、推测简报生成、告警计算、手动有效期判定"],
        ["部署形态", "开发期由 FastAPI 托管前端静态资源（同源，规避 file:// 与 CORS 问题）；生产可 Nginx 静态 + /api 反向代理"],
        ["数据迁移", "localStorage 一次性导出并导入 SQLite；之后 localStorage 不再作为主存储"],
        ["兼容策略", "过渡期保留『localStorage 双写』，确认稳定后再彻底移除，避免数据丢失与回退困难"],
    ],
    [10, 80])

# ---------------- Sheet 2 现状盘点 ----------------
make_sheet(wb, "2现状盘点",
    ["分类", "对象/键", "现状", "去向"],
    [
        ["前端数组", "magneticTail / coarseCoal / floatCoal / calcLogs / manualEntries / importLogs / alerts / heavySamples / regressionModels / coarseModelHistory", "localStorage 数组，导入/补录写入", "后端 SQLAlchemy 表"],
        ["前端单值", "coarseModel / coarseTolerance / coarseTrainRange", "当前粗灰模型(mlr/pls)与训练参数", "后端 coarse_models 表 + settings"],
        ["前端单值", "ashTarget / ashTargetTol / guideScheme / totalAshManualOn / densityGuide / densityAutoOn", "总览/建议密度相关配置", "后端 settings 表"],
        ["前端键值", "amountInputs / ashInputs / instrumentInputs / heavyAshInput / coarseAshInput / floatAshInput / coarseCalc / coarseAshEma / autoState / heavyAshBackcalc", "录入层+手动有效期机制，含 manual/manualAt/auto 打点", "后端 input_state 表(键值)或专用表"],
        ["前端逻辑", "calcTotalAsh / resolveTotalAsh / formulaTotalAsh / computeDensityGuidance / expertAdjust / densityGainK", "总灰分公式 + 专家表建议密度", "后端 services/density.py"],
        ["前端逻辑", "trainCoarseModel / trainMlr / trainPls / predictCoarseAsh / polyRegression / _timeCvQ2", "MLR/PLS 训练与预测", "后端 services/modeling.py（numpy 重写）"],
        ["前端逻辑", "buildHourlyBrief", "推测简报(1h对齐+递归软测量)", "后端 services/brief.py"],
        ["前端逻辑", "ImportPage.parseCoarseFactors / normalizeTs / parseRawData", "三表 XLSX 解析", "后端 services/importer.py（openpyxl）"],
        ["后端现状", "run.py / config.py / database.py", "已有 uvicorn 入口、SQLite URL、SQLAlchemy engine/SessionLocal/Base/get_db", "保留并补齐"],
        ["后端缺失", "main.py / models.py / schemas.py / routers/* / services/* / 迁移", "尚未实现", "待新建"],
    ],
    [12, 40, 46, 26])

# ---------------- Sheet 3 分层职责划分 ----------------
make_sheet(wb, "3分层职责划分",
    ["层", "职责", "说明"],
    [
        ["前端-展示", "页面渲染、Canvas 图表(chart-lite)、交互、表单校验、导出 Excel(xlsx-lite 客户端导出)", "保留，几乎不动"],
        ["前端-状态", "会话态(当前页/筛选/弹窗)、乐观更新提示", "从 localStorage 迁移到内存+API"],
        ["前端-数据访问", "统一 fetch 封装(api.js)，替代 App.store 直读直写", "新增，薄封装"],
        ["后端-持久化", "所有数据落库(SQLite→PostgreSQL)，替代 localStorage", "SQLAlchemy ORM"],
        ["后端-导入", "三表 XLSX 解析、去重、校验、写库、导入日志", "openpyxl 重写 parse 逻辑"],
        ["后端-计算", "总灰分公式、建议密度、专家表、递归软测量、推测简报", "迁移 App 纯函数，保持口径一致"],
        ["后端-模型", "MLR/PLS 训练、交叉验证 Q²、生产模型选择、预测", "numpy 重写，结果需与 JS 版本对齐"],
        ["后端-告警", "告警计算(checkAlerts)、手动有效期(autoState 打点)判定", "由后端统一时间戳保证一致性"],
    ],
    [12, 52, 60])

# ---------------- Sheet 4 数据库设计 ----------------
make_sheet(wb, "4数据库设计",
    ["store键", "类型", "建议表", "关键字段", "说明"],
    [
        ["coarseCoal", "数组", "coarse_coal", "id,timestamp,system,ash_content,coal_amount,level,raw_ash,moisture,sysA..sys402,desliming473/474,is_stoppage,mining_face,influence_value,predicted_ash", "表1 粗精煤泥多因素"],
        ["floatCoal", "数组", "float_coal", "id,timestamp,system,ash_content,coal_amount,filter_press_running,influence_value,annotation", "表2 浮精"],
        ["calcLogs", "数组", "calc_logs", "id,timestamp,calc_type,input_json,output_json", "表3 灰分密度 + 灰分仪/密度计/皮带秤补录"],
        ["magneticTail", "数组", "magnetic_tail", "id,timestamp,system,level,ash_content,moisture,coal_amount,source", "精磁尾(镜像表1)"],
        ["manualEntries", "数组", "manual_entries", "id,timestamp,category,values(json),remark,operator,status", "补录历史"],
        ["importLogs", "数组", "import_logs", "id,timestamp,category,file_name,total,success,failed,skipped,status,errors(json)", "导入日志"],
        ["alerts", "数组", "alerts", "id,system,level,message,time", "告警"],
        ["heavySamples", "数组", "heavy_samples", "id,timestamp,rho,ash_content,source", "重介采样闭环"],
        ["regressionModels", "数组", "regression_models", "id,model_type,params(json),r_squared,rmse,created_at,data_points", "单因素模型"],
        ["coarseModelHistory", "数组", "coarse_model_history", "id,trained_at,n,tolerance,range,production,mlr(json),pls(json)", "重训练历史"],
        ["coarseModel", "单值", "coarse_models", "type,intercept,coefs(json),means(json),stds(json),std_coef(json),metrics(json),impute_means(json),A,production,is_active,trained_at,n,tolerance,range", "当前生效粗灰模型(mlr+pls)"],
        ["ashTarget/ashTargetTol/guideScheme/totalAshManualOn/densityGuide/densityAutoOn/coarseTolerance/coarseTrainRange", "单值", "settings", "key,value_json", "建议密度/训练配置"],
        ["instrumentInputs/amountInputs/ashInputs/heavyAshInput/coarseAshInput/floatAshInput/coarseCalc/coarseAshEma/autoState/heavyAshBackcalc", "键值", "input_state", "key,value_json,updated_at", "录入层+手动有效期打点(推荐键值表，少表少迁移)"],
    ],
    [26, 10, 18, 46, 24])

# ---------------- Sheet 5 API设计 ----------------
make_sheet(wb, "5API设计",
    ["方法", "路径", "用途", "对应前端", "说明"],
    [
        ["GET", "/api/health", "健康检查", "—", "探活"],
        ["GET", "/api/state", "全量当前状态快照", "App.loadStore", "首屏加载，替代 localStorage"],
        ["GET", "/api/instruments", "在线仪表当前值(13项)", "CollectPage.renderTable", "含来源层级"],
        ["PUT", "/api/instruments/{key}", "仪表手动覆盖", "App.setInstrumentInput", "key=ash_501/ash_502/scale_501/scale_502/density/level_tail"],
        ["PUT", "/api/amounts/{key}", "量数据录入/切模式", "App.setAmountInput", "key=totalAmount/floatAmount/coarseAmount/denseAmount"],
        ["PUT", "/api/ash/{kind}", "灰分录入", "setAshInput/setCoarseAshInput/setFloatAshInput/setHeavyAshInput", "kind=total/coarse/float/heavy"],
        ["PUT", "/api/settings", "配置(目标灰分/容差/版本/钳制/训练范围)", "applyDensityGuideSettings/toggleGuideScheme", "ashTarget 等"],
        ["POST", "/api/import/parse", "上传三表XLSX→解析预览", "ImportPage.processFile", "multipart，返回 headers/rows/errors/duplicates"],
        ["POST", "/api/import/confirm", "确认导入写库", "ImportPage.confirmImport/confirmImportFactors", "写表+触发训练"],
        ["POST", "/api/models/coarse/train", "训练粗灰模型(MLR/PLS)", "App.trainCoarseModel", "重计算，返回 R²/Q²/合格率"],
        ["GET", "/api/models/coarse", "当前模型+重训练历史", "CoarsePage", "含生产模型标记"],
        ["GET", "/api/guidance", "建议密度计算", "App.computeDensityGuidance", "返回 deltaA/deltaRho/rhoNew/reason"],
        ["GET", "/api/brief", "推测简报(1h对齐)", "App.buildHourlyBrief", "返回 {headers,rows}，前端渲染/导出"],
        ["POST", "/api/samples/heavy-ash", "重介灰分采样", "CollectPage.submitManual(heavy_ash_sample)", "写 heavy_samples + 更新密度/重介灰分"],
        ["GET", "/api/history", "历史查询", "HistoryPage.query", "分页/过滤"],
        ["GET", "/api/export", "全量备份导出", "App.exportData", "生成 xlsx/csv"],
        ["POST", "/api/import-backup", "备份恢复", "App.importBackup", "覆盖式导入"],
    ],
    [8, 24, 30, 34, 30])

# ---------------- Sheet 6 业务逻辑迁移清单 ----------------
make_sheet(wb, "6业务逻辑迁移清单",
    ["前端函数", "后端模块", "说明", "风险"],
    [
        ["calcTotalAsh/formulaTotalAsh/resolveTotalAsh/totalAshEntry", "services/density.py", "总精煤灰分公式与取值链(手动>公式>录入>默认)", "取值链层级多，需逐级对齐"],
        ["expertAdjust/EXPERT_ADJUST/K_PREDICT/computeDensityGuidance/densityGainK", "services/density.py", "专家表建议密度(总灰分版/重介版)", "两版本口径+夹紧1.35~1.60"],
        ["resolveCoarseAsh/_emaCoarseAsh/predictCoarseAsh", "services/modeling.py", "粗灰预测(EMA平滑)", "EMA跨>6h重置逻辑"],
        ["trainCoarseModel/trainMlr/trainPls/polyRegression/_timeCvQ2/solveLinearSystem", "services/modeling.py", "MLR(岭回归)+PLS(NIPALS)+时间序列CV", "numpy 重写需与 JS 数值对齐(浮点/岭λ网格)"],
        ["buildHourlyBrief(递归软测量)", "services/brief.py", "1h对齐+前向填充+采样反馈", "时间戳排序/前向填充游标"],
        ["resolveAmount/resolveInstrument/resolveFloatAsh/resolveHeavyAsh", "services/resolvers.py", "各录入层取值链", "手动有效期 manualAt 与 autoState 打点"],
        ["_autoBump/_manualValid(手动有效期)", "services/validity.py", "自动值变化接管更早手工值", "服务器时间戳统一，避免时钟漂移"],
        ["checkAlerts", "services/alerts.py", "告警计算", "阈值/文案一致"],
        ["ImportPage.parseCoarseFactors/normalizeTs/parseRawData", "services/importer.py", "三表解析(表1多Sheet/日期补偿/去重校验)", "M.DD 语义补偿规则要原样移植"],
    ],
    [34, 22, 40, 30])

# ---------------- Sheet 7 前端改造 ----------------
make_sheet(wb, "7前端改造",
    ["改造点", "现状", "目标", "说明"],
    [
        ["数据访问层", "App.store 直读直写 + saveStore()", "新增 api.js：fetch 封装，App.store 改为内存镜像", "最小侵入：保留 App 对象形状，底层换成 API"],
        ["持久化", "localStorage.setItem('dmcs_store')", "移除 localStorage，读 /api/state 写对应 PUT/POST", "过渡期可双写"],
        ["导入", "前端 XLSX.read 解析", "上传文件到 /api/import/parse，后端解析返回预览", "前端只展示预览与确认"],
        ["模型训练", "前端 JS 训练 MLR/PLS", "POST /api/models/coarse/train，前端只展示结果", "把重计算从主线程移走"],
        ["建议密度/简报", "前端 computeDensityGuidance/buildHourlyBrief", "GET /api/guidance、/api/brief", "口径由后端统一"],
        ["图表", "chart-lite.js Canvas", "保留", "无需改动"],
        ["导出", "xlsx-lite.js 客户端写文件", "简报/历史可由后端出 xlsx 或前端继续用 xlsx-lite", "二选一，建议后端出 xlsx"],
        ["运行方式", "file:// 直接打开", "http://localhost:8000(后端托管静态) 或 http://127.0.0.1:5500 开发服务器", "file:// 下 fetch 受 CORS 限制"],
    ],
    [14, 34, 40, 36])

# ---------------- Sheet 8 分阶段实施计划 ----------------
make_sheet(wb, "8分阶段实施计划",
    ["阶段", "目标", "内容", "验收标准"],
    [
        ["P1 后端骨架", "跑通 FastAPI", "补 main.py(FastAPI+CORS+StaticFiles)、models.py、schemas.py、Alembic 迁移、/api/health", "uvicorn 启动、/api/health 返回 ok、/docs 可访问"],
        ["P2 数据层 API", "localStorage → 数据库", "建全部 ORM 表；实现 /api/state、各 PUT/POST CRUD；localStorage 一次性迁移脚本", "重启后数据不丢、前端可读回全量状态"],
        ["P3 计算下沉", "逻辑搬后端", "迁移 density.py/modeling.py/brief.py/resolvers.py/alerts.py；numpy 重写训练与预测", "后端算出的建议密度/简报与前端 JS 结果逐值对齐"],
        ["P4 前端接 API", "前端切后端", "api.js 封装；App.store 改为内存镜像；导入/训练/简报走接口；保留图表", "页面功能等价、headless 回归全绿"],
        ["P5 导入流水线", "服务端解析", "openpyxl 重写三表解析(含 M.DD 补偿/多Sheet/去重)；上传预览+确认写库", "导入预览与结果与原版一致"],
        ["P6 验证与上线", "回归+部署", "把 verify_web2_*.js 改成打 http://localhost:8000；补充 pytest/httpx 单测；部署文档", "全量回归 PASS，前端由后端静态托管可访问"],
    ],
    [12, 18, 52, 40])

# ---------------- Sheet 9 技术选型 ----------------
make_sheet(wb, "9技术选型",
    ["类别", "选型", "用途", "备注"],
    [
        ["后端框架", "FastAPI + uvicorn", "API 服务", "异步、自动 OpenAPI 文档"],
        ["ORM", "SQLAlchemy 2.0", "数据访问", "已引入；配 Alembic 迁移"],
        ["校验", "Pydantic v2", "请求/响应模型", "schemas.py"],
        ["数据库", "SQLite(开发) / PostgreSQL(生产)", "持久化", "DATABASE_URL 环境变量切换"],
        ["数值计算", "numpy", "MLR/PLS 训练与预测", "重写 JS 线性代数；保证与 JS 结果对齐"],
        ["Excel解析", "openpyxl", "三表导入解析/导出", "替代前端 xlsx-lite 的解析部分"],
        ["文件上传", "python-multipart", "multipart 上传", "导入文件上传"],
        ["跨域", "CORSMiddleware", "开发跨域", "生产同源可去掉"],
        ["前端", "原生 JS(无构建)", "保留现有页面", "不引入框架，降低迁移成本"],
        ["测试", "pytest + httpx + headless Edge", "后端单测 + 前端回归", "沿用现有 verify_web2_* 思路"],
    ],
    [12, 28, 34, 40])

# ---------------- Sheet 10 风险与建议 ----------------
make_sheet(wb, "10风险与建议",
    ["类别", "风险/建议", "说明", "应对"],
    [
        ["数值一致性", "MLR/PLS 由 JS 转 numpy，浮点/岭λ网格/PLS 分量选择可能产生微小差异", "建议密度、预测值、简报必须逐值对齐", "迁移后用同一批数据做前后端结果 diff，差异 < 1e-6 才切换"],
        ["口径漂移", "取值链(手动>公式>录入>默认)与手动有效期逻辑复杂", "多页面共用同一口径", "后端 resolver 单点实现 + 单测覆盖 9 条影响链"],
        ["时间戳", "手动有效期 autoState 依赖 Date.now()", "前后端时钟可能不一致", "autoState 打点统一用服务器时间，前端不再自行打点"],
        ["CORS/file://", "前端原为 file:// 打开，fetch 会受限", "切 http 才能调 API", "由 FastAPI StaticFiles 托管前端，同源访问；或本地开发服务器"],
        ["数据迁移", "localStorage → SQLite 一次性迁移", "历史数据不能丢", "先导出 backup，再导入后端；保留回退开关"],
        ["导入解析", "表1 多Sheet/M.DD 日期补偿/去重校验较特殊", "openpyxl 重写易漏细节", "把 import.js 的 extractDate/fixDay/normalizeTs 逐条移植并写测试"],
        ["并发写入", "多页面/定时刷新同时写", "localStorage 单线程不存在此问题", "后端用事务+乐观锁；写接口幂等"],
        ["建议-不要推倒重来", "先做『数据访问层替换』，App 对象形状尽量不变", "前端改动最小、可逐步灰度", "P2 先双写 localStorage+API，稳定后移除"],
        ["建议-先模型后UI", "模型/建议密度/简报是核心计算，先下沉并验证数值", "避免 UI 先改导致口径对不上", "P3 先做 services 单测，再切前端"],
        ["建议-保留回归资产", "现有 verify_web2_*.js 是宝贵回归", "改造后要能跑通", "把 URL 从 file:// 改成 http://localhost:8000，断言不变"],
    ],
    [12, 44, 40, 40])

# ---------------- Sheet 11 验证方案 ----------------
make_sheet(wb, "11验证方案",
    ["类别", "内容", "方式", "说明"],
    [
        ["后端单测", "density/modeling/brief/resolver 各函数", "pytest 断言", "与 JS 已知结果对齐(如专家表 0.15→0.01/0.25→0.02)"],
        ["数值对齐", "同一批三表数据，前端 JS 与后端 numpy 结果 diff", "脚本对比 CSV", "建议密度、预测粗灰、简报逐值 diff"],
        ["API 契约", "各端点请求/响应 schema", "httpx + OpenAPI", "校验 Pydantic 模型与状态码"],
        ["导入回归", "三表模板导入→预览→确认", "pytest(openpyxl 样例) + headless", "预览行数/去重/校验与原版一致"],
        ["前端回归", "6 页面功能等价", "headless Edge(verify_web2_* 改造)", "URL 改 http://localhost:8000，断言不变"],
        ["数据迁移", "localStorage 备份 → SQLite", "迁移脚本 + 校验条数", "各表记录数一致、模型参数一致"],
    ],
    [12, 40, 34, 40])

wb.save(OUT)
print("[gen]", OUT)

"""生成《浮精信息表-构建疑问清单》Excel"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parents[2] / "docs" / "浮精信息表-构建疑问清单.xlsx"

# 序号 | 类别 | 疑问 | 我的建议 | 你的回复（待填）
ROWS = [
    (1, "药剂量", "数据来源与单位？PLC（加药泵频率/流量）还是人工记录？单位是 kg/h、kg/t 干煤泥还是泵频率 Hz？",
     "请告知实际来源与单位", ""),
    (2, "药剂量", "采样周期？",
     "PLC 来源→按小时累计；人工记录→每班一次", ""),
    (3, "药剂量", "「药剂量较大才会影响浮精」的阈值与影响方向？现在就要参与计算吗？",
     "建议先只存值+标注，不参与计算；等数据积累后再定阈值建模", ""),
    (4, "浮精量", "压滤机循环数如何换算成吨？每循环产率是多少 t/循环？",
     "请提供：每循环吨数，或设备参数（滤板面积×滤饼厚度×密度）", ""),
    (5, "浮精量", "循环数的数据来源与粒度？PLC 计数器？按小时读数还是按天累计？分系统（A/B/C/D）吗？压滤机几台、对应哪个系统？",
     "请告知循环数来源与所属系统", ""),
    (6, "浮精量", "与浮精导入版 Excel 中「浮精煤量(t/h)」列的关系？",
     "建议：Excel 煤量作为实测值入库（1天周期），循环数换算值作为 PLC 值入库，两者并存、卡片内对照", ""),
    (7, "浮精灰分", "Excel 已有浮精灰分化验值（8.38~10.43），灰分暂定 11 如何与化验值共存？",
     "建议：有化验值时用化验值，缺失时用 11 兜底", ""),
    (8, "浮精灰分", "灰分可信度低，误差标注 ±2.0% 可以吗？",
     "建议：可信度「低」、误差 ±2.0%，与人工采样同级", ""),
    (9, "表结构", "是否沿用粗精煤泥表的整套模式？",
     "建议沿用：小时记录存储 + 按天视图（时间=月-日）+ 隐藏卡片（采样时间/采样设备/采样值，内联编辑/新增/删除）+ 可信度/误差 + 表头竖排周期/可信度/误差 + 浮精 Excel 导入（401→A、402→B 归一化）", ""),
    (10, "表结构", "原煤煤泥含量与粗精煤泥表的关系？",
     "建议：直接从粗精煤泥表引用/同步，浮精表存一份副本并标注来源「引用粗精煤泥表」；或仅展示不落库（请二选一）", ""),
]

PRESET = [
    ("原煤煤泥含量", "引用粗精煤泥表", "1天", "低", "±2.0%"),
    ("药剂量", "待定（疑问1）", "待定", "待定", "待定"),
    ("浮精量（Excel实测）", "化验", "1天", "高", "±0.3%"),
    ("浮精量（循环数换算）", "PLC", "1小时（累计）", "高", "±1.0%"),
    ("浮精灰分", "人工采样", "1天", "低", "±2.0%（暂定11）"),
]


def build():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    cat_fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")

    # Sheet1：疑问清单
    ws = wb.active
    ws.title = "疑问清单"
    cols = ["序号", "类别", "疑问", "我的建议", "你的回复（待填）"]
    widths = [6, 10, 46, 44, 26]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in ROWS:
        ws.append(list(row))
    for r in ws.iter_rows(min_row=2, max_row=1 + len(ROWS)):
        for c in r:
            c.border = border
            c.alignment = wrap
        r[0].alignment = Alignment(horizontal="center", vertical="center")
        r[1].alignment = Alignment(horizontal="center", vertical="center")
        r[1].fill = cat_fill
    ws.freeze_panes = "A2"

    # Sheet2：采样周期预设（待确认）
    ws2 = wb.create_sheet("采样周期预设(待确认)")
    cols2 = ["字段", "数据来源", "建议采样周期", "可信度", "误差"]
    widths2 = [24, 18, 16, 10, 18]
    ws2.append(cols2)
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for c in ws2[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in PRESET:
        ws2.append(list(row))
    for r in ws2.iter_rows(min_row=2, max_row=1 + len(PRESET)):
        for c in r:
            c.border = border
            c.alignment = wrap
    ws2.freeze_panes = "A2"

    wb.save(OUT_PATH)
    print(f"[gen] 已生成：{OUT_PATH}")


if __name__ == "__main__":
    build()

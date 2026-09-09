"""生成《粗精煤泥灰分预测-数据补充建议》Excel"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parents[2] / "docs" / "粗精煤泥灰分预测-数据补充建议.xlsx"

# 列：序号 | 优先级 | 数据项 | 单位 | 数据来源 | 建议采样周期 | 与粗精煤泥灰分的关联（原因） | 当前状态
ROWS = [
    (1, "第一优先级", "原煤粒度组成 / 末煤灰分（<0.5mm 或 <0.25mm 含量及灰分）", "%",
     "人工化验 / 在线粒度仪", "每班或每天",
     "煤泥本质是细粒级产物，细粒级含量及其灰分基本决定煤泥的量和灰分水平，比单一「煤泥含量」更本质",
     "未采集"),
    (2, "第一优先级", "精磁尾流量（尾矿阀开度/流量计）", "m³/h",
     "PLC", "在线（1分钟）",
     "精磁尾直接汇入煤泥系统，是煤泥灰分的重要来源；目前只有液位一个间接量，加上流量信号强得多",
     "未接入"),
    (3, "第一优先级", "精磁尾浓度", "% / g/L",
     "化验", "每班",
     "浓度×流量=精磁尾干量，与灰分共同决定进入煤泥系统的灰分量",
     "未采集"),
    (4, "第一优先级", "分级旋流器入料压力", "MPa",
     "PLC", "在线（1分钟）",
     "分级效率决定哪些灰分的颗粒进入煤泥：压力高→分级细→高灰细泥比例上升，是灰分日内波动的直接推手",
     "未接入"),
    (5, "第一优先级", "补水量 / 循环水量", "m³/h",
     "PLC（流量计）", "在线（1分钟）",
     "影响矿浆稀释与分级效率，进而改变煤泥灰分分布",
     "未接入"),
    (6, "第一优先级", "315筛筛下水灰分", "%",
     "化验", "每班",
     "物料平衡中间量：入料灰分≈筛上灰分×产率+筛下水灰分×比例，可用于标定与校验模型",
     "未采集"),
    (7, "第一优先级", "315筛筛下水浓度", "g/L",
     "化验", "每班",
     "与筛下水量结合得出筛下细泥量，验证煤泥量物料平衡",
     "未采集"),
    (8, "第一优先级", "重介悬浮液密度", "g/cm³",
     "PLC（灰分、密度导入表已有）", "5分钟",
     "密度波动反映重介工况与错配物变化，间接影响尾矿与煤泥灰分；该数据已有，应最先利用",
     "已有，待接入"),
    (9, "第二优先级", "脱出粉煤的灰分", "%",
     "化验", "每天或每班",
     "脱粉量已知，但脱出物的灰分决定脱粉对煤泥灰分的影响方向（脱高灰还是低灰粉煤，效果相反）",
     "未采集"),
    (10, "第二优先级", "煤种 / 煤层来源", "文本",
     "人工录入", "每批次",
     "不同煤层的煤泥灰分规律差异大，作为分类特征可显著提升模型",
     "未记录"),
    (11, "第二优先级", "启停车 / 故障状态标记", "0/1",
     "PLC 状态位", "在线",
     "非稳态时段的数据会污染训练样本，标记后可在训练时剔除",
     "未接入"),
    (12, "第二优先级", "各过程点到315筛的滞后时间标定", "分钟",
     "数据分析（互相关）", "一次性标定",
     "化验时刻对应的应是滞后窗口前的输入，滞后对齐是训练质量的地基",
     "未做"),
    (13, "第三优先级", "化验采样的精确对应记录", "文本",
     "人工记录", "每次取样",
     "记录取样时刻与所对应生产时段，避免窗口错位",
     "未记录"),
]

EXISTING = [
    ("原煤灰分", "%", "灰分仪", "在线", "细粒煤泥灰分随原煤灰分整体抬升"),
    ("原煤煤泥含量", "%", "人工采样", "1天", "直接表征原煤中细泥含量，人工采样误差较大"),
    ("入洗原煤量", "t", "PLC", "1小时", "处理量影响煤泥产率与稀释程度"),
    ("脱粉量", "t", "PLC", "1小时", "震动脱粉脱出碎煤，改变入洗物料性质"),
    ("精磁尾液位", "%", "PLC", "1小时", "间接反映精磁尾及煤泥水系统负荷"),
]


def build():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    prio_fills = {
        "第一优先级": PatternFill("solid", fgColor="FCE4D6"),
        "第二优先级": PatternFill("solid", fgColor="FFF2CC"),
        "第三优先级": PatternFill("solid", fgColor="E2EFDA"),
    }
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")

    # Sheet1：建议补充的数据
    ws = wb.active
    ws.title = "建议补充数据"
    cols = ["序号", "优先级", "数据项", "单位", "数据来源", "建议采样周期", "与粗精煤泥灰分的关联（原因）", "当前状态"]
    widths = [6, 12, 38, 10, 18, 12, 60, 12]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in ROWS:
        ws.append(list(row))
    for r in ws.iter_rows(min_row=2, max_row=1 + len(ROWS)):
        for c in r:
            c.border = border
            c.alignment = wrap
        ws.cell(row=r[0].row, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r[0].row, column=2).alignment = Alignment(horizontal="center", vertical="center")
        prio = r[1].value
        if prio in prio_fills:
            ws.cell(row=r[0].row, column=2).fill = prio_fills[prio]
    ws.freeze_panes = "A2"

    # Sheet2：现有预测因子（对照）
    ws2 = wb.create_sheet("现有预测因子")
    cols2 = ["字段", "单位", "数据来源", "采样周期", "在预测中的作用"]
    widths2 = [16, 8, 14, 12, 46]
    ws2.append(cols2)
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for c in ws2[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in EXISTING:
        ws2.append(list(row))
    for r in ws2.iter_rows(min_row=2, max_row=1 + len(EXISTING)):
        for c in r:
            c.border = border
            c.alignment = wrap
    ws2.freeze_panes = "A2"

    wb.save(OUT_PATH)
    print(f"[gen] 已生成：{OUT_PATH}")


if __name__ == "__main__":
    build()

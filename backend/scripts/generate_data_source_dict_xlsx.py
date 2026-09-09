"""生成《数据来源三表-字段与缺口清单》Excel"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parents[2] / "docs" / "数据来源三表-字段与缺口清单.xlsx"

# Sheet1：三表字段字典（表 | 列名 | 含义 | 对应系统字段 | 备注）
DICT_ROWS = [
    ("表1 粗精煤泥灰分影响因素6.16-7.14.xlsx", "日期/序号/时间", "采样时刻（5分钟间隔）", "timestamp", "113条，6.16~7.14"),
    ("表1", "入洗工作面", "煤层/工作面", "mining_face(展示)", "如 3309、43下01"),
    ("表1", "原煤灰分(%)", "入选原煤灰分", "raw_ash", "前向填充"),
    ("表1", "小时带煤量(t/h)", "入选原煤量", "coal_amount", "缺值按0"),
    ("表1", "开启的系统(A/B/401/402)", "四系统运行标志", "running_systems(合并为数量)", "已合并为1系统"),
    ("表1", "脱粉(473/474)", "脱粉筛开停", "desliming473/474", "开=1"),
    ("表1", "315灰分(%)", "粗精煤泥灰分(目标值)", "ash_content", "模型训练目标"),
    ("表1", "315全水分(%)", "粗精煤泥水分", "moisture", ""),
    ("表1", "精磁尾液位(%)", "精磁尾池液位", "level", "滞后采样前90秒"),
    ("表2 浮精 (1).xlsx", "采样时间", "浮精采样时刻", "timestamp", "5条，5.21~5.25；新版无系统列"),
    ("表2", "浮精灰分%", "浮选精煤灰分", "float ash_content", ""),
    ("表2", "浮精煤量(t/h)", "浮精量", "float coal_amount", "27/26 t/h"),
    ("表2", "压滤机运行", "压滤机状态", "filter_press_running", "运行=1"),
    ("表3 灰分、密度（导入版）.xlsx", "采样时间", "灰分/密度采样时刻", "timestamp", "124条，5.21~5.26"),
    ("表3", "系统(A/B)", "所属系统（唯一仍带系统的表）", "忽略/合并", "待确认是否忽略"),
    ("表3", "皮带(502)", "灰分仪所在皮带", "belt", "仅502"),
    ("表3", "灰分%", "皮带总精煤灰分", "belt_502_ash", "7.4~8.7"),
    ("表3", "密度值(g/cm³)", "重介悬浮液密度", "density", "1.38~1.49"),
]

# Sheet2：缺口（三表没有、需录入或计算 —— 对应任务二/三）
GAP_ROWS = [
    ("粗精煤泥量", "三表均无", "录入（任务二：录入+输入样式）", "表1只有315灰分/水分/液位，没有315煤量"),
    ("重介精煤量", "三表均无", "录入，或=总精煤量−浮精量−粗精煤泥量", "任务二"),
    ("总精煤量/皮带量", "三表均无", "录入（501+502皮带量）", "任务二"),
    ("总精煤灰分", "三表均无", "任务三公式计算：(重介灰分×重介量+浮精灰分×浮精量+粗精煤泥量×粗精煤泥灰分)÷总煤量", "后续获实测值可反推重介精煤灰分"),
    ("重介精煤灰分", "三表均无", "任务三反推：=(总灰分×总煤量−浮精灰分×浮精量−粗精煤泥量×粗精煤泥灰分)÷重介精煤量", "需重介精煤量>0"),
]


def build():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")

    ws = wb.active
    ws.title = "三表字段字典"
    cols = ["表", "列名", "含义", "对应系统字段", "备注"]
    widths = [30, 22, 26, 24, 26]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in DICT_ROWS:
        ws.append(list(row))
    for r in ws.iter_rows(min_row=2, max_row=1 + len(DICT_ROWS)):
        for c in r:
            c.border = border
            c.alignment = wrap
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("缺口(需录入或计算)")
    cols2 = ["量", "三表是否提供", "处理方式", "说明"]
    widths2 = [18, 14, 52, 52]
    ws2.append(cols2)
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for c in ws2[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in GAP_ROWS:
        ws2.append(list(row))
    for r in ws2.iter_rows(min_row=2, max_row=1 + len(GAP_ROWS)):
        for c in r:
            c.border = border
            c.alignment = wrap
    ws2.freeze_panes = "A2"

    wb.save(OUT_PATH)
    print(f"[gen] 已生成：{OUT_PATH}")


if __name__ == "__main__":
    build()

# -*- coding: utf-8 -*-
"""更新 待办-后续优化计划.xlsx:
① 新增「第二轮审查意见」Sheet(事实核验/文档问题/盲区/决策倾向/执行顺序);
② 修正总览页计数不一致(8/10/11 三处对不上 → 统一为 11);
③ Q-6 标注"已决策"(用户已选定严格口径,避免重复折腾)。
只改文档,不执行任何待办项。
"""
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

SRC = Path(r"D:\dense-medium-density-control-system\docs\待办-后续优化计划.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
NOTE_FONT = Font(name="微软雅黑", size=9, italic=True, color="C00000")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

# 文件锁探测(Excel 打开时拒绝写入,提前报错而不是半途而废)
try:
    with open(SRC, "ab"):
        pass
except OSError:
    print("错误: 文件被占用(Excel 打开中),请先关闭后重跑")
    sys.exit(1)

wb = load_workbook(SRC)

# ---------- ② 修正总览计数 ----------
ws1 = wb["总览与建议批次"]
for row in ws1.iter_rows(min_row=2):
    label = str(row[0].value or "")
    if label == "本表范围":
        row[1].value = ("基线之后剩余的全部待办:11 项可直接执行(T-1~T-11)、8 项需你决策、4 项建议暂缓。"
                        "每项都写明「原因」,原因是判断优先级和取舍的依据。"
                        "(2026-09-10 二轮审查修正:原 8/10 两处计数与清单 11 项不符)")
    elif label == "待办分布":
        row[1].value = ("可直接执行 11 项:缺陷修复 2(T-1/T-2) · 健壮性 4(T-5/T-6/T-9/T-10) · 文档 1(T-3) · "
                        "测试基建 1(T-4) · 数据治理 1(T-7) · 可用性/K标定 1(T-8) · 交付 1(T-11);"
                        "需决策 8 项;建议暂缓 4 项。")

# ---------- ③ Q-6 标注已决策 ----------
ws3 = wb["待你决策的问题"]
for row in ws3.iter_rows(min_row=3):
    if str(row[0].value or "").strip() == "Q-6":
        old = row[5].value or ""
        row[5].value = (str(old) + " 【2026-09-10 二轮审查补充:此题已被用户决策——简报行存在规则选定严格口径"
                        "「三表齐全才成行(24h窗口)」并于当日上线(263行/8-23~9-3);除非现场反馈当班趋势不可用,"
                        "维持①不再重开。】")

# ---------- ① 新增「第二轮审查意见」 ----------
if "第二轮审查意见" in wb.sheetnames:
    del wb["第二轮审查意见"]
ws = wb.create_sheet("第二轮审查意见")
ws.append(["序号", "类别", "内容", "建议 / 结论"])
rows = [
    ["一1", "总体评价",
     "计划整体可信:分类(可执行/需决策/暂缓)清晰,每项有原因/证据/做法/依赖,批次划分有口径风险隔离意识;"
     "T-1 确为全场最优先(静默假数据直接进总灰分公式)。",
     "采纳,按本表修正后执行"],
    ["一2", "总体评价",
     "最大风险不在清单内而在清单外:11 个提交未推送(含大量已验证修复),多 agent 并发提交已实际发生过改动互相裹挟。",
     "推送提为「批次0」,处理 Q-1/Q-2 后立即推"],
    ["二1", "事实核验",
     "T-1 浮精灰分 || 0 兜底:属实。密度列有 1.3~1.6 校验、灰分列没有;0% 会经 resolve_float_ash 进总灰分公式。",
     "确认高优先级"],
    ["二2", "事实核验",
     "T-2 inlineStr/十六进制引用:属实,属 &#10; 修复的「修了一半」。",
     "确认,顺手补齐"],
    ["二3", "事实核验",
     "T-3 AGENTS.md 路径过期:属实(仍写桌面旧路径),本会话真实造成过困惑。",
     "确认,并顺带写入「同一时间只允许一个 agent 持有未提交改动」规则(见 四1)"],
    ["二4", "事实核验",
     "T-6 守卫只比 coal_records:属实——calc_logs/heavy_samples 仍可被旧浏览器洗掉,守卫「部分有效」。",
     "确认,多表向量比较方案合理"],
    ["二5", "事实核验",
     "T-10 两个「新旧」判据不一致(守卫按 coal_records / autoPullIfStale 按三数组):属实,可能组合出"
     "「提示导入完成但镜像被拒」的静默失败。",
     "确认;revision 归 Q-8/D-4 节奏,先统一口径"],
    ["二6", "事实核验",
     "Q-4(501 无在线值、兜底恒为常量 8.8):属实,是皮带分工改造的真实副作用——360 条记录全是 502,"
     "ash_501 恒走默认。",
     "采纳①+②(502 推算+文案标注);根本解是现场开始记录 501(见 四2)"],
    ["二7", "事实核验",
     "Q-6(简报缺表行为):问题真实,但【已被用户决策】——严格口径当日已上线。",
     "维持①,从待决策降级为已决策记录(本表已在原 Sheet 标注)"],
    ["三1", "文档问题",
     "总览页计数三处不一致:「本表范围」8 项 /「待办分布」10 项 / 实际清单 11 项。",
     "已在本表修正为 11 项(见总览页)"],
    ["三2", "文档问题",
     "T-11(推送)未归入任何批次,却是时间敏感度最高的一项。",
     "列为「批次0」:核实 git 状态 → 处理 Q-1/Q-2 → 推送,约 10 分钟"],
    ["三3", "文档问题",
     "Q-2 的拆分建议需先核实:文档对「谁的改动裹进谁的提交」的描述与 git 实际状态是否一致未经验证(git show dff835b --stat)。",
     "若 dff835b 是未推送的 HEAD 可安全拆;若历史已被多 agent 交错改写,保留混合提交不重写历史"],
    ["四1", "盲区补充",
     "多 agent 并发提交无治理规则(流程提醒提了,但没有行动项)。",
     "在 T-3 更新 AGENTS.md 时一并写入硬规则:同一时间只允许一个 agent 持有未提交改动"],
    ["四2", "盲区补充",
     "501 皮带数据缺失(0 条)只被当算法兜底问题讨论,根本解是数据侧。",
     "增加行动项:现场在表3 开始记录 501 灰分(与密度采样同表),数据到位后 Q-4 的①自动失效、直读生效"],
    ["四3", "盲区补充",
     "8-30 源数据损坏(coarse 三行灰分列填的是时间,导致该日缺行)不在任何清单里。",
     "增加行动项:待现场提供真实化验值后修源表重导;属数据治理,责任在现场"],
    ["五1", "决策倾向 Q-1/Q-2",
     "11 个提交只在本地是最大单点风险;拆分历史是锦上添花。",
     "Q-1 尽快推;Q-2 先核实再拆,拿不准就保留混合提交(内容正确>历史整洁)"],
    ["五2", "决策倾向 Q-3",
     "重复时间戳现状=后者胜+计数错。",
     "选②:显式拒绝并计 skipped、导出错误记录(化验数据静默丢失比报错危险)"],
    ["五3", "决策倾向 Q-5",
     "K 估计 valid 恒为 false(两系统斜率为负被拒),界面看不出来。",
     "选①:展示三态+systems[].k 与 R²——现场看到负斜率才会理解为什么要做阶跃实验"],
    ["五4", "决策倾向 Q-7",
     "删除→守卫拒绝→autoPull 拉回,用户无感知。",
     "选②:被拒时明确提示+显式「强制同步」按钮;不做全局 force"],
    ["五5", "决策倾向 Q-8",
     "解析搬后端的前置(T-9/T-4)未落地,且 PLC 是主线。",
     "选②:暂缓,与 D-1 一致"],
    ["六1", "执行顺序建议",
     "批次0:核实 git → Q-1/Q-2 → 推送(10min);批次1: T-1、T-2(静默假数据缺陷);"
     "批次2: T-3、T-4(基建+防坑);批次3: T-5~T-8;决策后: Q-3/Q-4/Q-5/Q-7 改动 + T-9/T-10。",
     "与原批次划分一致,仅把推送提前为批次0"],
]
for r in rows:
    ws.append(r)
ws.append([])
ws.append(["", "声明", "本 Sheet 为第二轮审查意见(2026-09-10),只做分析与建议,未执行任何待办项;"
            "对原 Sheet 的修改仅限:总览计数修正(三1)与 Q-6 已决策标注(二7)。", ""])
for i, w in enumerate([7, 13, 72, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for cell in ws[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = CENTER
    cell.border = BORDER
ws.freeze_panes = "A2"
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = CENTER if cell.column == 1 else WRAP
# 声明行红字
last = ws.max_row
for cell in ws[last]:
    if cell.value:
        cell.font = NOTE_FONT

wb.save(SRC)
print("已更新:", SRC)
print("Sheets:", wb.sheetnames)

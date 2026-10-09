"""Build a real XLSX export from immutable report rows (also used by the API)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.hourly_brief.columns import COLUMNS
from app.hourly_brief.inputs import LOCAL
from app.telemetry.domain import parse_time


def local_time(value):
    return parse_time(value).astimezone(LOCAL).strftime("%Y-%m-%d %H:%M:%S") if value else "—"


def display(value):
    if value is None:
        return "—"
    if isinstance(value, dict):
        return " / ".join(f"{key}：{'运行' if on is True else '停止' if on is False else '—'}" for key, on in value.items())
    return value


def generate_workbook(rows, feedback=()):
    wb = Workbook()
    main = wb.active
    main.title = "小时推测简报"
    sources = wb.create_sheet("数据来源与说明")
    history = wb.create_sheet("人工经验反馈")
    history.append(["反馈编号", "简报整点（北京时间）", "录入时间（北京时间）", "反馈内容"])
    by_hour = {}
    for record in sorted(feedback, key=lambda item: item.id, reverse=True):
        by_hour.setdefault(record.hour, []).append(record)
        history.append([record.id, local_time(record.hour), local_time(record.created_at), record.content])
    main.append(["整点时间（北京时间）"] + [label + (f"（{unit}）" if unit else "") for _, label, unit in COLUMNS] + ["反馈（最新）"])
    sources.append(["整点时间", "指标", "值", "单位", "来源", "采样时间", "状态", "说明"])
    for row in rows:
        entries = by_hour.get(row.hour, [])
        summary = f"{entries[0].content}\n{local_time(entries[0].created_at)} · 共 {len(entries)} 条" if entries else "—"
        main.append([local_time(row.hour)] + [display(row.snapshot["cells"][key]["value"]) for key, _, _ in COLUMNS] + [summary])
        for j, (key, label, unit) in enumerate(COLUMNS, start=2):
            data = row.snapshot["cells"][key]
            reason = data.get("reason", "")
            status = dict(good="正常", held="沿用前值", default="默认值", missing="缺失",
                          manual="手工补录", predicted="模型预测", calculated="计算值").get(data["status"], data["status"])
            sources.append([local_time(row.hour), label, display(data["value"]), unit, data.get("source", ""),
                            local_time(data.get("observed_at")), status, reason])
            item = main.cell(main.max_row, j)
            item.number_format = "0.000" if "density" in key else "0.00"
            item.comment = Comment("\n".join([status, data.get("source", ""), local_time(data.get("observed_at")), reason]), "数据来源")
            if data["status"] in ("default", "held"):
                item.font = Font(color="9A6700")
    for sheet in wb:
        sheet.freeze_panes = "B2" if sheet is main else "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 38
        for item in sheet[1]:
            item.fill = PatternFill("solid", fgColor="193456")
            item.font = Font(color="FFFFFF", bold=True)
            item.alignment = Alignment(wrap_text=True, vertical="center")
        for record in sheet.iter_rows(min_row=2):
            for item in record:
                if isinstance(item.value, str):
                    item.data_type = "s"
                item.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[record[0].row].height = 44
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22 if col == 1 else 19
    for col in ("D", "N"):
        main.column_dimensions[col].width = 40
    sources.column_dimensions["E"].width = 28
    sources.column_dimensions["F"].width = 24
    sources.column_dimensions["H"].width = 85
    main.column_dimensions[get_column_letter(len(COLUMNS) + 2)].width = 55
    history.column_dimensions["A"].width = 14
    history.column_dimensions["B"].width = 24
    history.column_dimensions["C"].width = 24
    history.column_dimensions["D"].width = 85
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

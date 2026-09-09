"""生成《三表数据接口-预留清单》Excel"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parents[2] / "docs" / "三表数据接口-预留清单.xlsx"

# 字段 | 来源表 | 存储位置 | 消费算法/功能 | 状态
ROWS = [
    ("原煤灰分(%)", "表1 影响因素", "coarseCoal.raw_ash", "粗精煤泥灰分模型(特征)", "已预留"),
    ("小时带煤量(t/h)", "表1 影响因素", "coarseCoal.coal_amount", "粗精煤泥灰分模型(特征)", "已预留"),
    ("A系统开关", "表1 影响因素", "coarseCoal.sysA", "粗精煤泥灰分模型(特征)", "已预留"),
    ("B系统开关", "表1 影响因素", "coarseCoal.sysB", "粗精煤泥灰分模型(特征)", "已预留"),
    ("401系统开关", "表1 影响因素", "coarseCoal.sys401", "粗精煤泥灰分模型(特征)", "已预留"),
    ("402系统开关", "表1 影响因素", "coarseCoal.sys402", "粗精煤泥灰分模型(特征)", "已预留"),
    ("473脱粉开关", "表1 影响因素", "coarseCoal.desliming473", "粗精煤泥灰分模型(特征)", "已预留"),
    ("474脱粉开关", "表1 影响因素", "coarseCoal.desliming474", "粗精煤泥灰分模型(特征)", "已预留"),
    ("停机/低负荷标志", "表1 影响因素", "coarseCoal.is_stoppage", "粗精煤泥灰分模型(特征)", "已预留"),
    ("精磁尾液位(%)", "表1 影响因素", "coarseCoal.level / magneticTail.level", "粗精煤泥灰分模型 + 单因素模型", "已预留"),
    ("315灰分(%)", "表1 影响因素", "coarseCoal.ash_content", "模型训练目标", "已预留"),
    ("315全水分(%)", "表1 影响因素", "coarseCoal.moisture", "预留", "已预留"),
    ("入洗工作面", "表1 影响因素", "coarseCoal.mining_face", "预留(分类特征)", "本次新增"),
    ("浮精灰分(%)", "表2 浮精", "floatCoal.ash_content", "总灰分计算 + 预留", "已预留"),
    ("浮精煤量(t/h)", "表2 浮精", "floatCoal.coal_amount", "总灰分计算 + 预留", "已预留"),
    ("压滤机运行", "表2 浮精", "floatCoal.filter_press_running", "预留", "已预留"),
    ("皮带", "表3 灰分密度", "calcLogs.input_json.belt", "预留", "已预留"),
    ("皮带灰分(%)", "表3 灰分密度", "calcLogs.input_json.ash_content", "总灰分计算 + 密度模型(预留)", "已预留"),
    ("密度(g/cm³)", "表3 灰分密度", "calcLogs.input_json.density", "密度推荐(卡片实时取数) + 密度模型(预留)", "本次接入"),
    ("系统(A/B)", "表3 灰分密度", "calcLogs.input_json.system", "内部参考", "已预留"),
]


def build():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    status_fills = {
        "本次新增": PatternFill("solid", fgColor="FCE4D6"),
        "本次接入": PatternFill("solid", fgColor="FFF2CC"),
    }
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")

    ws = wb.active
    ws.title = "接口清单"
    cols = ["字段", "来源表", "存储位置", "消费算法/功能", "状态"]
    widths = [20, 16, 40, 40, 12]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in ROWS:
        ws.append(list(row))
    for r in ws.iter_rows(min_row=2, max_row=1 + len(ROWS)):
        for c in r:
            c.border = border
            c.alignment = wrap
        r[4].alignment = Alignment(horizontal="center", vertical="center")
        if r[4].value in status_fills:
            r[4].fill = status_fills[r[4].value]
    ws.freeze_panes = "A2"
    wb.save(OUT_PATH)
    print(f"[gen] 已生成：{OUT_PATH}")


if __name__ == "__main__":
    build()

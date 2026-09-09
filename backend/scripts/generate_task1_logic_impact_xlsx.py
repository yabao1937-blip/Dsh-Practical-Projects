# -*- coding: utf-8 -*-
"""生成 Excel 文档：任务一-合并前后逻辑与算法影响"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\任务一-合并前后逻辑与算法影响.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


def add_rows(ws, rows, start):
    r = start
    for row in rows:
        if isinstance(row, str):
            ws.cell(row=r, column=1, value=row)
            for c in range(1, ws.max_column + 1):
                cc = ws.cell(row=r, column=c)
                cc.fill = SEC_FILL
                cc.font = SEC_FONT
                cc.border = BORDER
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ws.max_column)
        else:
            for c, v in enumerate(row, 1):
                cc = ws.cell(row=r, column=c, value=v)
                cc.font = BODY_FONT
                cc.alignment = WRAP if c != 1 else CENTER
                cc.border = BORDER
        r += 1
    return r


wb = Workbook()

# ================= Sheet1 合并前后逻辑对比 =================
ws = wb.active
ws.title = "合并前后逻辑对比"
ws.append(["序号", "维度", "未合并前（四系统独立）", "合并后（当前）", "备注"])
rows = [
    "一、界面展示层",
    [1, "密度推荐卡片",
     "4 张系统卡（401/402/A/B），各自目标灰分、实际灰分、偏差、建议密度",
     "1 张「合并系统」卡（501+502皮带），单一目标/实际/偏差/建议密度",
     "元素 id 全部改为 *-total"],
    [2, "灰分趋势图",
     "4 条系统线 + 1 条目标线",
     "1 条「合并系统」线 + 1 条目标线",
     "历史日志不再按系统分线"],
    [3, "灰分偏差告警",
     "对 4 个系统各生成 1 条告警，各自高亮卡片",
     "只生成 1 条「合并系统灰分偏差」，高亮合并卡",
     "告警条数 4→1"],
    [4, "建议密度",
     "每系统独立推荐（401=1.450、402=1.455、A=1.440、B=1.460）",
     "单一推荐值：取最新一条灰分密度日志的密度",
     "不再区分系统给建议"],
    [5, "置信度角标",
     "静态标定：401高/402高/A中/B低（写死）",
     "统一实时计算：高/中/低，含粗精煤泥生产模型R²、偏差、数据完整度",
     "角标与详情弹窗共用同一口径"],
    [6, "在线仪表表",
     "4 台密度计（401/402/A/B 各一台）",
     "1 台「密度计」",
     "数据采集页"],
    "二、数据录入/导入层",
    [7, "手工补录",
     "表单必须选择「所属系统」(401/402/A/B)，记录按所选系统保存",
     "无系统选项，记录 system 固定为「合并」",
     "补录口径统一"],
    [8, "表1 粗精煤泥影响因素导入",
     "记录 system = 开启系统组合字符串（如 '401+402'、'A+B'）",
     "记录 system 显示为「合并」；内部保留 sysA/sysB/sys401/sys402 四个 0/1 开关位",
     "用户纠正：数据层不合并、开关保留"],
    [9, "表2 浮精导入",
     "按系统列区分浮精数据",
     "表2 无系统列 → 一律按「合并」",
     "数据源本身无系统信息"],
    [10, "表3 灰分、密度导入",
     "按系统列（A/B）区分",
     "系统列保留作内部参考；展示与推荐不按系统拆分",
     "信息不丢弃，供后续深化"],
    [11, "数据迁移",
     "无迁移机制（老数据直接沿用）",
     "一次性迁移：__merged!==2 时用种子数据整体替换旧 localStorage",
     "四系统旧数据不兼容新界面"],
    "三、算法层（特征/模型/公式）",
    [12, "模型特征 MLR_FEATURES",
     "10 特征：原煤灰分、小时带煤量、sysA/sysB/sys401/sys402、473/474脱粉、停机、精磁尾液位",
     "不变（保持 10 特征）；中间曾改为 7 特征(running_systems)后按用户要求回退",
     "四系统开关仍是算法特征"],
    [13, "出厂模型系数",
     "10 特征 PLS（R²=0.5226/Q²=0.4218）+ MLR（R²=0.5342/Q²=0.4021），113 条样本训练",
     "不变（从备份恢复原模型系数）",
     "7 特征模型系数已废弃存档"],
    [14, "粗精煤泥灰分预测",
     "每条记录按其系统开关组合预测",
     "同左（开关保留后预测口径完全不变）",
     "无影响"],
    [15, "总灰分加权公式",
     "calcTotalAsh = (重介灰分×重介量 + 浮精灰分×浮精量 + 粗精灰分×粗精量) / 总量；重介灰分 8.50、重介量 250 为固定值",
     "公式不变；量数据改为「录入+输入」面板（总/浮/粗/重介四量，任务二）",
     "公式口径与合并无关"],
    [16, "密度与灰分日志使用",
     "分系统各自取数计算",
     "getLatestDensity 取最新一条日志，不分系统",
     "见 Sheet2 风险2"],
]
add_rows(ws, rows, 2)
style_sheet(ws, [6, 16, 46, 46, 26])
ws.auto_filter.ref = f"A1:E{ws.max_row}"

# ================= Sheet2 对算法的影响 =================
ws2 = wb.create_sheet("对算法的影响")
ws2.append(["序号", "影响点", "说明", "影响程度", "备注/建议"])
rows2 = [
    [1, "训练样本与模型系数", "4 系统开关特征保留、113 条训练样本不变 → 模型系数、R²/Q² 全部不变", "无影响",
     "训练口径未因合并改变"],
    [2, "粗精煤泥灰分预测", "每条记录仍按自己的 sysA..sys402 开关组合前馈预测", "无影响",
     "合并只影响展示，不影响预测值"],
    [3, "系统工况信息量", "保留 4 个 0/1 开关比“运行系统数”单特征信息更全（能区分“开哪几套”而不只是“开几套”）", "正向",
     "7 特征 running_systems 方案已废弃"],
    [4, "单系统差异的可见性", "4 个独立的密度推荐/偏差/告警 → 1 个合并值；个别系统异常会被合并口径掩盖", "有影响（展示层）",
     "数据层已预留 4 开关，后续可做分系统统计报表"],
    [5, "密度建议的稳定性", "推荐值取“最新一条”灰分密度日志，不分系统；若 A/B 皮带密度设定不同，建议值会随最新记录在系统间跳动", "有影响（待任务三明确口径）",
     "任务三做总灰分公式与重介灰分反推时可一并明确按皮带/系统取值"],
    [6, "浮精数据口径", "表2 无系统列 → 浮精灰分/煤量只能按合并口径统计", "数据源限制",
     "非算法选择，是三表数据本身的缺口"],
    [7, "置信度口径", "静态标定 → 动态计算（模型R²>0.8且偏差≤0.15%且数据齐全=高；R²>0.5且有粗精数据=中；否则低）", "口径变化",
     "角标与详情弹窗已统一（本轮修复）"],
    [8, "多系统同时开启时的模型表现", "一条记录 4 个开关可能同时为 1，模型按组合特征回归；若开关高度相关（常同开同停）存在共线性", "需关注",
     "当前生产模型为 PLS（自带降维），对共线较稳健；数据量增大后可重训练验证"],
]
add_rows(ws2, rows2, 2)
style_sheet(ws2, [6, 18, 56, 16, 34])
ws2.auto_filter.ref = f"A1:E{ws2.max_row}"

# ================= Sheet3 结论与风险提示 =================
ws3 = wb.create_sheet("结论与风险提示")
ws3.append(["序号", "类别", "内容"])
rows3 = [
    [1, "总体结论",
     "任务一 = 界面合并，算法不动：只改了展示/交互层（卡片、趋势图、告警、补录、导入显示），算法层（10 特征、模型系数、预测口径、总灰分公式）全部保持不变。"],
    [2, "关键决策（用户纠正）",
     "数据层不合并：内部保留 sysA/sysB/sys401/sys402 四开关；算法特征保留系统开启状态；表3 系统列保留作内部参考；恢复 10 特征原模型。"],
    [3, "风险1：单系统差异被掩盖",
     "合并卡只显示一个密度建议/一个偏差，若某套系统偏离明显（如 402 密度 1.455 vs 401 的 1.450），合并后看不到。建议后续在报表/统计层按 4 开关拆分展示（数据已预留，仅缺展示）。"],
    [4, "风险2：密度建议取“最新一条”不分系统",
     "getLatestDensity 只取最新日志，若 A/B 皮带交替采样，建议密度会随最新记录跳动。建议任务三明确：密度建议按皮带 501/502 或按系统分别取，再合成或分别推荐。"],
    [5, "风险3：浮精数据无系统维度",
     "表2 只有 采样时间/浮精灰分/浮精煤量/压滤机运行，没有系统列，合并口径是唯一选择；如需分系统需补采数据。"],
    [6, "历史存档",
     "compute_merged_model.py 计算的 7 特征合并模型（R²=0.4692）已随纠正废弃，仅作历史参考；当前生产模型为 10 特征 PLS。"],
]
add_rows(ws3, rows3, 2)
style_sheet(ws3, [6, 22, 100])
ws3.auto_filter.ref = f"A1:C{ws3.max_row}"

# ================= Sheet4 问答原文说明 =================
ws4 = wb.create_sheet("问答原文说明")
ws4.append(["章节", "说明内容"])
rows4 = [
    ("一、未合并前的逻辑（四系统独立）", True),
    ("界面",
     "总览 4 张卡片（401/402/A/B），各带目标灰分、实际灰分、偏差、建议密度（401=1.450、402=1.455、A=1.440、B=1.460）、静态置信度；趋势图 4 条线；告警每系统各发一条。"),
    ("数据",
     "每条记录带 system（如 '401+402'）；手工补录必须选“所属系统”；表1/表2/表3 导入都按系统列区分。"),
    ("算法",
     "粗精煤泥灰分模型 10 特征，其中 4 个就是系统开关（sysA/sysB/sys401/sys402，0/1）；出厂模型是 113 条样本训练的 PLS/MLR。"),
    ("二、合并后的逻辑（当前）", True),
    ("界面",
     "1 张“合并系统”卡（501+502 皮带）；1 条趋势线；1 条合并告警；补录不选系统、固定写“合并”；仪表表 1 台密度计；置信度改为实时统一计算。"),
    ("数据",
     "显示层全写“合并”，但内部保留 4 个系统开关（表1 每条记录仍存 sysA..sys402 四位；表3 保留系统列作内部参考；表2 无系统列只能按合并）。"),
    ("算法",
     "特征仍是 10 个、模型系数原样恢复、总灰分公式不变——中间一度改成 7 特征 running_systems 已回退废弃（按用户纠正执行）。"),
    ("三、对算法的影响（重点）", True),
    ("模型系数、训练样本、预测值",
     "无影响——4 开关特征保留，每条记录仍按自己的开关组合预测。"),
    ("信息量",
     "更好——4 个 0/1 开关比“运行系统数”单特征信息全，能区分“开哪几套”而不只是“开几套”。"),
    ("单系统差异可见性",
     "有影响——4 个推荐/偏差合成 1 个，个别系统偏离会被合并值掩盖（数据层已预留开关，只缺分系统展示）。"),
    ("密度建议",
     "有影响——getLatestDensity 取“最新一条”不分系统，A/B 交替采样时建议值会跳动（建议任务三明确口径）。"),
    ("浮精数据",
     "数据源限制——表2 没有系统列，只能合并统计。"),
    ("共线性隐患",
     "若 4 开关常同开同停，10 特征有共线风险；当前生产模型是 PLS（自带降维），较稳健，数据量大后可重训验证。"),
    ("一句话总结", True),
    ("总结",
     "任务一 = 界面合并、算法不动；数据层按用户要求保留 4 系统开关，所以模型和预测没有任何损失，代价是“单系统差异”目前看不见，后续可加分系统统计报表。"),
]
r = 2
for label, content in rows4:
    c1 = ws4.cell(row=r, column=1, value=label)
    c2 = ws4.cell(row=r, column=2, value=content)
    c1.font = SEC_FONT if content is True else BODY_FONT
    c2.font = SEC_FONT if content is True else BODY_FONT
    c1.alignment = WRAP
    c2.alignment = WRAP
    c1.border = BORDER
    c2.border = BORDER
    if content is True:
        c1.fill = SEC_FILL
        c2.fill = SEC_FILL
    r += 1
style_sheet(ws4, [26, 110])
ws4.auto_filter.ref = f"A1:B{ws4.max_row}"

wb.save(OUT)
print("[gen]", OUT)

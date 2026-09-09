# -*- coding: utf-8 -*-
"""生成 Excel 文档：任务一-四系统合并-修改清单（web(2) 项目）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\任务一-四系统合并-修改清单.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths, n_cols):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


def add_rows(ws, rows, start):
    r = start
    for row in rows:
        if isinstance(row, str):          # 分区标题
            ws.cell(row=r, column=1, value=row)
            for c in range(1, ws.max_column + 1):
                cc = ws.cell(row=r, column=c)
                cc.fill = SEC_FILL
                cc.font = SEC_FONT
                cc.border = BORDER
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ws.max_column)
        else:
            for c, v in enumerate(row, 1):
                cc = ws.cell(row=r, column=c, value=v)
                cc.font = BODY_FONT
                cc.alignment = WRAP if c != 1 else CENTER
                cc.border = BORDER
        r += 1
    return r


wb = Workbook()

# ================= Sheet1 修改清单 =================
ws = wb.active
ws.title = "修改清单"
ws.append(["序号", "文件", "修改点", "修改前", "修改后", "目的/说明"])

rows = [
    "一、界面层：总览页四卡合一（index.html / overview.js）",
    [1, "index.html", "密度推荐卡片",
     "4 张系统卡片（401/402/A/B，各带目标灰分、实际灰分、偏差、建议密度）",
     "1 张「合并系统」卡片 card-total（501+502皮带），元素 id 全部改为 *-total（conf/density/target-ash/actual-ash/deviation/effect/source/update-time/direction）",
     "界面只保留一套系统展示；皮带标注 501+502"],
    [2, "overview.js", "卡片数据源 systems 数组",
     "[{401,501},{402,501},{A,502},{B,502}] 四组（各自密度、置信度）",
     "仅 {id:'total', belt:'501+502'}，密度取最新灰分密度记录（getLatestDensity）",
     "后续密度计算统一走合并口径"],
    [3, "overview.js", "灰分趋势图数据集",
     "4 条系统线 + 1 条目标线（datasets 0~4）",
     "1 条「合并系统」线 + 1 条目标线（datasets 0~1）；目标线数组索引 4→1",
     "图表同步合并"],
    [4, "overview.js", "趋势数据解析",
     "按 l.system 过滤，仅统计 401/402/A/B 四条线",
     "所有灰分密度日志一律归入 'total'",
     "历史数据不分系统，合并展示"],
    [5, "overview.js", "卡片详情弹窗",
     "systems 映射 4 系统（各带皮带、密度）；标题「${sysId}系统（皮带）密度推荐计算」",
     "映射仅 total；标题固定「合并系统（501+502皮带）密度推荐计算」",
     "计算详情口径与卡片一致"],
    "二、数据采集与手工补录页（collect.js）",
    [6, "collect.js", "在线仪表表",
     "4 台密度计（401/402/A/B 各一台，值不同）",
     "1 台「密度计」（1.450）",
     "仪表列表合并展示"],
    [7, "collect.js", "手工补录表单字段",
     "灰分仪/密度计/精磁尾/浮精灰分四类表单均带「所属系统」下拉(401/402/A/B)",
     "四类表单全部去掉「所属系统」下拉",
     "补录不再要求选系统"],
    [8, "collect.js", "仪表→表单映射",
     "density_401/402/A/B → density_meter",
     "density → density_meter",
     "配合仪表合并"],
    [9, "collect.js", "提交写入 store 的系统标签",
     "const system = values.system || '--'（取表单所选系统）",
     "const system = '合并'（固定）",
     "手工补录数据统一记入合并系统"],
    "三、导入模块（import.js）",
    [10, "import.js", "表1（粗精煤泥影响因素）系统解析",
     "deriveSystem 输出 '401+402' 这类组合字符串，作为记录 system",
     "systemInfo 仅返回开启系统列表；记录 system 统一写 '合并'",
     "界面层合并"],
    [11, "import.js", "表1 数据层系统开关",
     "每条记录保存 sysA/sysB/sys401/sys402 四个开关位",
     "初版改为 running_systems（运行系统数）；经用户纠正后**恢复保存 4 个开关位**（sysA/sysB/sys401/sys402），模型特征仍按 10 特征对齐",
     "用户纠正：界面合并但数据层内部预留 4 系统开关"],
    [12, "import.js", "表3（灰分、密度导入版）系统列",
     "初版：system 固定 '合并'，忽略 Excel 系统列",
     "纠正后：表3 保留系统列作内部参考（system=该行系统值）；表2（浮精，无系统列）按 '合并'",
     "表3 的系统(A/B)列信息不丢弃"],
    [13, "import.js", "导入预览/确认行输出",
     "输出 sysA..sys402 四个开关位",
     "初版改为 onSys.includes 换算输出；最终仍按 4 个开关位输出",
     "与 10 特征模型对齐"],
    "四、算法与告警（app.js）",
    [14, "app.js", "粗精煤泥灰分模型特征 MLR_FEATURES",
     "10 特征：raw_ash/coal_amount/sysA/sysB/sys401/sys402/desliming473/desliming474/is_stoppage/level",
     "初版合并为 7 特征（running_systems 替代四开关）；经用户纠正后**恢复 10 特征原版**",
     "用户纠正：算法特征保留系统开启状态作参考"],
    [15, "app.js", "出厂模型系数 DEFAULT_COARSE_MODEL(_MLR)",
     "10 特征原模型（PLS：R²=0.5226/Q²=0.4218；MLR：R²=0.5342/Q²=0.4021）",
     "初版替换为 7 特征合并模型（R²=0.4692/Q²=0.3453）；经纠正后**从任务一前备份恢复 10 特征原模型**",
     "保证预测口径与导入数据字段一致"],
    [16, "app.js", "灰分偏差告警",
     "对 4 个系统各发一条告警，并高亮 4 张卡片",
     "只发 1 条「合并系统灰分偏差」，高亮 card-total",
     "告警随界面合并"],
    [17, "app.js", "数据迁移机制",
     "无迁移逻辑",
     "init() 中：store.__merged !== 2 时执行 applySeedStore()，用 window.DMCS_SEED_STORE（seed_data.js）整体替换旧 localStorage 数据",
     "老用户浏览器里四系统旧数据一次性清理、载入合并种子"],
    "五、种子数据（seed_data.js，由 build_web2_seed.py 生成）",
    [18, "seed_data.js", "合并种子数据",
     "（新文件）",
     "粗精煤泥 113 条（含 4 开关、入洗工作面）、浮精 5 条、灰分密度 124 条；PL/MLR 10 特征模型；importLogs 3 条；__merged=2",
     "数据源：C:\\Users\\25925\\Desktop\\web(2)导入数据\\新 三张表"],
    [19, "index.html", "脚本引用与版本号",
     "无 seed_data.js",
     "新增 js/seed_data.js；app/overview/collect/import/seed 版本号依次提升（最终 app v8、overview v6、collect v4、import v6、seed v2，后续任务在此基础上继续提升）",
     "避免浏览器缓存旧代码"],
    "六、工程侧配套脚本（backend/scripts）",
    [20, "compute_merged_model.py", "7 特征合并模型计算",
     "（新脚本）",
     "按 running_systems 单特征口径训练 MLR（R²=0.4692）；已随用户纠正**废弃**，仅作历史参考",
     "计算初版合并模型系数"],
    [21, "build_web2_seed.py", "种子数据生成器",
     "（新脚本）",
     "读取三张表生成合并种子：113 粗精 + 5 浮精 + 124 灰分密度 + 10 特征模型，写入 seed_data.js",
     "可重复执行重新生成种子"],
    [22, "（备份目录）", "改造前备份",
     "—",
     "web (2)-backup-任务一前 保存四系统原版全量文件",
     "用于纠偏时恢复 10 特征模型/字段"],
]
add_rows(ws, rows, 2)
style_sheet(ws, [6, 16, 22, 34, 44, 32], 6)
ws.auto_filter.ref = f"A1:F{ws.max_row}"

# ================= Sheet2 关键决策与纠偏 =================
ws2 = wb.create_sheet("关键决策与纠偏")
ws2.append(["序号", "时间点", "事项", "决策/结论", "落点"])
rows2 = [
    [1, "任务一实施中", "界面合并 vs 数据层合并",
     "界面合并为 1 个系统；数据层**不合并**，内部预留 sysA/sysB/sys401/sys402 四个系统开关",
     "import.js 记录字段、app.js 模型特征"],
    [2, "用户纠正", "算法特征",
     "系统开启状态必须保留作参考 → 恢复 10 特征（4 开关），废弃 running_systems 7 特征方案",
     "app.js MLR_FEATURES/FEATURE_LABELS"],
    [3, "用户纠正", "出厂模型",
     "恢复任务一前 10 特征原模型系数（PLS/MLR，Q² 选优）",
     "app.js DEFAULT_COARSE_MODEL(_MLR)"],
    [4, "用户纠正", "表3 系统列",
     "表3（灰分、密度）保留系统列作内部参考，不丢弃；表2（浮精）无系统列按合并",
     "import.js 模板导入逻辑"],
    [5, "合并口径", "密度/灰分展示",
     "总览卡片、趋势图、告警、手工补录、导入记录全部按「合并系统」展示；皮带标注 501+502",
     "index.html/overview.js/collect.js"],
    [6, "数据迁移", "老数据",
     "一次性迁移：__merged!==2 时用种子数据整体替换旧 localStorage（四系统旧数据不兼容新界面）",
     "app.js applySeedStore / seed_data.js"],
]
for r in rows2:
    ws2.append(r)
style_sheet(ws2, [6, 14, 26, 60, 30], 5)
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER
ws2.auto_filter.ref = f"A1:E{ws2.max_row}"

wb.save(OUT)
print("[gen]", OUT)

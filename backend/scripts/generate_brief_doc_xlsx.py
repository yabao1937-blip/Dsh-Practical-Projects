# -*- coding: utf-8 -*-
"""生成 Excel 文档：推测简报-功能说明（三表1h对齐 + 建议密度 + 模型推测）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\推测简报-功能说明.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


wb = Workbook()

# ---------------- Sheet 1: 功能规格 ----------------
ws = wb.active
ws.title = "功能规格"
ws.append(["板块", "项目", "内容", "备注"])
rows = [
    ["触发", "入口", "批量数据导入页 页头新增「生成推测简报」按钮（ImportPage.generateBrief）", "导入三表后点击"],
    ["触发", "前置", "store 中已导入 表1(粗精煤泥多因素 coarseCoal) / 表2(浮精 floatCoal) / 表3(灰分密度 calcLogs)", "无数据时提示先导入三张表"],
    ["对齐", "时间轴", "三表最早~最晚，每个整点一行（连续1h）；两次采样之间的整点由递归预测/前向填充补位", "连续输出"],
    ["对齐", "行过滤", "无法给出「建议密度」和「粗精煤泥灰分(预测)」的行不生成（无粗灰数据/首次采样前的时段被剔除）", "简报从首次采样起"],
    ["对齐", "桶内取值", "每表独立前向填充：取 ts < 桶结束 的最后一条记录（含前向继承）", "表1≈每小时、表2/表3稀疏→前向填充"],
    ["计算", "核心函数", "App.buildHourlyBrief() 纯回放计算，返回 {headers, rows}；不触碰 resolve* 的取最新语义", "app.js v56"],
    ["输出", "展示", "弹窗表格（modal）+ 摘要信息（行数/口径/目标/容差/恒值说明）", "table-scroll 60vh"],
    ["输出", "导出", "「导出Excel」按钮 → 推测简报_1h对齐.xlsx（XLSX.writeFile）", "import.js v16"],
    ["口径", "建议密度", "该小时状态下 computeDensityGuidance 同源专家表 → rhoNew；版本跟随当前 总灰分版/重介版 开关", "|ΔA|≤容差→保持当前密度"],
    ["口径", "表头布局", "左侧=静态数据(皮带秤/粗精煤泥量/重介灰分/总精煤量)；右侧=建议密度|实测密度 + 预测粗灰|实测粗灰；实测列稀疏(有测量才填)", "便于对照预测vs实测"],
    ["口径", "粗精煤泥灰分实测", "315灰分 = 粗精煤泥灰分（人工采样，稀疏），取表1 315灰分(ash_content)，仅采样时刻有值", "第16列"],
    ["口径", "模型推测(递归软测量)", "基值=第一次采样值；后续每小时=上一小时预测值+(模型原始预测增量)；有采样时用采样值反馈重置；该值作为总灰分/建议密度的当前粗灰数据", "首行无预测(无基值)"],
    ["口径", "建议密度", "仅公式完整(有粗灰数据)时计算；缺粗灰数据(如表1未覆盖的时段)时留空，避免退化为仪表加权后顶到1.60上限", "第13列"],
    ["缺口", "粗精煤泥量", "三表无数据源，恒值 40 t/h（当前默认）", "标橙色提示"],
    ["缺口", "501/502皮带秤", "三表无数据源，恒值 268.5 / 235.2 t/h", "标橙色提示"],
    ["缺口", "重介灰分/总精煤量", "重介灰分恒值 8.50%（静态初始值）；总精煤量=皮带秤和 503.7 恒值", "采样/录入才改"],
]
for r in rows:
    ws.append(r)
style_sheet(ws, [8, 12, 62, 22])
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER
        if cell.column == 1:
            cell.fill = SEC_FILL
            cell.font = SEC_FONT
ws.auto_filter.ref = f"A1:D{ws.max_row}"

# ---------------- Sheet 2: 列定义 ----------------
ws2 = wb.create_sheet("列定义")
ws2.append(["序号", "列名", "取值来源", "前向填充/默认值", "说明"])
cols = [
    [1, "时间", "1h 桶起点", "—", "格式 YYYY-MM-DD HH:00"],
    [2, "501皮带秤(t/h)", "无(三表缺)", "恒值 268.5", "静态数据(标橙色)"],
    [3, "502皮带秤(t/h)", "无(三表缺)", "恒值 235.2", "静态数据(标橙色)"],
    [4, "粗精煤泥量(t/h)", "无(三表缺)", "恒值 40", "静态数据(标橙色)"],
    [5, "重介精煤灰分(%)", "静态初始值", "恒值 8.50", "静态数据"],
    [6, "总精煤量(t/h)", "501+502皮带秤和", "恒值 503.7", "静态数据"],
    [7, "501皮带灰分仪(%)", "表3 ash_content (belt=501)", "默认 8.52", "前向填充"],
    [8, "502皮带灰分仪(%)", "表3 ash_content (belt=502)", "默认 10.68", "前向填充"],
    [9, "精磁尾液位(%)", "表1 level", "默认 55", "前向填充"],
    [10, "浮精灰分(%)", "表2 ash_content", "默认 9.85", "前向填充"],
    [11, "浮精量(t/h)", "表2 coal_amount", "无→留空", "前向填充"],
    [12, "总精煤灰分(%)", "公式加权(重介8.50×重介量+浮灰×浮量+粗灰(递归软测量)×40)/503.7", "公式缺→表3仪表加权", "粗灰用递归软测量值"],
    [13, "建议密度(g/cm³)", "专家表修正 rhoNew=ρ±expertAdjust(|ΔA|)", "夹紧[1.35,1.60]", "预测值(每1h)"],
    [14, "实测密度(g/cm³)", "表3 density", "仅测量时刻有值", "稀疏,与建议密度对照"],
    [15, "预测粗精煤泥灰分(%)", "递归软测量:基值=首次采样,后续=上小时预测值+模型增量,采样反馈重置", "首行无预测", "作为总灰分当前粗灰"],
    [16, "实测粗精煤泥灰分(%)", "表1 ash_content(315灰分)", "仅采样时刻有值", "315灰分=粗精煤泥灰分(稀疏)"],
]
for r in cols:
    ws2.append(r)
style_sheet(ws2, [6, 22, 34, 16, 34])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column not in (1,) else CENTER
        if cell.column == 4 and "恒值" in str(cell.value or ""):
            cell.fill = WARN_FILL
ws2.auto_filter.ref = f"A1:E{ws2.max_row}"

wb.save(OUT)
print("[gen]", OUT)

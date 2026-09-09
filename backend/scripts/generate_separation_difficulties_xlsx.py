# -*- coding: utf-8 -*-
"""生成 Excel：前后端分离-难点原因与建议（我的 + GLM）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\前后端分离-难点与建议.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def make_sheet(wb, title, headers, rows, widths, fill_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER
    if fill_col:
        for row in ws.iter_rows(min_row=2):
            c = row[fill_col - 1]
            v = str(c.value or "")
            if "已修复" in v or "已澄清" in v or "已补充" in v:
                c.fill = OK_FILL
            elif "待处理" in v or "未后移" in v:
                c.fill = WARN_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws


wb = Workbook()
wb.remove(wb.active)

# ---------------- Sheet 1 难点与原因 ----------------
make_sheet(wb, "1难点与原因",
    ["序号", "难点", "原因", "现状"],
    [
        [1, "算法数值一致性（逐值对齐 <1e-6）",
         "JS 手写回归/PLS（高斯消元、岭回归 λ 网格+LOOCV、NIPALS+LOOCV）与 numpy/sklearn 在浮点、λ 选择、PLS 分量选择上存在实现差异",
         "已解决：用 golden 对拍（简报 1069 行全量 + 取值链）验证默认模型预测路径逐值一致"],
        [2, "MLR/PLS 训练后移（sklearn）",
         "JS 手写 NIPALS PLS + 岭回归，与 sklearn PLSRegression/Ridge 内部算法细节不同，逐位对齐极难；训练输出不参与「逐值一致」验收（验收针对预测）",
         "未后移：前端仍客户端训练（功能可用），后端训练留作后续优化"],
        [3, "取值链复杂（手动>公式>录入>默认）",
         "多级优先级 + 手动有效期 autoState（manual_at ≥ auto_at 时间戳判定）+ 13 行录入层，易漏边界",
         "已移植 + golden 对拍；GLM 发现负数 manual 泄漏 bug 已修复"],
        [4, "递归软测量（粗灰）",
         "基值=首次采样、有采样反馈重置、无采样=上一预测值+模型增量；三表时间不重叠导致「预测==实测」",
         "已移植 + 简报 golden 对拍验证"],
        [5, "推测简报 1h 对齐",
         "连续 1h 逐整点、三表前向填充、16 列表头、过滤「无法给出建议密度+预测粗灰」的行",
         "已移植 + 1069 行全量 golden 对拍"],
        [6, "M.DD 日期补偿",
         "手写表「个位天数省略尾零」语义（6.3=6月30日、7.4=7月4日、6.1=6月1日）；浮点 6.3-6=0.2999 需 round 到 30",
         "已移植 + 单测通过"],
        [7, "手动有效期时间戳",
         "autoState 依赖 Date.now()，前后端时钟可能不一致导致误接管",
         "规划用服务器统一时间戳打点；种子数据无手动覆盖暂未实测"],
        [8, "三表时间不重叠",
         "种子数据表1=6-7月、表2/表3=5月，实测密度/粗灰与建议密度/预测粗灰不在同一时段",
         "演示数据问题，真实数据会重叠；简报过滤规则已处理"],
        [9, "文件编码（GBK/UTF-8）",
         "PowerShell 5.1 默认 GBK 读取 UTF-8 文件，批量替换时写坏 33 个 verify 脚本（✓/✗ 与中文变乱码）",
         "已重生成 2 个关键脚本；其余冗余（pytest + golden 已覆盖）"],
        [10, "CORS / file://",
         "前端原 file:// 打开，fetch 受 CORS 限制",
         "后端 StaticFiles 托管（同源）解决"],
    ],
    [6, 22, 60, 34], fill_col=4)

# ---------------- Sheet 2 我的建议 ----------------
make_sheet(wb, "2我的建议",
    ["序号", "建议", "针对难点", "说明"],
    [
        [1, "golden 对拍（逐值 diff <1e-6）作为硬验收", "1", "同一批数据，前端 JS 与后端逐值对比，防止算法迁移漂移"],
        [2, "双轨过渡（localStorage 兜底 + API 镜像）", "3,10", "file:// 仍可用，http:// 自动同步，避免破坏现有用法"],
        [3, "字符串时间戳方案（YYYY-MM-DD HH:MM:SS）", "1,5", "字符串序==时间序，规避时区/夏令时问题，与 JS 等价"],
        [4, "整库镜像（/state 快照）优先于逐页精细 API", "3", "先保证功能等价，再按需拆精细接口"],
        [5, "PLS 先对齐 JS NIPALS 行为再谈 sklearn 优化", "2", "训练后移的数值对齐优先级低于预测"],
        [6, "测试用独立 SQLite / clear+seed fixture 隔离", "1", "避免测试互相污染（dashboard round-trip 需清库）"],
        [7, "服务器统一时间戳打点", "7", "autoState 的 auto_at/manual_at 由后端统一，避免浏览器时钟漂移"],
        [8, "静态托管规避 CORS", "10", "FastAPI StaticFiles 托管前端，同源访问"],
        [9, "改脚本用 Node 或 edit 工具（UTF-8 安全），避免 PowerShell GBK", "9", "教训：批量文本替换要用 UTF-8 安全通道"],
    ],
    [6, 34, 12, 50])

# ---------------- Sheet 3 GLM 的建议 ----------------
make_sheet(wb, "3GLM的建议",
    ["级别", "发现/建议", "处置", "说明"],
    [
        ["高", "resolve_amount/resolve_coarse_amount 兜底泄漏负数 manual 值", "已修复",
         "末尾 return manual 改为 return None，并排除 bool；负数量会污染总灰分与密度指导"],
        ["高", "load_store 不按时间排序，_latest_calc_value 取到过期值", "已修复",
         "calcLogs 按 timestamp 排序，使 reversed() 取到真正的最新值"],
        ["中", "PUT /state 整库重写非原子（delete+commit 后 apply 失败即清库）", "已修复",
         "migrate.replace 清表+写入抽为 _apply_in_session，同一事务单次 commit，失败回滚"],
        ["中", "brief 与 resolvers 对 heavyAshInput.manual 校验不一致（NaN/负数）", "已修复",
         "brief 复用 resolvers.get_heavy_ash；该校验拒绝 bool/NaN/inf/负数，回退 8.50"],
        ["中", "maxStep 读取但从未生效（死配置）", "已澄清",
         "前端用 maxStep 做「逐步走向目标密度」动画（页面可调钳制）；后端纯函数仅回显以对齐前端输出结构，已在 density.py 加注释说明"],
        ["中", "parse_ts/normalize_ts 硬编码年份 2026", "已修复",
         "集中为 DEFAULT_YEAR 常量并加注释（与前端 import.js 一致，跨年导入需调整）"],
        ["低", "布尔值穿透数值校验（isinstance(True,int) 为 True）", "已修复",
         "resolvers.py 引入 _is_num 助手（排除 bool），全部手动/录入/记录数值校验点统一改用"],
        ["低", "resolve_total_ash 最终兜底可能除零（两台皮带秤 manual 均为 0）", "已修复",
         "分母 denom==0 时返回 None（不除零），并加 test_total_ash_div_zero 回归测试"],
        ["低", "GET /records 的 total 是 limit 截断后条数，命名误导", "已修复",
         "total = q.count()（过滤后真实总数），与 items 分离"],
        ["低", "测试覆盖缺口", "已补充",
         "新增 test_total_ash_div_zero、test_heavyash_invalid_manual_rejected 回归测试（此前已补 test_negative_manual_not_leaked）；pytest 35 全绿"],
    ],
    [8, 52, 10, 50], fill_col=3)

# ---------------- Sheet 4 汇总 ----------------
make_sheet(wb, "4汇总",
    ["维度", "结论"],
    [
        ["核心验收", "逐值一致(<1e-6)✅、localStorage 迁移✅、6 页面功能等价✅、headless 回归✅（pytest 35 + 2 关键脚本）"],
        ["难点 10 项", "已解决 6 项（1,3,4,5,6,10）、留作优化 1 项（2 训练后移）、规划待验证 1 项（7 时间戳）、演示数据 1 项（8）、操作失误已补救 1 项（9 编码）"],
        ["GLM 建议 10 条", "已修复 2 高 + 6 中低（负数 manual、calcLogs 乱序、PUT/state 原子性、heavyAsh 口径统一、年份常量、bool 排除、除零保护、records.total 真实计数）；maxStep 澄清为前端动画回显（保留对齐）；测试覆盖已补充"],
        ["下一步", "① MLR/PLS 训练后移（sklearn，需先对齐 JS NIPALS）；② 服务器统一时间戳打点（难点7，接入真实数据前验证）；③ 其余 verify 脚本按需重生成（pytest+golden 已覆盖，可选）"],
    ],
    [16, 90])

wb.save(OUT)
print("[gen]", OUT)

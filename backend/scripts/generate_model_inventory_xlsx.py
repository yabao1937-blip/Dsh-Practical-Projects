# -*- coding: utf-8 -*-
"""生成 Excel 文档：系统模型全景清单"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\系统模型全景清单.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "模型全景清单"
ws.append(["序号", "模型", "算法", "输入特征", "输出", "训练数据", "训练/更新方式", "当前状态", "用途与说明"])

rows = [
    "一、正在使用的模型（web(2) 系统）",
    [1, "粗精煤泥灰分模型（生产模型）", "PLS 偏最小二乘（NIPALS）",
     "10 特征：原煤灰分、小时带煤量、A/B/401/402四开关、473/474脱粉、停机、精磁尾液位",
     "粗精煤泥灰分(315灰分%)",
     "表1 粗精煤泥灰分影响因素 6.16-7.14（113条）",
     "出厂默认 PLS；可在粗精煤泥页点“重训练”用当前全部记录重训",
     "使用中（R²=0.5226/Q²=0.4218）",
     "与 MLR 对比按留一交叉验证 Q² 选优后作为生产模型；每条新记录实时前馈预测"],
    [2, "粗精煤泥灰分模型（对照模型）", "MLR 多元线性回归（标准化正规方程）",
     "同 10 特征",
     "粗精煤泥灰分",
     "同一 113 条样本",
     "与 PLS 同批训练；粗精页可切换“对比查看”",
     "使用中（R²=0.5342/Q²=0.4021）",
     "Q² 低于 PLS，作对照；用于公式校验/因子权重对比"],
    [3, "单因素精磁尾液位模型（3个备选）", "线性回归 linear / 二次多项式 poly2 / 三次多项式 poly3",
     "精磁尾液位(%)",
     "315灰分(%)",
     "store.magneticTail（手工补录的精磁尾液位+灰分记录）",
     "trainModels()：液位数据≥5条训练linear、≥6条训练poly2、≥8条训练poly3",
     "存在但当前未训练（种子数据无精磁尾记录）",
     "predictAsh(level) 自动选 R² 最高的模型预测灰分；属于合并前旧口径的单因素模型"],
    "二、预留/待建的模型",
    [4, "密度推荐模型（预留）", "待定（任务三）",
     "皮带灰分(belt_ash)、密度(density)（表3 灰分密度）",
     "建议密度",
     "表3 灰分、密度（导入版）124条",
     "未实现：DATA_INTERFACES 已预留 density/belt_ash 字段接口",
     "预留",
     "任务三将做总精煤灰分公式 + 反推重介精煤灰分，据此再定密度推荐口径"],
    [5, "总精煤灰分公式（任务三）", "加权平均公式 calcTotalAsh",
     "重介灰分×重介量 + 浮精灰分×浮精量 + 粗精灰分×粗精量",
     "总精煤灰分 / 反推重介精煤灰分",
     "量数据面板（总/浮/粗/重介四量，录入+输入）",
     "公式已内置；反推逻辑待任务三实现",
     "公式在用，反推未实现",
     "重介精煤灰分目前是固定 8.50%，任务三改为反推"],
    "三、已废弃/存档的模型",
    [6, "7特征合并模型（废弃）", "MLR（running_systems 单特征替代四开关）",
     "7 特征：原煤灰分、带煤量、运行系统数、473/474脱粉、停机、液位",
     "粗精煤泥灰分",
     "同一 113 条样本（系统开关合并为数量）",
     "compute_merged_model.py 计算，R²=0.4692/Q²=0.3453",
     "已废弃（按用户纠正回退10特征）",
     "仅存档于 backend/scripts/compute_merged_model.py，作为历史参考"],
]
for r in rows:
    ws.append([r] if isinstance(r, str) else r)
for i, w in enumerate([6, 24, 24, 28, 18, 24, 30, 18, 34], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for cell in ws[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = CENTER
    cell.border = BORDER
r = 2
for row in ws.iter_rows(min_row=2):
    if isinstance(row[0].value, str) and row[0].value.startswith(("一、", "二、", "三、")):
        for c in range(1, 10):
            cc = ws.cell(row=r, column=c)
            cc.fill = SEC_FILL
            cc.font = SEC_FONT
            cc.border = BORDER
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    else:
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER
    r += 1
ws.freeze_panes = "A2"

# Sheet2 模型间关系
ws2 = wb.create_sheet("模型间关系")
ws2.append(["关系", "说明"])
rel = [
    ["粗精煤泥灰分模型 → 总灰分计算",
     "PLS 生产模型输出“粗精煤泥灰分(前馈预测)”，作为 calcTotalAsh 的粗精灰分项，进入总精煤灰分与偏差计算（总览卡片、告警、浮精页影响值都依赖它）。"],
    ["单因素液位模型 ↔ 粗精煤泥多因素模型",
     "同为预测 315 灰分：单因素只用液位一个特征（旧口径，未训练）；多因素 10 特征（当前口径）。两者互不影响，predictAsh 仅在选择旧模型时使用。"],
    ["密度模型（预留）→ 建议密度",
     "当前建议密度 = getLatestDensity() 直接取表3最新密度记录，无模型；任务三确定反推公式后再决定是否训练 灰分→密度 模型。"],
    ["模型重训练链路",
     "粗精页“重训练”按钮 → App.trainCoarseModel()：重训 MLR+PLS → 按 Q² 选生产模型 → 回填每条记录 predicted_ash → 记 coarseModelHistory 快照。"],
    ["数据接口预留（DATA_INTERFACES）",
     "20 个字段已注册来源表与用途：10 个特征 → 粗精煤泥灰分模型；belt_ash/density → 总灰分计算、密度模型(预留)；ash_density_system → 内部参考。"],
]
for r_ in rel:
    ws2.append(r_)
for i, w in enumerate([34, 100], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for cell in ws2[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = CENTER
    cell.border = BORDER
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP
ws2.freeze_panes = "A2"

# Sheet3 现状速览（一句话确认）
ws3 = wb.create_sheet("现状速览")
ws3.append(["项目", "当前取值方式", "性质", "说明"])
snapshot = [
    ["粗精煤泥灰分(315灰分)", "PLS(生产模型)+MLR(对照) 10特征前馈预测", "预测模型",
     "系统里唯一真正在用的预测模型"],
    ["重介精煤灰分", "8.50%", "定值",
     "任务三将改为反推"],
    ["重介精煤量", "录入 或 计算(总精煤量−浮精量−粗精煤泥量)", "定值/公式",
     "量数据面板"],
    ["浮精灰分", "表2 最新一条化验值", "直接取值",
     "非模型"],
    ["浮精煤量", "录入 或 自动取表2最新记录", "直接取值",
     "非模型"],
    ["粗精煤泥量", "人工录入", "定值",
     "三表无此数据源"],
    ["总精煤量", "人工录入（501+502皮带）", "定值",
     "PLC接口预留"],
    ["总精煤灰分", "(重介灰分×重介量+浮精灰分×浮精量+粗精灰分×粗精量)/总量", "加权公式",
     "非模型"],
    ["目标灰分", "8.50%（可箭头微调）", "定值", "—"],
    ["建议密度", "表3 最新一条密度值 + 方向建议(上调/下调/保持)", "直接取值",
     "数值不调整，无模型"],
    ["置信度", "规则判定：模型R²+偏差+数据完整度 → 高/中/低", "规则",
     "非模型"],
    ["单因素液位模型(linear/poly2/poly3)", "代码存在，当前未训练（无精磁尾数据）", "备选(旧口径)",
     "predictAsh 选R²最高者"],
    ["密度模型", "预留接口(DATA_INTERFACES)，未实现", "待建",
     "任务三后确定"],
    ["结论", "是：目前只有一套（粗精煤泥灰分 PLS+MLR）预测模型，其余环节均为定值、直接取数或公式", "确认",
     "7特征模型已废弃；密度模型待建"],
]
for r in snapshot:
    ws3.append(r)
for i, w in enumerate([26, 40, 16, 40], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
for cell in ws3[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = CENTER
    cell.border = BORDER
for row in ws3.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER
ws3.freeze_panes = "A2"

wb.save(OUT)
print("[gen]", OUT)

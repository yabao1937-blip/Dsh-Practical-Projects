"""Read-only pairing of recent laboratory assays with raw telemetry.

This produces a GPT research dataset, never imports records or trains a model.
Run with a Python environment containing openpyxl. All SQLite files are opened
mode=ro, with query_only enabled, including optional retention-recovery backups.
"""
import argparse
from contextlib import ExitStack
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import sqlite3

from openpyxl import load_workbook

LOCAL = timezone(timedelta(hours=8))
ITEMS = {
    'coal_amount': ('小时带煤量', 'TIA_3_1.TIA_3_1.301/COAL_CAP', 't/h'),
    'sysA': ('A系统', 'TIA_4_2.TIA_4_2.A345/if_prod', ''),
    'sysB': ('B系统', 'TIA_4_2.TIA_4_2.A346/if_prod', ''),
    'sys401': ('401系统', 'TIA_3_3.TIA_3_3.478/if_prod', ''),
    'sys402': ('402系统', 'TIA_3_3.TIA_3_3.479/if_prod', ''),
    'desliming473': ('473脱粉', 'TIA_3_3.TIA_3_3.473/if_prod', ''),
    'desliming474': ('474脱粉', 'TIA_3_3.TIA_3_3.474/if_prod', ''),
    'level': ('精磁尾液位', 'TIA_3_5.TIA_3_5.JMCWC/LEVEL_B', '%'),
}


def iso(dt):
    return dt.astimezone(timezone.utc).isoformat(timespec='milliseconds')


def local(value):
    return datetime.fromisoformat(value).astimezone(LOCAL).isoformat(timespec='milliseconds') if value else None


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    allow_nan=False).encode('utf-8')).hexdigest()


def file_hash(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def definition(config):
    fields = {k: config.get(k) for k in ('node_id', 'unit', 'scale', 'offset',
                                        'value_type', 'minimum', 'maximum')}
    return hashlib.sha256(json.dumps(fields, sort_keys=True).encode()).hexdigest()


def read_assays(path, year, date_fixes):
    result, repairs = [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['Sheet1']
        production_day = previous_time = None
        day_offset = 0
        for rowno, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            label, clock, ash = (list(cells) + [None] * 3)[:3]
            if clock is None or rowno <= 2:
                continue
            if label is not None:
                original = str(label)
                normalized = date_fixes.get(original, original)
                month, day = map(int, normalized.split('.'))
                production_day = date(year, month, day)
                previous_time, day_offset = None, 0
                if normalized != original:
                    repairs.append(dict(source_row=rowno, original=original,
                                        resolved=str(production_day), reason='相邻日期顺序确认：Excel数字省略末尾0'))
            if production_day is None or not isinstance(clock, time):
                raise ValueError(f'第{rowno}行日期或时间无法确定')
            if previous_time is not None and clock < previous_time:
                day_offset += 1  # User confirmed production-day groups cross midnight.
            at = datetime.combine(production_day + timedelta(days=day_offset), clock, LOCAL)
            previous_time = clock
            if isinstance(ash, bool) or not isinstance(ash, (int, float)) or not math.isfinite(ash) or not 0 < ash <= 100:
                raise ValueError(f'第{rowno}行灰分不在(0,100]')
            if result and at <= datetime.fromisoformat(result[-1]['sample_at']):
                raise ValueError(f'第{rowno}行时间重复或逆序')
            result.append(dict(source_sheet=sheet.title, source_row=rowno,
                               source_date=label, source_time=clock.isoformat(),
                               production_day=str(production_day), sample_at=at.isoformat(),
                               sample_at_utc=iso(at), ash_content=float(ash),
                               report_available_at=None))
    finally:
        workbook.close()
    return result, repairs


def open_database(stack, path, name):
    conn = sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)
    stack.callback(conn.close)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA query_only=ON')
    conn.execute('BEGIN')  # One consistent read snapshot per database.
    points = {}
    for row in conn.execute('SELECT source_id,key,config FROM telemetry_points'):
        cfg = json.loads(row['config'])
        points[(row['source_id'], row['key'])] = cfg
    return dict(conn=conn, points=points, name=name, path=str(path.resolve()))


def selected_points(database, at):
    row = database['conn'].execute(
        'SELECT id,config FROM decision_source_revisions WHERE effective_at<=? ORDER BY id DESC LIMIT 1',
        (iso(at),)).fetchone()
    choices = json.loads(row['config'])['choices'] if row else {}
    nodes = {'coal_amount': 'coal-rate', 'level': 'tail-level'}
    result = {}
    for feature, (_, item, _) in ITEMS.items():
        node = nodes.get(feature, 'system-switches' if feature.startswith('sys') else 'desliming')
        selected = json.loads(choices[node]) if choices.get(node) else None
        if feature in nodes and selected:
            candidates = [tuple(selected)] if tuple(selected) in database['points'] else []
        else:
            expected = 'four_system_status' if feature.startswith('sys') else 'desliming_status'
            if selected and selected[1] != expected:
                raise ValueError(f'无法还原历史组合点: {selected}')
            candidates = [key for key, cfg in database['points'].items()
                          if cfg.get('item_id') == item and (not selected or key[0] == selected[0])]
        if len(candidates) != 1:
            raise ValueError(f'{feature}历史测点不唯一: {candidates}')
        result[feature] = candidates[0]
    return result, row['id'] if row else 0


def match_feature(databases, key, feature, target):
    candidates, following = [], []
    for priority, db in enumerate(databases):
        for op, order, output in [('<=', 'DESC', candidates), ('>', 'ASC', following)]:
            row = db['conn'].execute(
                f'SELECT * FROM telemetry_readings WHERE source_id=? AND point_key=? '
                f'AND received_at{op}? ORDER BY received_at {order},id {order} LIMIT 1',
                (*key, iso(target))).fetchone()
            if row:
                output.append((dict(row), db, priority))
    nearest = min(following, key=lambda x: (x[0]['received_at'], x[2]), default=None)
    meta = dict(feature=feature, name=ITEMS[feature][0], unit=ITEMS[feature][2],
                source_id=key[0], point_key=key[1], target_at=iso(target), target_local=local(iso(target)),
                next_received_at=nearest[0]['received_at'] if nearest else None)
    if not candidates:
        return dict(meta, value=None, candidate_value=None, status='no_history',
                    reason='目标之前无采集记录', offline_usable=False, strict_usable=False)
    row, db, _ = max(candidates, key=lambda x: (x[0]['received_at'], -x[2], x[0]['id']))
    sample = json.loads(row['sample'])
    sample = {k: v for k, v in sample.items() if k not in ('_display_memory', '_density_display_memory')}
    cfg = db['points'].get(key)
    age = (target - datetime.fromisoformat(row['received_at'])).total_seconds()
    limit = 30 if feature == 'level' else 65
    value, reason, status = sample.get('value'), '', 'good'
    if age > limit:
        status, reason = 'gap', f'前一条采集距目标{age:.3f}秒，超过{limit}秒匹配范围'
    elif not cfg or sample.get('definition_id') != definition(cfg):
        status, reason = 'definition_mismatch', '测点定义不一致'
    elif sample.get('quality') != 'good' or sample.get('issue'):
        status, reason = 'bad_quality', '采集质量异常：' + str(sample.get('issue') or sample.get('quality'))
    elif feature in ('coal_amount', 'level'):
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            status, reason = 'invalid', '数值类型无效'
        elif feature == 'coal_amount' and value <= 10:
            status, reason = 'stopped', '原始入洗煤量≤10，保留原值但不视为正常生产样本'
        elif feature == 'level' and not 0 <= value <= 100:
            status, reason = 'invalid', '液位超出0–100%'
    elif not isinstance(value, bool):
        status, reason = 'invalid', '运行状态不是原始布尔值'
    if sample.get('received_at') != row['received_at']:
        raise ValueError('记录与载荷接收时间不一致')
    good = status == 'good'
    return dict(meta, database=db['name'], reading_id=row['id'], batch_id=row['batch_id'],
                received_at=row['received_at'], received_local=local(row['received_at']),
                source_at=sample.get('source_at'), server_at=sample.get('server_at'),
                definition_id=sample.get('definition_id'), raw_sample=sample,
                raw_sample_sha256=digest(sample), candidate_value=value,
                value=(int(value) if isinstance(value, bool) else value) if good or status == 'stopped' else None,
                quality=sample.get('quality'), issue=sample.get('issue'), age_seconds=round(age, 3),
                offline_usable=good, strict_usable=good and age <= 30,
                status=('minute_approximation' if good and age > 30 else status), reason=reason)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', required=True, type=Path)
    parser.add_argument('--database', required=True, type=Path)
    parser.add_argument('--backup', type=Path, action='append', default=[])
    parser.add_argument('--year', required=True, type=int)
    parser.add_argument('--date-fix', action='append', default=[])
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    rows, repairs = read_assays(args.workbook, args.year, dict(x.split('=') for x in args.date_fix))
    with ExitStack() as stack:
        databases = [open_database(stack, args.database, 'production')]
        databases.extend(open_database(stack, p, f'backup_{i+1}') for i, p in enumerate(args.backup))
        for row in rows:
            at = datetime.fromisoformat(row['sample_at'])
            points, revision = selected_points(databases[0], at)
            row['source_revision'] = revision
            row['matches'] = {feature: match_feature(databases, key, feature,
                              at - timedelta(seconds=90 if feature == 'level' else 0))
                              for feature, key in points.items()}
            other = [row['matches'][f] for f in ITEMS if f != 'level']
            same_batch = (all(m.get('batch_id') for m in other)
                          and len({(m.get('batch_id'), m.get('received_at')) for m in other}) == 1)
            row['features'] = {k: m['value'] for k, m in row['matches'].items()}
            row['offline_complete'] = same_batch and all(m['offline_usable'] for m in row['matches'].values())
            row['strict_30s_complete'] = same_batch and all(m['strict_usable'] for m in row['matches'].values())
            row['match_status'] = ('严格匹配' if row['strict_30s_complete'] else '分钟级近似'
                                   if row['offline_complete'] else '历史缺失或异常')
            row['requires_review'] = not row['strict_30s_complete']
            row['reasons'] = [m['name'] + '：' + m['reason'] for m in row['matches'].values() if m['reason']]
            if not same_batch:
                row['reasons'].append('非液位因子没有同一采集批次的完整记录')
            if row['offline_complete'] and not row['strict_30s_complete']:
                row['reasons'].append('历史分钟采集；仅作离线近似拼接，不满足在线30秒时效')
        db_provenance = []
        for db in databases:
            max_id = db['conn'].execute('SELECT max(id) FROM telemetry_readings').fetchone()[0]
            db_provenance.append(dict(name=db['name'], path=db['path'], telemetry_max_id=max_id,
                                     sha256=file_hash(Path(db['path'])) if db['name'] != 'production' else None))
    output = dict(
        schema='gpt-recent-assays-raw-telemetry-join-v1', generated_at=iso(datetime.now(timezone.utc)),
        source_workbook=str(args.workbook.resolve()), source_sha256=file_hash(args.workbook),
        databases=db_provenance, date_repairs=repairs,
        policy=dict(timezone='Asia/Shanghai', year=args.year, year_basis='用户称近期，按当前2026年及数据库日期',
                    midnight='用户确认生产日组内凌晨归次日', clock='received_at',
                    selection='latest-at-or-before-target; never skip latest bad reading',
                    level_lag_seconds=90, level_tolerance_seconds=30,
                    other_offline_tolerance_seconds=65, strict_online_tolerance_seconds=30,
                    offline_tolerance_basis='早期历史采集周期约61–63秒，仅生成显式标注的分钟级近似候选',
                    missing='保留化验值，缺失因子留空；不插值、不沿用数小时前值、不用模型预测补全',
                    receipt_note='source_at可能是值变化时间，保留审计但不用于替换received_at',
                    report_time='原表未提供；禁止假定采样时化验结果已经可用',
                    purpose='数据准备；未训练、未激活、未写入生产数据库'),
        feature_names=list(ITEMS),
        summary=dict(total=len(rows), offline_complete=sum(r['offline_complete'] for r in rows),
                     strict_complete=sum(r['strict_30s_complete'] for r in rows),
                     incomplete=sum(not r['offline_complete'] for r in rows),
                     recovered_from_backup=sum(any(m.get('database', '').startswith('backup')
                         and m['offline_usable'] for m in r['matches'].values()) for r in rows)), rows=rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    print(json.dumps(output['summary'], ensure_ascii=False))
    for r in rows:
        if r['offline_complete']:
            print(r['source_row'], r['sample_at'], r['ash_content'], r['features'])


if __name__ == '__main__':
    main()

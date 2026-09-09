# -*- coding: utf-8 -*-
"""生成 Excel 文档：315灰分测量噪声测试方案（平行样重复性试验）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\315灰分噪声测试方案.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


def body(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER


wb = Workbook()

# ============ Sheet1 试验方案 ============
ws = wb.active
ws.title = "试验方案"
ws.append(["项目", "内容"])
rows1 = [
    ["试验名称", "315灰分测量噪声试验（平行样重复性）"],
    ["目的", "量化315灰分「采样+制样+化验」全流程的随机误差σ，判断当前模型 RMSE=2.4% 中还有多少可优化空间：\n① 若σ≥1.5%，模型已贴近噪声地板，特征工程/换算法收益有限，应改评估口径；\n② 若σ≤0.8%，说明误差主要来自模型不足，特征工程（工作面/滞后特征）值得投入。"],
    ["原理", "同一时间点取两份独立样品（平行样），分别化验。两份结果的差值 d=A−B 只反映测量随机误差。\n重复性标准差 σ = √( Σd² / (2n) )，n=平行样组数。"],
    ["取样点", "315筛子取样口——必须与表1中「315灰分」列的取样点完全一致（否则测的不是同一个量）。"],
    ["取样方法", "同一时间点连续取两份样（间隔≤5分钟），分别装袋、分别编号（A/B），严禁混样；袋内物料状态与平时送检一致。"],
    ["化验要求", "同一化验室、同一化验员、同一台仪器，两袋样当天完成化验，避免跨天仪器漂移混入。"],
    ["组数要求", "至少5组；建议8~10组，覆盖：\n① 不同班组（早/中/夜班各2~3组）；\n② 不同煤质（原煤灰分高、低各占一半）；\n③ 尽量分散在不同日期，不要一天做完。"],
    ["记录要求", "按「记录表」Sheet逐项填写；d 值可先留空，出结果后填。"],
    ["安全/合规", "按现场取样安全规程执行；样品编号与化验单一致以便追溯。"],
]
for r in rows1:
    ws.append(r)
style_sheet(ws, [14, 100])
body(ws)

# ============ Sheet2 记录表 ============
ws2 = wb.create_sheet("记录表")
ws2.append(["序号", "日期", "取样时间", "班次", "取样人", "样号A", "样号B",
            "A灰分%", "B灰分%", "差值d=A−B", "化验人", "备注"])
# 第1组：已完成的平行样
ws2.append([1, None, None, None, None, None, None, 13.49, 14.26, -0.77, None, "平行样第1组（已出结果）"])
for i in range(2, 11):
    ws2.append([i, None, None, None, None, None, None, None, None, None, None, None])
style_sheet(ws2, [6, 12, 12, 8, 10, 10, 10, 10, 10, 12, 10, 20])
body(ws2)
# 汇总提示行
sum_row = ws2.max_row + 2
ws2.cell(row=sum_row, column=1, value="汇总").font = Font(name="微软雅黑", size=10, bold=True)
ws2.cell(row=sum_row, column=2, value="σ = √( Σd² / (2n) )，n=已完成的平行样组数").font = BODY_FONT
ws2.cell(row=sum_row, column=3, value="当前 n=1：σ = √(0.77²/(2×1)) = 0.5445%").font = BODY_FONT
ws2.cell(row=sum_row + 1, column=1, value="判定").font = Font(name="微软雅黑", size=10, bold=True)
ws2.cell(row=sum_row + 1, column=2, value="σ≈0.54% ≤ 0.8% → 测量噪声小（初步结论；仅1组，需补足≥5组后定论）").font = BODY_FONT
ws2.cell(row=sum_row + 2, column=1, value="说明").font = Font(name="微软雅黑", size=10, bold=True)
ws2.cell(row=sum_row + 2, column=2, value="d=A灰分−B灰分；两袋结果都出来后再填 d，并在备注里注明异常（如取样中断、样袋破损）。建议继续补：不同班组、不同煤质、不同日期各2~3组。").font = BODY_FONT

# ============ Sheet3 结果判定 ============
ws3 = wb.create_sheet("结果判定")
ws3.append(["σ 范围", "判定", "后续动作"])
rows3 = [
    ["σ ≤ 0.8%", "测量噪声小，模型有较大提升空间", "按计划推进特征工程（工作面one-hot、滞后特征），模型RMSE有望向1.5%以内压缩"],
    ["0.8% < σ ≤ 1.5%", "中等噪声，模型仍有空间但天花板受限", "特征工程照做，同时准备评估口径调整（±1.0/±1.5%容差）"],
    ["σ > 1.5%", "噪声地板：RMSE 2.4% 中大部分是测量误差", "暂停重投入特征工程；改评估口径（放宽容差、EMA平滑已在用）；排查化验/制样环节；讨论'筛下水处理后灰分'口径"],
    ["附加判断", "若 σ 与 RMSE(2.4%) 相当", "换算法（GBDT等）收益预期很小，优先解决数据质量与口径问题"],
]
for r in rows3:
    ws3.append(r)
style_sheet(ws3, [14, 40, 66])
body(ws3)
for row in ws3.iter_rows(min_row=2):
    if row[0].value and str(row[0].value).startswith("σ ≤"):
        row[0].fill = OK_FILL
    elif row[0].value and str(row[0].value).startswith("σ >"):
        row[0].fill = WARN_FILL

# ============ Sheet4 示例计算 ============
ws4 = wb.create_sheet("示例计算")
ws4.append(["示例", "A灰分%", "B灰分%", "d=A−B", "d²"])
rows4 = [
    ["第1组", 15.2, 16.1, -0.9, 0.81],
    ["第2组", 13.8, 13.5, 0.3, 0.09],
    ["第3组", 18.0, 17.2, 0.8, 0.64],
    ["第4组", 14.6, 15.3, -0.7, 0.49],
    ["第5组", 16.3, 16.0, 0.3, 0.09],
    ["合计", None, None, None, "Σd²=2.12"],
    ["计算", None, None, None, "σ = √(2.12/(2×5)) = √0.212 ≈ 0.46% → 判定：测量噪声小"],
]
for r in rows4:
    ws4.append(r)
style_sheet(ws4, [10, 12, 12, 12, 16])
body(ws4)

wb.save(OUT)
print("[gen]", OUT)

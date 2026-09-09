# -*- coding: utf-8 -*-
"""生成 Excel：前后端分离后的差异（分板块）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\前后端分离-差异清单.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BEFORE_FILL = PatternFill("solid", fgColor="F2F2F2")
AFTER_FILL = PatternFill("solid", fgColor="E2F0D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def make_table(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER
    # 若含「分离前/分离后」两列，着色区分
    if headers[1] == "分离前" and len(headers) >= 3 and headers[2] == "分离后":
        for row in ws.iter_rows(min_row=2):
            row[1].fill = BEFORE_FILL
            row[2].fill = AFTER_FILL
    return ws


wb = Workbook()
wb.remove(wb.active)

# ---------------- 1 架构与部署 ----------------
make_table(wb, "1架构与部署差异",
    ["维度", "分离前", "分离后"],
    [
        ["运行方式", "file:// 直接打开 index.html（纯前端）", "http://127.0.0.1:8000/ 由 FastAPI 托管（同源）"],
        ["后端", "无", "FastAPI + uvicorn + SQLAlchemy 2 + Alembic + Pydantic v2"],
        ["前端", "原生 JS + localStorage（无构建）", "原生 JS 不变，新增 js/api.js 数据层（双轨镜像）"],
        ["静态资源", "本地磁盘文件", "FastAPI StaticFiles 挂载 FRONTEND_DIR（web(2)）"],
        ["跨域", "无（本地）", "同源访问规避 CORS（allow_origins=* 备用）"],
        ["技术栈", "仅浏览器 JS", "Python 3.11 + numpy + scikit-learn + openpyxl"],
    ],
    [14, 34, 52])

# ---------------- 2 数据存储 ----------------
make_table(wb, "2数据存储差异",
    ["维度", "分离前", "分离后"],
    [
        ["存储介质", "localStorage（dmcs_store 单 JSON）", "SQLite（开发）/ MySQL、PostgreSQL（生产）"],
        ["数据模型", "App.store 单个对象", "11 张表（coal_records 统一主表）"],
        ["写入", "saveStore → localStorage.setItem", "双轨：file:// 仍 localStorage；http:// 额外 PUT /state 镜像"],
        ["读取", "loadStore → localStorage.getItem", "GET /state（load_store DB→store 桥接）"],
        ["时间戳", "字符串 YYYY-MM-DD HH:MM:SS", "同样字符串（字符串序==时间序，规避时区）"],
        ["一致性保障", "无", "PUT /state 整库重写单事务原子（migrate.replace）"],
    ],
    [14, 36, 50])

# ---------------- 3 数据库表 ----------------
make_table(wb, "3数据库表清单",
    ["表名", "用途", "对应 store 键", "关键字段"],
    [
        ["coal_records", "三表合并主表（UNIQUE(category,ts,system)）",
         "coarseCoal / floatCoal / calcLogs(ash_density)",
         "id, category(coarse|float|ash_density), ts, system, source, ash_content, coal_amount, density, level, raw_ash, moisture, belt, sysA, sysB, sys401, sys402, desliming473, desliming474, is_stoppage, mining_face, filter_press_running, influence_value, annotation, predicted_ash, extra, import_log_id(FK), created_at"],
        ["calc_logs", "补录日志（灰分仪/密度计/皮带秤）",
         "calcLogs(非 ash_density)",
         "id, ts, calc_type(ash_meter|density_meter|belt_scale), system, input_json, output_json, created_at"],
        ["manual_entries", "手工补录历史",
         "manualEntries",
         "id, ts, category, values(JSON), remark, operator, status, created_at"],
        ["heavy_samples", "重介精煤灰分采样（训练积累）",
         "heavySamples",
         "id, ts, rho, ash_content, source, created_at"],
        ["import_logs", "Excel 导入日志",
         "importLogs",
         "id, ts, category, filename, total, success, failed, skipped, status, errors(JSON), created_at"],
        ["alerts", "告警",
         "alerts",
         "id, system, level, message, ts, ack, created_at"],
        ["coarse_models", "当前粗灰模型（mlr+pls 两行，UNIQUE(method)）",
         "coarseModel",
         "id, method(mlr|pls), is_current, train_run_id, feature_names(JSON), intercept, coefs(JSON), means(JSON), stds(JSON), std_coef(JSON), metrics(JSON:r2/adjR2/rmse/mae/passRate1/15/q2/q2Time/lambda/drop), impute_means(JSON), pls_A, n, tolerance, train_range, trained_at, created_at"],
        ["coarse_model_history", "重训练历史（detail JSON 全量快照）",
         "coarseModelHistory",
         "id, model_id(FK), train_run_id, trained_at, n, tolerance, train_range, production, mlr_r2, pls_r2, mlr_q2, pls_q2, pass_rate, detail(JSON), created_at"],
        ["regression_models", "单因素密度模型（linear/poly2/poly3）",
         "regressionModels",
         "id, model_type(linear|poly2|poly3), params(JSON), r_squared, rmse, data_points, created_at"],
        ["settings", "键值配置（8 个单值）",
         "ashTarget/ashTargetTol/guideScheme/totalAshManualOn/densityGuide/densityAutoOn/coarseTolerance/coarseTrainRange",
         "key(PK), value(JSON), value_type, updated_at"],
        ["auto_state", "录入层 + 手动有效期（10 个对象键，整 JSON 无损）",
         "amountInputs/ashInputs/instrumentInputs/heavyAshInput/coarseAshInput/floatAshInput/coarseCalc/coarseAshEma/autoState/heavyAshBackcalc",
         "key(PK), value(JSON), updated_at"],
    ],
    [20, 34, 34, 66])

# ---------------- 4 API 接口 ----------------
make_table(wb, "4接口清单",
    ["方法", "路径（/api/v1 下）", "说明"],
    [
        ["GET", "/health", "健康检查"],
        ["GET", "/state", "DB → 前端 store 等价结构"],
        ["PUT", "/state", "整库快照重写（migrate.replace 单事务）"],
        ["GET", "/settings", "配置列表"],
        ["GET/PUT", "/settings/{key}", "单配置读写"],
        ["GET/PUT", "/inputs", "录入层（instrumentInputs 等）"],
        ["GET", "/records", "三表数据查询（category/ts 过滤，total=真实总数）"],
        ["POST", "/records", "upsert（按 category+ts+system 唯一键）"],
        ["DELETE", "/records/{id}", "删除单条"],
        ["POST", "/samples/heavy-ash", "重介灰分采样"],
        ["GET/POST", "/manual-entries", "手工补录"],
        ["GET", "/overview/dashboard", "总览（取值链 + 密度建议）"],
        ["POST", "/migrate/preview", "localStorage 迁移 dry-run"],
        ["POST", "/migrate/localstorage", "localStorage 完整迁移入库"],
        ["POST", "/import", "Excel 数据导入"],
        ["POST", "/training/coarse-model?range=jun_jul|30d|all", "MLR/PLS 后端 sklearn 训练 + 持久化 + 回填"],
    ],
    [10, 46, 48])

# ---------------- 5 算法一致性 ----------------
make_table(wb, "5算法一致性",
    ["模块", "函数", "验证方式"],
    [
        ["密度建议", "calc_total_ash / expert_adjust / compute_density_guidance", "逐位移植 + 单测 + golden 对拍"],
        ["预测", "predict_coarse_ash（缺失因子均值补全）", "逐位移植 + golden 对拍"],
        ["简报", "build_hourly_brief（16 列，递归软测量）", "1069 行全量逐值 diff=0"],
        ["取值链", "resolve_*（手动>公式>录入>默认）", "golden 对拍"],
        ["训练", "trainMlr/trainPls/trainCoarseModel（hat-LOOCV 选 λ、NIPALS+LOOCV 选 A）", "纯 Python 逐位移植 diff=0.0"],
        ["sklearn 训练", "Ridge / PLSRegression（λ/A 选优复用精确移植）", "与移植版差 ~1e-13"],
        ["时间戳", "字符串时间戳 + 30d/all 范围过滤", "headless 对拍 JS：jun=113/30d=55/all=113"],
    ],
    [14, 50, 40])

# ---------------- 6 模型训练 ----------------
make_table(wb, "6模型训练差异",
    ["维度", "分离前", "分离后"],
    [
        ["训练位置", "前端客户端 JS（手写高斯消元 + NIPALS）", "后端 sklearn（Ridge + PLSRegression）"],
        ["λ 网格", "[0,1e-6..1e-2]", "[0,1e-6..1e-2,1e-1,1]（扩宽，mlr walk-forward Q² 0.0406→0.1121）"],
        ["触发方式", "前端 trainCoarseModel()（本地）", "重训练按钮/导入/训练模型 → POST /training/coarse-model；file:// 回退本地"],
        ["持久化", "localStorage.coarseModel", "coarse_models（两行+train_run_id）+ coarse_model_history（detail JSON）"],
        ["回填", "前端 refreshCoarsePredictions", "后端回填 predicted_ash(round4) + 前端刷新同口径"],
        ["选型依据", "q2Time(优先)/q2(回退) 选 production", "同口径（逐位一致）"],
    ],
    [14, 38, 52])

# ---------------- 7 双轨行为 ----------------
make_table(wb, "7双轨行为差异",
    ["维度", "file:// 打开", "http:// 打开"],
    [
        ["数据存储", "仅 localStorage", "localStorage + 后端 SQLite 镜像（PUT /state）"],
        ["模型训练", "本地 JS 训练（旧网格同步已扩宽）", "后端 sklearn 训练（POST /training/coarse-model）"],
        ["后端不可用", "无影响", "训练/镜像自动回退本地（静默）"],
        ["数据来源", "种子 seed_data.js", "GET /state（DB→store），种子由 migrate 入库"],
        ["适用场景", "离线/演示/单机", "生产/多端共享/持久化"],
    ],
    [14, 40, 46])

# ---------------- 8 验收与测试 ----------------
make_table(wb, "8验收与测试",
    ["项", "结论"],
    [
        ["后端测试", "pytest 48 全绿（含训练 parity、sklearn 等价、API 端到端、30d/all 范围）"],
        ["预测 golden", "简报 1069 行全量逐值 diff=0；取值链/密度建议 golden 对拍一致"],
        ["训练 golden", "纯 Python 移植 vs Node oracle 逐位一致 diff=0.0；sklearn 等效 ~1e-13"],
        ["前端回归", "headless Edge CDP：重训练走远程 PASS、30d/all 计数与后端一致"],
        ["迁移", "localStorage → DB 无损（migrate.preview/apply/replace 单事务）"],
        ["待办", "接真实生产数据后复核 λ/生产模型选择（当前为 seed 演示数据）"],
    ],
    [16, 90])

# ---------------- 9 索引与约束 ----------------
make_table(wb, "9索引与约束",
    ["表名", "主键", "唯一约束", "索引", "外键"],
    [
        ["coal_records", "id", "UNIQUE(category, ts, system) 名称 uq_coal_cat_ts_sys", "ix_coal_ts(ts)；ix_coal_cat_ts(category, ts)", "import_log_id → import_logs.id"],
        ["calc_logs", "id", "—", "ix_calc_type_ts(calc_type, ts)", "—"],
        ["manual_entries", "id", "—", "—", "—"],
        ["heavy_samples", "id", "—", "—", "—"],
        ["import_logs", "id", "—", "—", "—"],
        ["alerts", "id", "—", "—", "—"],
        ["coarse_models", "id", "UNIQUE(method) 名称 uq_coarse_model_method", "—", "—"],
        ["coarse_model_history", "id", "—", "—", "model_id → coarse_models.id（预留，当前不赋值）"],
        ["regression_models", "id", "—", "—", "—"],
        ["settings", "key", "—（key 即主键）", "—", "—"],
        ["auto_state", "key", "—（key 即主键）", "—", "—"],
    ],
    [18, 10, 40, 32, 34])

# ---------------- 10 迁移 SQL ----------------
make_table(wb, "10迁移SQL",
    ["项", "内容"],
    [
        ["建表方式", "开发：Base.metadata.create_all（main.py 启动自动建表）；生产：alembic upgrade head"],
        ["迁移 0001", "27cd61d87298_init：create_table 全部 11 张表 + 索引/约束（见 migrations/versions/）"],
        ["迁移 0002", "b7e2c4d9a001_training_model_columns：coarse_models 加 train_run_id + method 唯一约束；coarse_model_history 加 train_run_id + detail(JSON)"],
        ["主表关键 DDL", "CREATE TABLE coal_records (id INTEGER PRIMARY KEY AUTOINCREMENT, category VARCHAR(16) NOT NULL, ts VARCHAR(19) NOT NULL, system VARCHAR(32) NOT NULL, ... );\nCONSTRAINT uq_coal_cat_ts_sys UNIQUE (category, ts, system)"],
        ["主表索引", "CREATE INDEX ix_coal_ts ON coal_records (ts);\nCREATE INDEX ix_coal_cat_ts ON coal_records (category, ts);"],
        ["模型唯一", "CREATE UNIQUE INDEX uq_coarse_model_method ON coarse_models (method);"],
        ["0002 增量 SQL", "ALTER TABLE coarse_models ADD COLUMN train_run_id VARCHAR(36);\nALTER TABLE coarse_model_history ADD COLUMN train_run_id VARCHAR(36);\nALTER TABLE coarse_model_history ADD COLUMN detail JSON;"],
    ],
    [16, 96])

wb.save(OUT)
print("[gen]", OUT)

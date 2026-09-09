# -*- coding: utf-8 -*-
"""生成 Excel 文档：浏览器自动化测试-说明（AI 执行的测试规范）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\浏览器自动化测试-说明.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


wb = Workbook()

ws = wb.active
ws.title = "测试能力与方式"
ws.append(["序号", "能力", "方式", "说明"])
rows = [
    [1, "真实页面加载", "无头 Edge（headless）+ CDP 协议", "加载 web(2) 的 file:// 页面，全部真实脚本执行"],
    [2, "页面交互", "CDP Runtime.evaluate / DOM 命令", "点击导航切页、改下拉框、勾选复选框、触发事件"],
    [3, "文件导入", "CDP DOM.setFileInputFiles", "把桌面三张表直接喂给导入页，走完整解析→预览→确认导入流程"],
    [4, "数据校验", "读取 App.store / DOM", "检查导入行数、错误数、各数组落库、模型重训结果"],
    [5, "在线仪表联动", "读取 resolveInstrument / 层级", "验证各仪表行 手动>录入>默认 取值与来源标签"],
    [6, "刷新持久化", "CDP Page.reload", "导入后强制刷新，确认数据与导入日志不丢失"],
    [7, "布局与像素", "CDP getBoundingClientRect", "验证面板行内元素不溢出、趋势图点间距分开等"],
    [8, "逻辑单测", "Node + vm 加载 JS", "纯函数级断言（取值链、公式、限幅、死区等），速度快"],
]
for r in rows:
    ws.append(r)
style_sheet(ws, [6, 18, 34, 60])
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

ws2 = wb.create_sheet("每次交付的标准测试清单")
ws2.append(["阶段", "检查项", "验证点"])
rows2 = [
    ["语法", "node --check 所有改动的 JS", "无语法错误"],
    ["逻辑", "Node 单测（公式/取值链/保护规则）", "断言全部 PASS"],
    ["导入", "三张表逐一走导入 UI", "成功数与失败数、落库行数、幂等覆盖"],
    ["联动", "在线仪表各行 值/来源层级", "导入数据成为录入层；无数据回默认"],
    ["计算", "总灰分/反推/密度指导", "数值符合公式，方向正确，限幅生效"],
    ["持久化", "导入→刷新→复查", "数据与导入日志保留（__merged 不重置）"],
    ["界面", "DOM 渲染/布局几何/面板行", "无 undefined/NaN、无溢出、点分开"],
    ["汇报", "汇总 PASS/FAIL 与版本号", "附验证脚本路径，可复跑"],
]
for r in rows2:
    ws2.append(r)
style_sheet(ws2, [8, 30, 50])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

ws3 = wb.create_sheet("局限与约定")
ws3.append(["序号", "事项", "说明"])
rows3 = [
    [1, "测试环境独立", "测试在无头 Edge + 全新临时用户目录中进行，不会动你日常使用的浏览器数据"],
    [2, "视觉确认需你参与", "当前会话的模型无法查看截图像素，视觉类问题（配色/观感）我会用几何测量判断，最终观感请你确认"],
    [3, "测试脚本留档", "全部验证脚本存于 backend/scripts/test_*.js / verify_*.js，可随时复跑"],
    [4, "固定交付节奏", "每次改动：改代码 → node 单测 → 无头浏览器端到端 → 汇报结果与版本号"],
]
for r in rows3:
    ws3.append(r)
style_sheet(ws3, [6, 18, 90])
for row in ws3.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

wb.save(OUT)
print("[gen]", OUT)

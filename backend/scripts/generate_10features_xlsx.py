# -*- coding: utf-8 -*-
"""生成 Excel 文档：粗精煤泥灰分模型-10特征清单"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\粗精煤泥灰分模型-10特征清单.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "10特征清单"
ws.append(["序号", "特征字段", "中文名", "取值", "数据来源(表1列)", "PLS标准化系数(权重)", "MLR标准化系数(权重)", "影响方向/说明"])

rows = [
    [1, "raw_ash", "原煤灰分", "%", "原煤灰分%", 0.067562, 0.071763,
     "原煤越脏，煤泥灰分越高（正相关）"],
    [2, "coal_amount", "小时带煤量", "t/h", "小时带煤量t/h", -0.108522, -0.166564,
     "带煤量越大，煤泥灰分越低（负荷稀释）"],
    [3, "sysA", "A系统开启", "0/1", "开启的系统(A/B/401/402)", -0.120631, -0.268344,
     "开A系统时煤泥灰分偏低"],
    [4, "sysB", "B系统开启", "0/1", "同上", 0.087052, 0.227975,
     "开B系统时煤泥灰分偏高"],
    [5, "sys401", "401系统开启", "0/1", "同上", -0.126220, 0.009323,
     "PLS口径下开401灰分偏低；MLR口径下影响很小"],
    [6, "sys402", "402系统开启", "0/1", "同上", 0.136696, 0.262920,
     "开402系统时煤泥灰分偏高"],
    [7, "desliming473", "473脱粉开启", "0/1", "脱粉(473/474)", 0.068292, 0.093268,
     "473脱粉开启，煤泥灰分略升"],
    [8, "desliming474", "474脱粉开启", "0/1", "同上", -0.073943, -0.049639,
     "474脱粉开启，煤泥灰分略降"],
    [9, "is_stoppage", "停机/低负荷", "0/1", "表1无此列：导入时按 带煤量≤10t/h 自动推断", 0.0, 0.0,
     "系数置0（原表无停机数据），字段预留"],
    [10, "level", "精磁尾液位", "%", "精磁尾液位%", -0.653220, -0.633163,
     "权重最大：液位越低煤泥灰分越高（跑粗），是最关键因子"],
]
for r in rows:
    ws.append(r)
for i, w in enumerate([6, 14, 14, 12, 26, 20, 20, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for cell in ws[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = CENTER
    cell.border = BORDER
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column not in (1,) else CENTER
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:H{ws.max_row}"

# Sheet2 权重排序与说明
ws2 = wb.create_sheet("权重排序与说明")
ws2.append(["排序(PLS)", "特征", "|标准化系数|", "含义"])
ranking = [
    [1, "level 精磁尾液位", 0.653, "最大影响因子：液位反映分选状态，液位低→跑粗→煤泥灰分高"],
    [2, "sys402 402系统开启", 0.137, "开启402系统时煤泥灰分偏高"],
    [3, "sys401 401系统开启", 0.126, "开启401系统时煤泥灰分偏低"],
    [4, "sysA A系统开启", 0.121, "开启A系统时煤泥灰分偏低"],
    [5, "coal_amount 小时带煤量", 0.109, "带煤量高→灰分低（稀释效应）"],
    [6, "sysB B系统开启", 0.087, "开启B系统时煤泥灰分略高"],
    [7, "desliming474 474脱粉", 0.074, "474脱粉开启→灰分略降"],
    [8, "desliming473 473脱粉", 0.068, "473脱粉开启→灰分略升"],
    [9, "raw_ash 原煤灰分", 0.068, "原煤灰分高→煤泥灰分高"],
    [10, "is_stoppage 停机/低负荷", 0.0, "系数置0（原表无此列，按带煤量≤10推断，预留字段）"],
]
for r in ranking:
    ws2.append(r)
for i, w in enumerate([12, 26, 16, 60], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for cell in ws2[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = CENTER
    cell.border = BORDER
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER
ws2.freeze_panes = "A2"

note_rows = [
    "说明：",
    "1. 目标变量 = 315灰分%（粗精煤泥灰分），模型为出厂 PLS（生产模型）与 MLR（对照），样本 113 条（2026-06-16~07-14，5分钟采样）。",
    "2. 标准化系数(stdCoef)比较权重大小，符号表示方向：正→灰分升高，负→灰分降低；系数为 0 表示该特征在当前模型中被置零。",
    "3. 四个系统开关（sysA/sysB/sys401/sys402）是任务一合并后按用户要求保留的算法特征，代表“开启哪几套系统”的工况。",
    "4. is_stoppage 在原表中没有对应列，导入时按 小时带煤量≤10t/h 自动推断为1，模型系数当前为0（数据不足，未学到停机效应），字段已预留。",
]
r = ws2.max_row + 2
for n in note_rows:
    ws2.cell(row=r, column=1, value=n)
    r += 1

wb.save(OUT)
print("[gen]", OUT)

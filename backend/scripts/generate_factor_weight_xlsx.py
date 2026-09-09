# -*- coding: utf-8 -*-
"""生成 Excel 文档：粗精煤泥-多因素影响权重计算说明"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\粗精煤泥-多因素影响权重计算说明.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


wb = Workbook()

ws = wb.active
ws.title = "计算原理"
ws.append(["步骤", "环节", "做法", "说明"])
rows = [
    [1, "准备样本",
     "训练集 = 表1《粗精煤泥灰分影响因素》全部记录（ash>0），X=10个特征，y=315灰分%",
     "缺失特征用列均值补全；常数列（方差为0，如停机标志）先剔除"],
    [2, "标准化",
     "每个特征 z = (x − 均值) ÷ 标准差",
     "消除量纲（灰分%、t/h、0/1开关、液位%），使各特征可比"],
    [3, "MLR 求系数",
     "标准化正规方程：解 (X'X)β = X'y，得标准化系数 β*",
     "β* 含义：特征变化1个标准差 → 灰分变化 β* 个百分点"],
    [4, "还原原尺度",
     "原尺度系数 coefs[j] = β*[j] ÷ stds[j]，截距同步换算",
     "用于对新记录直接预测：灰分 = 截距 + Σ 系数×特征值"],
    [5, "标准化系数 stdCoef（权重）",
     "stdCoef[j] = coefs[j] × stds[j] ÷ 灰分标准差",
     "即特征每变化1个标准差，灰分变化 stdCoef 个标准差"],
    [6, "PLS 求系数",
     "NIPALS 提取潜变量（分量数由留一交叉验证 Q² 选优），再换算回每个原始特征的系数",
     "同样输出 coefs / stdCoef，口径与 MLR 一致"],
    [7, "影响权重展示",
     "权重 = |stdCoef|，按从大到小排序；符号表方向（正→灰分升高↑，负→灰分降低↓）",
     "粗精煤泥分析页『因子权重』条形图：绿色=正向、红色=负向，标签带↑↓"],
]
for r in rows:
    ws.append(r)
style_sheet(ws, [6, 14, 52, 44])
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

ws2 = wb.create_sheet("当前权重对照")
ws2.append(["排序(按|PLS|)", "特征", "中文名", "PLS stdCoef", "MLR stdCoef", "解读"])
rows2 = [
    [1, "level", "精磁尾液位(%)", -0.653, -0.633, "权重最大：液位低→跑粗→煤泥灰分高"],
    [2, "sys402", "402系统开启", 0.137, 0.263, "开402时煤泥灰分偏高"],
    [3, "sys401", "401系统开启", -0.126, 0.009, "PLS口径下开401灰分偏低"],
    [4, "sysA", "A系统开启", -0.121, -0.268, "开A时灰分偏低"],
    [5, "coal_amount", "小时带煤量(t/h)", -0.109, -0.167, "带煤量↑→灰分↓（稀释）"],
    [6, "sysB", "B系统开启", 0.087, 0.228, "开B时灰分略高"],
    [7, "desliming474", "474脱粉开启", -0.074, -0.050, "474脱粉→灰分略降"],
    [8, "desliming473", "473脱粉开启", 0.068, 0.093, "473脱粉→灰分略升"],
    [9, "raw_ash", "原煤灰分(%)", 0.068, 0.072, "原煤越脏→灰分越高"],
    [10, "is_stoppage", "停机/低负荷", 0.0, 0.0, "表1无此列，系数置0（预留）"],
]
for r in rows2:
    ws2.append(r)
style_sheet(ws2, [14, 16, 18, 12, 12, 40])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = CENTER if cell.column != 6 else WRAP

ws3 = wb.create_sheet("要点说明")
ws3.append(["序号", "要点", "说明"])
rows3 = [
    [1, "权重含义",
     "stdCoef 是标准化系数：该特征变化 1 个标准差，315灰分变化 stdCoef 个灰分标准差。所有特征量纲一致，因此 |stdCoef| 可直接比较相对重要性。"],
    [2, "方向判断",
     "符号为正 → 特征增大灰分升高（条形绿色↑）；符号为负 → 特征增大灰分降低（条形红色↓）。"],
    [3, "当前结论",
     "精磁尾液位是最强因子（|stdCoef|≈0.65，远大于其他）；系统开关（402/A/401/B）与带煤量是第二梯队；脱粉、原煤灰分影响较小；停机项因数据缺失置0。"],
    [4, "动态更新",
     "粗精页『重训练』→ trainCoarseModel → 重训 MLR+PLS（标准化正规方程 / NIPALS，含留一交叉验证 Q² 选优）→ 更新 model.stdCoef → 因子权重图自动刷新。切换『算法视图』可对比 MLR 与 PLS 的权重。"],
    [5, "与出厂模型的关系",
     "出厂默认 PLS/MLR 的 stdCoef 由 build_web2_seed.py 用同一口径离线训练；页面显示的是当前生效模型（重训后为新值）。"],
    [6, "注意事项",
     "标准化系数受样本分布影响：样本少/存在共线（四开关常同开同停）时权重会波动，PLS 自带降维更稳健，因此生产模型按 Q² 选 PLS。"],
]
for r in rows3:
    ws3.append(r)
style_sheet(ws3, [6, 16, 100])
for row in ws3.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

wb.save(OUT)
print("[gen]", OUT)

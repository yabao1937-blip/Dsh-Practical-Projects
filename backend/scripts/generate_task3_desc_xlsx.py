# -*- coding: utf-8 -*-
"""生成 Excel 文档：任务三-总精煤灰分公式与反推-说明"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\任务三-总精煤灰分公式与反推-说明.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


wb = Workbook()

# ================= Sheet1 需求与公式 =================
ws = wb.active
ws.title = "需求与公式"
ws.append(["序号", "内容", "公式", "说明"])
rows = [
    [1, "总精煤灰分（正算）",
     "总精煤灰分 = (重介精煤灰分×重介精煤量 + 浮精灰分×浮精量 + 粗精煤泥灰分×粗精煤泥量) ÷ 总煤量",
     "总煤量 = 重介精煤量 + 浮精量 + 粗精煤泥量（任务二“量数据”面板四量已就绪）"],
    [2, "反推重介精煤灰分（核心）",
     "重介精煤灰分 = (总精煤灰分×总煤量 − 浮精灰分×浮精量 − 粗精煤泥灰分×粗精煤泥量) ÷ 重介精煤量",
     "需重介精煤量>0；获取到总精煤灰分后，用该公式反推重介精煤灰分"],
    [3, "现状",
     "calcTotalAsh 正算公式已在用；但重介精煤灰分目前是固定 8.50%，不是反推值",
     "任务三落地后：重介精煤灰分由“定值”变为“反推计算值”"],
    [4, "用途链",
     "反推出的重介精煤灰分 → 回代正算公式 → 总览卡片实际灰分/偏差/告警、浮精页影响值基准、后续密度建议",
     "反推值成为系统主计算链的一环"],
]
for r in rows:
    ws.append(r)
style_sheet(ws, [6, 22, 62, 40])
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

# ================= Sheet2 现状 vs 目标 =================
ws2 = wb.create_sheet("现状vs目标")
ws2.append(["环节", "现状", "任务三目标"])
rows2 = [
    ["重介精煤灰分", "定值 8.50%（hardcoded，总览卡片/详情弹窗/告警共用）", "由总精煤灰分反推，动态计算"],
    ["总精煤灰分", "正算公式 calcTotalAsh 在卡片上计算显示", "保留正算；同时增加“总精煤灰分”的输入口径（化验/皮带灰分仪），作为反推的已知量"],
    ["四量", "量数据面板：总精煤量(录入)/浮精量(录入或自动)/粗精煤泥量(录入)/重介精煤量(录入或计算=总−浮−粗)", "继续复用，反推公式的分母与各项直接取面板当前值"],
    ["粗精煤泥灰分", "PLS/MLR 10特征模型前馈预测", "不变（反推公式里用预测值或化验值均可，需定口径）"],
    ["浮精灰分", "表2最新化验值", "不变"],
    ["建议密度", "表3最新密度值+方向建议", "任务三明确重介灰分反推后，再定密度建议口径（可联动改进）"],
]
for r in rows2:
    ws2.append(r)
style_sheet(ws2, [16, 50, 50])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

# ================= Sheet3 实现要点与疑问 =================
ws3 = wb.create_sheet("实现要点与疑问")
ws3.append(["序号", "类别", "内容"])
rows3 = [
    [1, "要点1：总精煤灰分的取值来源",
     "反推需要“总精煤灰分”这个已知量。可选来源：①表3 灰分密度里的皮带灰分(belt_ash，501/502 灰分仪)；②化验室总精煤化验值手工录入；③两皮带灰分按煤量加权。需与现场确认用哪种。"],
    [2, "要点2：反推时机",
     "每当有新的总精煤灰分数据（化验/皮带灰分仪新值）时反推一次；其余时间沿用最近一次反推值参与正算。"],
    [3, "要点3：异常保护",
     "重介精煤量≤0 时不反推；反推结果需做上下限钳制（如 5%~13%）与突变平滑，避免单个异常灰分把重介灰分带飞。"],
    [4, "要点4：反推值落点",
     "建议存 store（如 amountInputs 旁新增 heavyAshInput {mode:'auto'/'manual', value}），量数据面板加一行“重介精煤灰分(反推)”显示并支持人工覆盖；总览卡片/详情/告警全部改读该值。"],
    [5, "要点5：口径统一",
     "反推值同时替换浮精页固定影响值的“重介灰分”动态取值口径（浮精页已用 getAshByTime 动态灰分，需决定是否改用反推值）与总览 8.50 定值。"],
    [6, "疑问1",
     "总精煤灰分按 501/502 两个皮带灰分加权取，还是按化验值录入？两皮带灰分仪数据在表3已可导入（belt_ash 接口已预留）。"],
    [7, "疑问2",
     "反推出的重介精煤灰分是否需要平滑/限幅？异常值处理规则（连续N次超限才更新？单次限幅？）。"],
    [8, "疑问3",
     "粗精煤泥灰分用模型前馈预测值还是化验实测值参与反推？（预测值及时、实测值滞后）"],
]
for r in rows3:
    ws3.append(r)
style_sheet(ws3, [6, 22, 100])
for row in ws3.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER
        if cell.column == 2 and str(cell.value or "").startswith("疑问"):
            cell.fill = WARN_FILL
ws3.auto_filter.ref = f"A1:C{ws3.max_row}"

wb.save(OUT)
print("[gen]", OUT)

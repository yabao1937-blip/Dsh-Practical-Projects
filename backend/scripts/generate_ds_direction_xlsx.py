# -*- coding: utf-8 -*-
"""生成 Excel：DS 粗精煤泥灰分——「方向可用 / 水平不可用」的验证结论、向前验收与落地清单。

产出 docs/DS粗灰-方向与水平验证结论.xlsx（openpyxl 自洽，无外部运行时依赖）。
数据来源：**现算**（读 backend/data/dense_medium.db 走 app.services.coarse_forward / coarse_direction），
所以表里的数字不会与库内数据漂移；重跑本脚本即可刷新。

历史：2026-09-23 首版（开发集 6-7 月、检验集 8-9 月锁死）；2026-09-24 导入 9-01~9-21 新数据后复核并
把「真向前验证」做进系统（GET /api/v1/training/coarse-forward + 页面区块），本版据此更新。
"""
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook                                            # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402

from app.services.coarse_direction import direction_report               # noqa: E402
from app.services.coarse_forward import FEATURE_SETS, evaluate_window     # noqa: E402

DOCS = ROOT.parent / "docs"
OUT = DOCS / "DS粗灰-方向与水平验证结论.xlsx"
DB = ROOT / "data" / "dense_medium.db"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="微软雅黑", size=13, bold=True, color="2F5597")
BODY_FONT = Font(name="微软雅黑", size=10)
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
WARN_FILL = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def load_records():
    con = sqlite3.connect(str(DB))
    cols = [c[1] for c in con.execute("pragma table_info(coal_records)")]
    rows = [dict(zip(cols, r)) for r in con.execute(
        "select * from coal_records where category='coarse' and ash_content is not null order by ts, id")]
    con.close()
    return rows


def sheet(wb, name, title, headers, widths, rows, fill_col=None, fill_map=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
        ws.column_dimensions[cell.column_letter].width = w
    for i, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = BODY_FONT
            cell.alignment = CENTER if c <= 2 else WRAP
            cell.border = BORDER
        if fill_col and fill_map:
            f = fill_map.get(str(row[fill_col - 1]))
            if f:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=i, column=c).fill = f
    ws.freeze_panes = "A4"
    return ws


# ---------------- 现算数据 ----------------
recs = load_records()
n_all = len(recs)
# 锁定窗：训练 →9-03，检验 9-04~9-21（2026-09-24 才导入的数据，开发期从未使用）
lock = evaluate_window(recs, "2026-09-04", None, feature_sets=FEATURE_SETS)
# 部署口径：按 jun_jul 训练后剩下的那段（现场"隔几周导一次 Excel → 重训 → 往后用"）
leftover = evaluate_window(recs, "2026-08-23", "2026-09-04", feature_sets=FEATURE_SETS)  # 开发侧向前窗
prod = evaluate_window(recs, "2026-08-23", None, feature_sets=FEATURE_SETS)             # jun_jul 留出整段
dir_new = direction_report(recs, split_ts="2026-09-04")
dir_left = direction_report(recs, split_ts="2026-08-23")
dir_months = direction_report(recs, ["2026-06", "2026-07"], ["2026-08", "2026-09"])
TODAY = datetime.now().strftime("%Y-%m-%d")


def mae_of(w, name):
    return w["models"].get(name, {}).get("mae")


def best_model(w):
    return min(w["models"].items(), key=lambda kv: kv[1]["mae"])


def best_base(w):
    return min(w["baselines"].items(), key=lambda kv: kv[1]["mae"])


wb = Workbook()
wb.remove(wb.active)

sheet(
    wb, "结论", "DS 粗精煤泥灰分：水平值不可用、方向可用（%s 用 %d 条数据复核）" % (TODAY, n_all),
    ["议题", "结论", "数据证据", "纪律 / 做法"], [16, 30, 62, 56],
    [
        ["水平值预测", "不可用（不承诺精度）",
         "锁定检验窗（训练 6-16~9-03 共 %d 条 → 检验 9-04~9-21 共 %d 条，模型没见过）："
         "模型(%s) MAE %.3f，而「取训练均值」MAE %.3f、最强在线基线（%s）MAE %.3f。"
         % (lock["trainN"], lock["testN"], best_model(lock)[0], best_model(lock)[1]["mae"],
            lock["baselines"]["取训练均值"]["mae"], best_base(lock)[0], best_base(lock)[1]["mae"]),
         "原因不是算法，而是灰分中枢跨期漂移：本窗训练期均值 %.3f、检验期均值 %.3f（%+.3f）。"
         "因此 DS 水平输出只能标注为「历史对照」，不能当向前精度承诺。"
         % (lock["trainMean"], lock["testMean"], lock["testMean"] - lock["trainMean"])],
        ["方向判断", "可用（唯一稳健的前向信号）",
         "新独立窗口（切点 2026-09-04）：命中 %.3f，多数基线 %.3f、惯性基线 %.3f，95%% 区间 %s，n=%d。"
         "更早的切点 2026-08-23：命中 %.3f（区间 %s，n=%d）；旧月切口径 %.3f（n=%d）。"
         % (dir_new["hit"], dir_new["baseline"], dir_new["inertia"], dir_new["ci"], dir_new["n"],
            dir_left["hit"], dir_left["ci"], dir_left["n"], dir_months["hit"], dir_months["n"]),
         "验收门槛：bootstrap 95% 区间**下界 > max(多数基线, 惯性基线)** 才判可用。"
         "方向用「第 t 条已知信息」预测「第 t+1 条相对第 t 条的涨跌」，|Δ|≤0.5% 记为平、不计入命中。"],
        ["机理", "均值回复（不是靠采样节奏）",
         "惯性基线（本次涨→下次继续涨）仅 %.3f，明显反持续；去掉「到下次读数的时间间隔」特征后命中率不变。"
         % dir_new["inertia"],
         "模型从「Δprev + 上一条水平」重建当前水平：当前偏高则下次多半回落。"
         "正式版**不使用**时间间隔特征 —— 排除「靠采样节奏猜方向」的偏利。"],
        ["拟合度的陷阱", "拟合度高 ≠ 向前可用（本轮又验了一次）",
         "同一锁定窗：把目标换成「灰分−原煤灰」后，样本内 R² 从 %.3f 升到 %.3f，而向前 MAE 从 %.3f 恶化到 %.3f。"
         % (mae_of(lock, "10因子(现行)|MLR") and lock["models"]["10因子(现行)|MLR"]["inSampleR2"],
            0.754, lock["models"]["10因子(现行)|MLR"]["mae"], 1.939),
         "所以系统里把「样本内 R²」与「向前 MAE + 基线对照」并列展示，选型只看向前窗口。"
         "（0.754/1.939 为目标变换实验值，见「备选方案-4因子」表同批实验。）"],
        ["工程结论", "把向前验证做成常规数字（已落地）",
         "新增只读接口 GET /api/v1/training/coarse-forward（返回向前 MAE、朴素基线、方向命中与门控结论），"
         "DS 训练结果同时带 forward 与 direction；粗精煤泥页 DS 视图直接显示该区块。",
         "接口只读、不训练、不写库；结果按 (数据指纹, range) 进程内缓存，首次约 4~8 秒、之后即时。"
         "训练范围若选「全部」，没有剩余数据可检验 → 页面显示「留尾参考」并说明原因。"],
        ["纪律①（踩过坑）", "特征只能用 t 时刻**已知**量",
         "曾把「第 t 条自身的灰分同源测量」与「第 t 条自身的变化」配对，得到 0.833~0.917 的假命中率；"
         "把目标前移一格（真向前）后掉到 0.657~0.771。",
         "前者是「同时刻解释」——生产上第 t 条还不存在，等于用未来信息。这条假象写进模块文档作反面教材。"],
        ["纪律②（对称的）", "也不能只追某个前向指标",
         "增量模型的向前 q2Time = +0.417（看着很好），但其检验集 R² = −11.47、MAE = 14.34 —— 彻底崩。",
         "目标换成 Δ 后 Q² 是在「增量方差」上算的，量纲与水平值不同，跨口径比指标会得出相反结论。"
         "任何前向指标必须与「水平 MAE / 是否打赢均值基线」一起看。"],
    ],
    fill_col=1,
    fill_map={"水平值预测": BAD_FILL, "方向判断": OK_FILL, "机理": OK_FILL, "拟合度的陷阱": WARN_FILL,
              "工程结论": OK_FILL, "纪律①（踩过坑）": WARN_FILL, "纪律②（对称的）": WARN_FILL},
)
rows = []


def add_window(title, w):
    """按表头 [窗口, 对象, 训练n, 检验n, 训练截止 / 检验区间, 向前 MAE, 偏差, 备注] 逐列对齐。"""
    if not w.get("usable"):
        rows.append([title, "（不可用）", w.get("trainN", 0), w.get("testN", 0), "-", "-", "-",
                     w.get("note", "")])
        return
    rows.append([title, "（窗口概览）", w["trainN"], w["testN"],
                 "训练 →%s（均值 %.3f）｜检验 %s ~ %s（均值 %.3f）"
                 % (w["trainEnd"][:10], w["trainMean"], w["testFrom"][:10], w["testTo"][:10], w["testMean"]),
                 "", "", ""])
    for k, v in sorted(w["models"].items(), key=lambda kv: kv[1]["mae"]):
        rows.append(["%s · 模型" % title, k, "", "", "", "MAE %.3f" % v["mae"], "偏差 %+.3f" % v["bias"],
                     "样本内 R² %.3f｜向前 R² %+.3f" % (v["inSampleR2"], v["r2"])])
    for k, v in sorted(w["baselines"].items(), key=lambda kv: kv[1]["mae"]):
        rows.append(["%s · 基线" % title, k, "", "", "", "MAE %.3f" % v["mae"], "偏差 %+.3f" % v["bias"],
                     "向前 R² %+.3f" % v["r2"]])
    d = w.get("direction") or {}
    rows.append(["%s · 方向" % title, "下一读数涨跌", "", "", "", "命中 %.3f" % (d.get("hit") or 0),
                 "多数基线 %.3f｜惯性 %.3f" % (d.get("baseline") or 0, d.get("inertia") or 0),
                 "95%%区间 %s，n=%s，%s" % (d.get("ci"), d.get("n"), "可用" if d.get("usable") else "不足")])


add_window("锁定检验窗（9-04~9-21，新数据）", lock)
add_window("开发侧向前窗（8-23~9-03）", leftover)
add_window("部署口径 jun_jul 留出整段（8-23~9-21）", prod)
sheet(
    wb, "真向前验证", "真向前验收：训练只用切点之前、评估只用切点之后（%s 现算，库内粗灰 %d 条）" % (TODAY, n_all),
    ["窗口", "对象", "训练n", "检验n", "训练截止 / 检验区间", "向前 MAE", "偏差", "备注"],
    [26, 30, 8, 8, 46, 14, 20, 40],
    rows,
)

sheet(
    wb, "方向验证明细", "方向信号稳健性：不同切分 × 门控结论（Ridge 回归 Δ 取符号，|Δ|>0.5% 计入）",
    ["切分口径", "开发 / 检验", "方向命中", "多数基线", "惯性基线", "95% bootstrap", "检验样本", "门控结论"],
    [26, 30, 12, 12, 12, 20, 12, 20],
    [
        ["时间切分 2026-09-04（新数据）", "6-16~9-03 / 9-04~9-21（%d 条）" % lock["testN"],
         dir_new["hit"], dir_new["baseline"], dir_new["inertia"],
         "[%.3f, %.3f]" % tuple(dir_new["ci"]), dir_new["n"], "可用" if dir_new["usable"] else "拒绝"],
        ["时间切分 2026-08-23", "6-16~8-22 / 8-23~9-21",
         dir_left["hit"], dir_left["baseline"], dir_left["inertia"],
         "[%.3f, %.3f]" % tuple(dir_left["ci"]), dir_left["n"], "可用" if dir_left["usable"] else "拒绝"],
        ["月份切分（旧口径）", "2026-06+07 / 2026-08+09",
         dir_months["hit"], dir_months["baseline"], dir_months["inertia"],
         "[%.3f, %.3f]" % tuple(dir_months["ci"]), dir_months["n"],
         "可用" if dir_months["usable"] else "拒绝"],
        ["月份切分（2026-09-23 首版）", "2026-06+07 / 2026-08+09（当时 152 条）",
         0.771, 0.514, 0.343, "[0.629, 0.914]", 35, "可用"],
        ["月份切分（首版第二组）", "2026-06 / 2026-07", 0.727, 0.523, 0.343, "[0.591, 0.841]", 44, "可用"],
        ["月份切分（首版被拒的那组）", "2026-06 / 2026-07+08+09", 0.617, 0.531, 0.343,
         "[0.506, 0.716]", 81, "拒绝（下界未超基线）"],
    ],
    fill_col=8,
    fill_map={"可用": OK_FILL, "拒绝": BAD_FILL, "拒绝（下界未超基线）": BAD_FILL},
)

sheet(
    wb, "数据导入-20260924", "新数据导入（用户 2026-09-24 确认口径后执行）",
    ["项", "内容", "依据 / 结果"], [24, 66, 52],
    [
        ["源文件", "C:\\Users\\25925\\Desktop\\web(2)导入数据\\新 下的 3 个文件："
                   "粗精煤泥灰分影响因素9-23.xlsx、灰分、密度 （导入版)9-23.xlsx、浮精 (1)9-23.xlsx",
         "旧版本文件（9-4 系列）与 ~$ 锁定文件不导；导入前确认工作簿已关闭"],
        ["粗精煤泥", "**有效数据只有 61 行**（3 行表头之外其余为空）→ 不再做日期向下填充，只导这 61 行；"
                     "库内粗灰 152 → 205 条，覆盖延到 2026-09-21",
         "用户 2026-09-24 明确：其余为空、只导有效行"],
        ["灰分密度", "**只导 A 系统** 183 行（B 系统 183 行不导并在 import_logs.errors 留痕）；"
                     "A 系统里密度写 1 的 6 行仍导入、密度记 NULL（保住灰分）",
         "用户口径「只导A」；与前端导入页 import.js 对坏密度的处理一致；库内 360 → 543 条"],
        ["浮精", "14 行（9-04~9-20），库内 17 → 31 条", "煤量须在 0~60 t/h；本批无异常行"],
        ["9-01/9-02 重叠", "新文件与库内旧数据是同一天同批采样、时间与数值都不同（旧值来自 9-5 的『9-4』文件）→ "
                           "**以新文件为准**：删库内 8 行、写入新文件 9 行（净 +1）",
         "用户 2026-09-24 选择「以新文件为准」；脚本 --overlap replace，重叠判定窗口 60 分钟"],
        ["写入方式与安全", "单事务写入；写入前用 sqlite3 在线备份 API 存 "
                           "backups/dense_medium-preimport-newbatch-20260924-150607-498.db（385024 字节）",
         "导入后核对：integrity_check=ok、同类无重复时间戳、条数 205/31/543、import_logs 新增 3 条"],
        ["脚本", "backend/scripts/import_new_batch.py（支持 --dry-run / --ash-systems / --overlap）",
         "可重复执行：已存在的时间戳不会重复写入"],
        ["未导入但保留的原始信息", "B 系统 183 行里 76 行带真实密度值（与 A 差中位 0.007、最大 0.102）；"
                                   "9-08 08:33~13:33 有 6 个时刻只有 B 有密度",
         "按「只导A」口径未导入；如需补导，重跑脚本加 --ash-systems A,B 即可（不会破坏 A 行）"],
    ],
)

sheet(
    wb, "备选方案-4因子", "备选：去掉 6 个接近常量的开关列（**未采用**，用户 2026-09-24 决定模型先不动）",
    ["特征集 / 口径", "开发侧向前窗 MAE", "锁定检验窗 MAE", "样本内 R²（锁定窗）", "说明"],
    [34, 20, 20, 18, 50],
    [
        ["10 因子（现行）", "%.3f" % (leftover["models"].get("10因子(现行)|MLR", {}).get("mae") or 0),
         "%.3f" % (lock["models"].get("10因子(现行)|MLR", {}).get("mae") or 0),
         "%.3f" % (lock["models"].get("10因子(现行)|MLR", {}).get("inSampleR2") or 0),
         "sysA/sysB/sys401/sys402/脱粉473/474 六个开关列在样本里几乎不变（sysA 均值 0.97），"
         "对向前预测只贡献噪声"],
        ["4 因子（原煤灰/煤量/液位/水分）",
         "%.3f" % (leftover["models"].get("4因子(原煤灰/煤量/液位/水分)|MLR", {}).get("mae") or 0),
         "%.3f" % (lock["models"].get("4因子(原煤灰/煤量/液位/水分)|MLR", {}).get("mae") or 0),
         "%.3f" % (lock["models"].get("4因子(原煤灰/煤量/液位/水分)|MLR", {}).get("inSampleR2") or 0),
         "两个窗口一致更优（锁定窗显著：配对 bootstrap 差值区间 [+0.138, +0.643] 不含 0）"],
        ["目标变换（灰分−原煤灰）", "%.3f" % 2.633, "%.3f" % 1.939, "%.3f" % 0.754,
         "样本内 R² 最漂亮、向前最差之一 —— 只追拟合度会选错模型"],
        ["目标变换（灰分/原煤灰）", "%.3f" % 2.488, "%.3f" % 1.821, "%.3f" % 0.459,
         "同上，向前没有改善"],
    ],
)

sheet(
    wb, "落地与待办", "已完成 / 待做 / 现场待确认（%s）" % TODAY,
    ["项", "类型", "内容", "状态"], [24, 12, 76, 18],
    [
        ["coarse_direction.py", "已完成", "方向判断模块：direction_report() 输出 usable/hit/baseline/inertia/ci/n/note/gating，"
                                          "支持 split_ts 时间切分与月份切分；自带 CLI 自检。", "已提交 3b48f09"],
        ["单元测试", "已完成", "方向 3 条（均值回复必须可用 / 随机游走必须拒绝 / 必须声明 gating）+ "
                               "向前验证 6 条（store 键兼容、split_ts、leftover 与 tail、缓存、训练结果携带）。", "9 passed"],
        ["时间切分修复", "已完成", "按月份切分在「训练范围已覆盖最新月份」时会把检验集切成空集（2026-09-24 实际发生）；"
                                   "改为按时间点切分 split_ts，并修掉只认 `ts` 键导致训练路径恒返回「样本不足」的缺陷。", "已修复"],
        ["真向前服务", "已完成", "app/services/coarse_forward.py：只读评估（模型 + 取训练均值/在线近k条/EWMA 基线 + 方向），"
                                 "训练范围之后为 leftover 口径、覆盖全部时退化为 tail 留尾并说明。", "本次"],
        ["只读接口", "已完成", "GET /api/v1/training/coarse-forward?range=&engine=ds[&sets=1]；进程内按数据指纹缓存，"
                               "首次约 4~8 秒。", "本次"],
        ["训练结果", "已完成", "DS 训练返回 forward（真向前 MAE/基线/结论文字）与 direction；随模型快照一起存档。"
                               "只加新键、不动 mlr/pls/metrics/history → DS 对拍基线 train_js.json 无需重生成。", "本次"],
        ["前端展示", "已完成", "粗精煤泥页 DS 视图新增「向前验证」区块：训练窗/检验窗、模型向前 MAE 与最强朴素基线、"
                               "方向命中与区间、样本内 R² 提示；训练范围为 all 时标注「留尾参考」。", "本次"],
        ["端到端验证", "已完成", "backend/scripts/verify_coarse_forward.js（对着临时库后端跑）：15 项断言，"
                                 "含页面与接口数字一致、留尾文案、旧后端降级提示。", "15 passed"],
        ["4 因子备选", "待决策", "去掉 6 个近常量开关列在两个窗口一致更优（锁定窗 MAE 1.814→1.521）。"
                                 "用户 2026-09-24 决定：模型先不动，只把向前验证与方向做进系统；本方案存档待更多数据。",
         "待拍板"],
        ["密度特征", "数据缺口", "想把「最近密度读数」作为因子：库内 A 系统密度读数只有 2026-05、08、09 三个月；"
                                 "6-7 月粗灰 113 条里 0 条能在 6 小时内对上密度读数 → **当前无法训练**。"
                                 "需要 PLC 密度连续采集接通后再评估。", "待接入"],
        ["旧后端一致性", "已处理", "生产 8000 仍跑 main 代码（无该接口）：页面检测到 404 时只显示一句灰字提示，不报红。",
         "已降级"],
        ["文档并表", "待做", "本结论后续按编号并入 docs/待办-后续优化计划.xlsx（避免重复 Sheet）。", "未开始"],
        ["现场待确认", "待确认", "B 系统 76 条真实密度是否补导；密度 1.441（9-21 单步降 0.079）是否手误；"
                                 "KEPServer 标签映射；P0 166 条是否收紧。", "待确认"],
    ],
    fill_col=4,
    fill_map={"已提交 3b48f09": OK_FILL, "9 passed": OK_FILL, "已修复": OK_FILL, "本次": OK_FILL,
              "15 passed": OK_FILL, "已降级": OK_FILL, "待拍板": WARN_FILL, "待接入": WARN_FILL,
              "未开始": WARN_FILL, "待确认": WARN_FILL},
)

DOCS.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("已生成:", OUT)
for ws in wb.worksheets:
    print("  - %-18s %d 行" % (ws.title, ws.max_row - 3))

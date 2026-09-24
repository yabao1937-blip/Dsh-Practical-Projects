# -*- coding: utf-8 -*-
"""生成 Excel:KEPServer 在线仪表标签按 瞬时/累计/周期 分类。

来源:C:/Users/25925/Desktop/web(2)导入数据/kepserver/在线仪表关联标签_20260923_170617.xlsx(只读)
分类规则(按说明字段 + 标签名 + 数据类型,顺序判定):
  周期   —— 窗口平均值(10/30分钟)或按班/日/月/年清零的累计量
  累计   —— 不清零的单调累计:总累计、累计灰分(AAD_T/ASH_T)、r射线累计值、cheng_* 计数器(说明缺失,按 UInt32 大数值判定)
  设定与状态 —— 非测量类:设定值/上下限/阈值/联锁/选择/输入/优化/状态/启泵/弃用/运行信号/故障
  瞬时   —— 其余:实时物理量测量(瞬时带煤量/流量、在线灰分、当前密度/液位、浓度等,含人工录入的当前化验灰分)

Sheet 结构:总览(规则+计数) / 瞬时 / 累计 / 周期 / 设定与状态(非测量)。
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = r"C:/Users/25925/Desktop/web(2)导入数据/kepserver/在线仪表关联标签_20260923_170617.xlsx"
OUT = Path(__file__).resolve().parent.parent.parent / "docs" / "KEPServer标签分类-瞬时累计周期.xlsx"

# ---------------------------------------------------------------- 读源
src = load_workbook(SRC, read_only=True)
ws = src["标签清单"]
rows = list(ws.iter_rows(values_only=True))
HDR = [str(h) for h in rows[0]]
data = [dict(zip(HDR, r)) for r in rows[1:] if r[0] is not None]

# ---------------------------------------------------------------- 分类
KEY_SET = ["设定", "上限", "下限", "阈值", "联锁", "选择", "输入", "优化", "状态", "启泵", "弃用",
           "运行信号", "故障", "范围"]

def classify(d):
    """返回 (大类, 依据)。判定顺序:cheng 计数器 → 周期 → 累计 → 布尔/整数状态 → 设定关键词 → 瞬时。"""
    tag = str(d.get("标签名") or "")
    desc = str(d.get("说明") or "")
    dtype = str(d.get("数据类型") or "")
    if tag.startswith("cheng_"):
        return "累计", "cheng_* 计数器(说明缺失,UInt32 大数值,按累计判定)"
    if any(k in desc for k in ("10分钟平均", "30分钟平均")):
        return "周期", "窗口平均值"
    if any(k in desc for k in ("班累计", "日累计", "月累计", "年累计")):
        return "周期", "按周期清零的累计量(" + desc[:2] + ")"
    if "累计" in desc:
        return "累计", "说明含「累计」(不清零)"
    if dtype == "Boolean":
        return "设定与状态", "开关量(Boolean)"
    if dtype in ("Int16", "Int32", "UInt32"):
        return "设定与状态", "整数型(状态/设定)"
    if any(k in desc for k in KEY_SET):
        return "设定与状态", "说明含设定/状态类关键词"
    return "瞬时", "实时物理量测量"

for d in data:
    d["_类"], d["_据"] = classify(d)

buckets = {}
for d in data:
    buckets.setdefault(d["_类"], []).append(d)

# 周期粒度(周期 sheet 专用):从说明提取
def period_grain(desc):
    for g in ("10分钟", "30分钟", "班", "日", "月", "年"):
        if g in str(desc):
            return g
    return ""

# ---------------------------------------------------------------- 样式(与仓库其他报表一致)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL = PatternFill("solid", fgColor="F2F6FA")

def write_sheet(wb, title, header, records, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for i, rec in enumerate(records):
        ws.append(rec)
        r = i + 2
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(widths[c-1] >= 30))
            if i % 2 == 1:
                cell.fill = ALT_FILL
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(records) + 1}"
    return ws

wb = Workbook()
wb.remove(wb.active)

# ---- Sheet 1: 总览
order = ["瞬时", "累计", "周期", "设定与状态"]
ov = wb.create_sheet("总览")
ov_rows = [
    ["KEPServer 在线仪表标签分类 —— 瞬时 / 累计 / 周期", ""],
    ["来源文件", Path(SRC).name],
    ["标签总数", len(data)],
    ["生成时间", __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ["", ""],
    ["分类", "条数", "口径"],
    ["瞬时", len(buckets.get("瞬时", [])), "实时物理量测量:瞬时带煤量/流量、在线灰分(X光/r射线瞬时、在线测灰)、当前化验灰分、当前密度、煤泥/磁性物密度、液位、浓度等 —— 采集按轮询周期读当前值"],
    ["累计", len(buckets.get("累计", [])), "不清零的单调累计:总累计带煤量/流量、累计灰分(AAD_T/ASH_T)、r射线累计灰分、cheng_123/186/210/501/579 计数器(说明缺失,按 UInt32 大数值判定) —— 采集要记基线做差或直接入库累加"],
    ["周期", len(buckets.get("周期", [])), "按窗口或周期组织的值:X光 10/30 分钟平均值;班/日/月/年累计带煤量与流量(每到周期边界清零重计) —— 采集按对齐周期边界读数"],
    ["设定与状态(非测量)", len(buckets.get("设定与状态", [])), "不属以上三类的:设定值/上下限/阈值/联锁/选择/输入/优化/运行信号/故障/启泵/弃用等开关量与状态量。单独列出以免污染测量类;接入时按需订阅变化(状态)或作为参数快照(设定)"],
    ["", ""],
    ["判定顺序", "cheng_* 计数器 → 周期(窗口平均/按周期清零) → 累计(说明含「累计」) → Boolean/整数 → 设定/状态关键词 → 其余归瞬时"],
    ["备注", "「化验灰分」(AAD)为人工录入的当前有效化验值,按瞬时点值归类;「输入灰分」(A_IN)为控制器输入,归设定与状态"],
]
for r in ov_rows:
    ov.append(r)
ov.column_dimensions["A"].width = 22
ov.column_dimensions["B"].width = 12
ov.column_dimensions["C"].width = 110
for c in ("A", "B", "C"):
    ov[f"{c}1"].font = Font(name="微软雅黑", size=13, bold=True, color="1F4E79")
    ov[f"{c}6"].fill, ov[f"{c}6"].font = HDR_FILL, HDR_FONT
for r in range(2, len(ov_rows) + 1):
    for c in ("A", "B", "C"):
        cell = ov[f"{c}{r}"]
        if cell.font is None or cell.font.name != "微软雅黑":
            cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
for r in range(7, 11):
    ov.row_dimensions[r].height = 46
ov.row_dimensions[12].height = 30

# ---- 数据 sheets(保留源表全部 15 列,末尾加 归类依据;周期再加 周期粒度)
COLS = HDR + ["归类依据"]
W = [5, 10, 10, 14, 26, 20, 40, 40, 9, 9, 7, 19, 10, 10, 34, 34]

def rec(d, extra=None):
    return [d.get(h) for h in HDR] + ([d["_据"]] + (extra or []))

write_sheet(wb, "瞬时", COLS, [rec(d) for d in buckets.get("瞬时", [])], W)
write_sheet(wb, "累计", COLS, [rec(d) for d in buckets.get("累计", [])], W)
write_sheet(wb, "周期", COLS + ["周期粒度"],
            [rec(d, [period_grain(d["说明"])]) for d in buckets.get("周期", [])], W + [10])
write_sheet(wb, "设定与状态(非测量)", COLS, [rec(d) for d in buckets.get("设定与状态", [])], W)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)

# ---------------------------------------------------------------- 校验输出
print(f"输出: {OUT}")
print(f"总数 {len(data)} = " + " + ".join(f"{k} {len(buckets.get(k, []))}" for k in order))
assert sum(len(v) for v in buckets.values()) == len(data), "分类覆盖数不等于总数"
# 抽查:各类不应混入明显异类
for d in buckets["瞬时"]:
    assert d["数据类型"] == "Float", f"瞬时混入非Float: {d['标签名']} {d['数据类型']}"
    assert not any(k in str(d["说明"]) for k in ("累计", "平均", "设定", "阈值")), f"瞬时混入: {d['标签名']}"
for d in buckets["累计"]:
    assert "累计" in str(d["说明"]) or str(d["标签名"]).startswith("cheng_"), f"累计无据: {d['标签名']}"
for d in buckets["周期"]:
    assert period_grain(d["说明"]) != "", f"周期无粒度: {d['标签名']} {d['说明']}"
print("断言全部通过")
print("\n周期粒度分布:", {g: sum(1 for d in buckets['周期'] if period_grain(d['说明']) == g)
                              for g in ('10分钟', '30分钟', '班', '日', '月', '年')})

# -*- coding: utf-8 -*-
"""导入桌面新数据（粗精煤泥 9-23 / 灰分密度导入版 9-23 / 浮精 9-23）。

用户已确认的三处口径（2026-09-24）：
  · 粗精煤泥表**有效数据只有 61 行**，其余 3 行表头之外没有别的数据 → 不做日期向下填充，
    只导这 61 行（原先怀疑的"按天合并单元格"不成立）。
  · 灰分密度表**只导 A 系统**（`--ash-systems A`）：B 系统行一律不导，但写入 import_logs.errors 留痕。
    A 系统里密度缺失/越界（源文件写成 1）的行**仍然导入**，密度记 NULL（保住灰分，与前端导入页
    frontend/js/import.js 的处理一致），不整行丢弃。
  · 9-01/9-02 与库内旧数据重叠（同一天同批采样、时间与数值都不同，旧值来自 9-5 的
    「粗精煤泥灰分影响因素 9-4.xlsx」）→ **以新文件为准**：删掉库内那 8 行，写入新文件 9 行
    （`--overlap replace`，默认）。判定窗口见 OVERLAP_WINDOW。

安全做法：先校验并打印计划 → 单事务写入 → 导入后核对条数与时间范围。
用法：
  python backend/scripts/import_new_batch.py --dry-run          # 只看计划，不写库
  python backend/scripts/import_new_batch.py                    # 写库（只导 A 系统；重叠以新为准）
  python backend/scripts/import_new_batch.py --ash-systems A,B  # 若以后要补导 B 系统
  python backend/scripts/import_new_batch.py --overlap skip     # 若改成"重叠时保留库内旧值"
  # 固定供给（用户 2026-09-24 Q7=A）：把文件丢进投放目录 → 一条命令搞定
  python backend/scripts/import_new_batch.py --scan-dir "C:\\Users\\25925\\Desktop\\web(2)导入数据\\自动导入"
  python backend/scripts/import_new_batch.py --scan-dir <目录> --apply   # 真正写库并把源文件移入 已导入\\
固定供给约定：
  · 目录里放 xlsx，文件名含「粗精煤泥」= 粗灰表、含「灰分」+「密度」= 灰分密度表、含「浮精」= 浮精表；
  · 每类取**修改时间最新**的一个文件（旧文件不会盖新文件）；~$ 临时文件自动忽略；
  · 不带 --apply 时只看计划；带 --apply 写库成功后把用过的文件移到 已导入\\<时间戳>\\；
  · --db 可覆盖库路径（验证时指向临时库，避免碰生产）。
"""
import argparse
import json
import sqlite3
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "dense_medium.db"
SRC = Path(r"C:\Users\25925\Desktop\web(2)导入数据\新")
FILES = {
    "coarse": "粗精煤泥灰分影响因素9-23.xlsx",
    "ash_density": "灰分、密度 （导入版)9-23.xlsx",
    "float": "浮精 (1)9-23.xlsx",
}
# ---- 固定供给（用户 2026-09-24 选择 Q7=A）：投放目录 + 文件名识别 + 导入后归档 ----
# 约定：把三个文件放进投放目录（固定列序，见 docs/DS粗灰预测-疑问与决策.xlsx），
# 文件名只要含下列关键字就能被识别；每类取**修改时间最新**的一个；导入成功后移到 已导入\<时间戳>\。
DROP_DIR = Path(r"C:\Users\25925\Desktop\web(2)导入数据\自动导入")
KEYWORDS = (("coarse", ("粗精煤泥",)), ("ash_density", ("灰分", "密度")), ("float", ("浮精",)))
ARCHIVE_DIRNAME = "已导入"
# 与库内既有 import_logs 的 category 取值保持一致（见 sqlite: select distinct category from import_logs）
CAT_NAME = {"coarse": "coarse_factors", "ash_density": "ash_density", "float": "float_ash"}
YEAR = "2026"
# 「同一次采样、时间戳略有出入」的判定窗口（分钟）。粗/浮精采样间隔以小时计 → 60 分钟安全；
# 灰分密度表相邻只有 5 分钟，一律按精确时间戳判重（窗口 0）。
OVERLAP_WINDOW = {"coarse": 60, "float": 60, "ash_density": 0}


def classify(name):
    """文件名 → 类别（None=不认识）。粗精煤泥必须先判，否则会被『灰分』误抓。"""
    if name.startswith("~$") or name.startswith("."):
        return None
    for cat, keys in KEYWORDS:
        if all(k in name for k in keys):
            return cat
    return None


def resolve_sources(scan_dir, min_bytes=4096):
    """扫描投放目录 → {类别: 文件路径}（每类取 mtime 最新；已导入/临时文件跳过）。

    为什么要 min_bytes：2026-09-24 实测，一个 0 字节的同名垃圾文件会因为 mtime 更新而"赢得"
    选取，随后 load_workbook 抛异常把整轮导入打断。真实表格 ≥ 十几 KB，按体积先挡一道。
    """
    scan_dir = Path(scan_dir)
    if not scan_dir.is_dir():
        raise SystemExit("投放目录不存在：%s" % scan_dir)
    best, tiny = {}, []
    for p in sorted(scan_dir.iterdir()):
        if not p.is_file() or p.suffix.lower() not in (".xlsx", ".xls", ".xlsm"):
            continue
        cat = classify(p.name)
        if cat is None:
            continue
        if p.stat().st_size < min_bytes:
            tiny.append("%s（%.1f KB，太小，按临时/损坏文件跳过）" % (p.name, p.stat().st_size / 1024))
            continue
        if cat not in best or p.stat().st_mtime > best[cat].stat().st_mtime:
            best[cat] = p
    if tiny:
        print("   已忽略 %d 个体积异常的文件：" % len(tiny))
        for t in tiny:
            print("      · %s" % t)
    return best


def archive_sources(sources):
    """导入成功后把用过的文件移进 已导入\\<时间戳>\\，避免下一轮重复导入。"""
    if not sources:
        return None
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest_dir = Path(next(iter(sources.values()))).parent / ARCHIVE_DIRNAME / stamp
    dest_dir.mkdir(parents=True, exist_ok=True)
    moved = []
    for p in sources.values():
        target = dest_dir / Path(p).name
        Path(p).replace(target)
        moved.append(str(target))
    return moved


def norm_ts(v):
    """采样时间 → 'YYYY-MM-DD HH:MM:SS'；兼容 datetime 单元格与多种字符串写法。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    s = str(v or "").strip()
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M"):
        try:
            return datetime.strptime(s, f).strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            continue
    return None


def onoff(v):
    """停/开（含 0/1、True/False）→ (0/1, 说明)；说明 ∈ {'ok','空','无法识别'}。

    注意：不能用 `str(v or "")` —— Excel 里读到的**整数 0** 会被 `0 or ""` 变成空串，
    于是"停(0)"被误报成"无法解析"（2026-09-24 实测：61 行里 37 行假报）。
    """
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return 0, "空"
    if isinstance(v, bool):
        return (1 if v else 0), "ok"
    if isinstance(v, (int, float)):
        if v == 0:
            return 0, "ok"
        if v == 1:
            return 1, "ok"
        return 0, "无法识别"
    s = str(v).strip()
    if s in ("停", "0", "0.0", "false", "False", "关"):
        return 0, "ok"
    if s in ("开", "1", "1.0", "true", "True"):
        return 1, "ok"
    return 0, "无法识别"


def num(v, lo=None, hi=None):
    try:
        f = float(v)
    except Exception:
        return None
    if f != f or f in (float("inf"), float("-inf")):
        return None
    if lo is not None and f < lo:
        return None
    if hi is not None and f > hi:
        return None
    return f


def read_rows(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            out.append((ws.title, rows))
    wb.close()
    return out


def build(cat, ash_systems, path=None):
    """→ (可入库记录, 异常/口径跳过明细, 统计计数, 文件名)"""
    path = Path(path) if path else (SRC / FILES[cat])
    built, skipped = [], []
    st = Counter()
    for sheet, rows in read_rows(path):
        start = 1
        if cat == "coarse":                       # 前 3 行是双层表头
            for i, r in enumerate(rows[:4]):
                if r and str(r[0]).strip() not in ("粗精煤泥的灰分影响因素", "日期", "", "None"):
                    start = i
                    break
        for i, r in enumerate(rows[start:], start=start + 1):
            if not r or all(v is None or str(v).strip() == "" for v in r):
                st["全空行"] += 1
                continue
            if cat == "coarse":
                d, t = str(r[0] or "").strip(), str(r[2] or "").strip()
                if "." not in d or ":" not in t:
                    st["表头/无日期"] += 1
                    skipped.append(("表头/无日期", "第%d行 %s" % (i, str(r[:4])[:60])))
                    continue
                m, day = d.split(".")[:2]
                try:
                    ts = datetime.strptime("%s-%02d-%02d %s" % (YEAR, int(m), int(day), t), "%Y-%m-%d %H:%M:%S")
                except Exception:
                    st["日期非法"] += 1
                    skipped.append(("日期非法", "第%d行 %s %s" % (i, d, t)))
                    continue
                ash = num(r[12], 1, 60)            # M 315灰分 = 粗精煤泥灰分（模型目标）
                coal = num(r[5], 0, 3000)          # F 小时带煤量
                if ash is None or coal is None:
                    st["缺灰分/煤量"] += 1
                    skipped.append(("缺灰分/煤量", "第%d行 %s ash=%s coal=%s" % (i, d, r[12], r[5])))
                    continue
                flags = {}
                empty_flag, bad_flag = [], []
                for key, col in (("sysA", 6), ("sysB", 7), ("sys401", 8), ("sys402", 9),
                                 ("desliming473", 10), ("desliming474", 11)):
                    v, why = onoff(r[col])
                    if why == "空":
                        empty_flag.append(key)
                    elif why == "无法识别":
                        bad_flag.append("%s=%r" % (key, r[col]))
                    flags[key] = v
                if empty_flag:                     # 源表该格为空 → 按"停(0)"，只统计不报异常
                    st["开关列空格→按停(0)"] += 1
                if bad_flag:                       # 不静默：真正读不懂的值记进统计与日志明细
                    st["开关列无法解析→按停(0)"] += 1
                    skipped.append(("开关列无法解析→按0", "第%d行 %s" % (i, "、".join(bad_flag))))
                built.append(dict(
                    category="coarse", ts=ts.strftime("%Y-%m-%d %H:%M:%S"), system="合并", source="import",
                    ash_content=ash, coal_amount=coal, raw_ash=num(r[4], 0, 100), moisture=num(r[13], 0, 100),
                    level=num(r[14], 0, 100), mining_face=str(r[3] or "").replace("\n", " ").strip()[:40] or None,
                    is_stoppage=1 if coal <= 10 else 0, **flags))
            elif cat == "ash_density":
                ts = norm_ts(r[0])
                if not ts:
                    st["时间非法"] += 1
                    skipped.append(("时间非法", "第%d行 %s" % (i, r[0])))
                    continue
                sysname = str(r[1] or "").strip()
                if sysname not in ash_systems:      # 口径：只导 A 系统
                    st["非%s系统(口径不导)" % "/".join(sorted(ash_systems))] += 1
                    skipped.append(("口径不导：系统%s" % (sysname or "空"),
                                    "%s 皮带%s 灰分%s 密度%s" % (ts, r[2], r[3], r[4])))
                    continue
                den = num(r[4], 1.3, 1.6)           # E 密度值：源文件里 1 是占位，按缺失处理
                if den is None:
                    st["A系统密度缺失→记NULL"] += 1
                    skipped.append(("A系统密度缺失→导入并记 NULL",
                                    "%s 皮带%s 灰分%s 密度=%s" % (ts, r[2], r[3], r[4])))
                built.append(dict(category="ash_density", ts=ts, system=sysname,
                                  belt=str(r[2] or "").strip() or None, source="import",
                                  ash_content=num(r[3], 0, 100), density=den))
            else:
                ts = norm_ts(r[0])
                if not ts:
                    st["时间非法"] += 1
                    skipped.append(("时间非法", "第%d行 %s" % (i, r[0])))
                    continue
                amt = num(r[2], 0, 60)              # C 浮精煤量 t/h
                if amt is None:
                    st["煤量越界/坏值"] += 1
                    skipped.append(("煤量越界/坏值", "第%d行 %s 煤量=%s 灰分=%s 压滤=%s"
                                    % (i, ts, r[2], r[1], r[3])))
                    continue
                built.append(dict(category="float", ts=ts, system="合并", source="import",
                                  ash_content=num(r[1], 0, 100), coal_amount=amt,
                                  filter_press_running=(str(r[3] or "").strip() == "运行")))
    # 文件内按 (时间, 系统) 去重（保留第一条）
    seen, uniq = set(), []
    for b in built:
        k = (b["ts"], b["system"])
        if k in seen:
            st["文件内重复(同刻同系统)"] += 1
            continue
        seen.add(k)
        uniq.append(b)
    return uniq, skipped, st, path.name


def _sec(ts):
    return datetime.strptime(ts, "%Y-%m-%d %H:%M:%S").timestamp()


def match_overlap(cat, new_rows, db_ts, window_min):
    """识别「同一次采样、时间戳略有出入」的重叠（贪心一对一，先配最近的一对）。

    → (pairs, matched_new_idx)：pairs = [(新行ts, 库内ts, 差分钟)]，matched_new_idx = 命中的新行下标集合。
    窗口 ≤0（灰分密度表：相邻仅 5 分钟）时不做近似匹配，只靠"精确时间戳已存在"判重。
    """
    if window_min <= 0 or not new_rows or not db_ts:
        return [], set()
    cands = []
    for i, b in enumerate(new_rows):
        t = _sec(b["ts"])
        for d in db_ts:
            delta = abs(_sec(d) - t) / 60.0
            if delta <= window_min:
                cands.append((delta, i, d))
    cands.sort()
    used_new, used_old, pairs = set(), set(), []
    for delta, i, d in cands:
        if i in used_new or d in used_old:
            continue
        used_new.add(i)
        used_old.add(d)
        pairs.append((new_rows[i]["ts"], d, round(delta, 1)))
    return sorted(pairs), used_new


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--ash-systems", default="A", help="灰分密度表导入哪些系统（默认只导 A）")
    ap.add_argument("--overlap", default="replace", choices=("replace", "skip", "both"),
                    help="与库内旧行近似重叠时：replace=以新文件为准换掉旧行（默认）/ skip=保留旧行 / both=两套都留")
    ap.add_argument("--scan-dir", default=None,
                    help="固定供给模式：扫描该投放目录（每类取 mtime 最新的文件）。"
                         "默认只报告，必须加 --apply 才会写库并归档")
    ap.add_argument("--apply", action="store_true", help="配合 --scan-dir：确认写库并把源文件移入 已导入\\")
    ap.add_argument("--db", default=None, help="覆盖数据库路径（验证用临时库时使用）")
    a = ap.parse_args()
    ash_systems = {s.strip().upper() for s in a.ash_systems.split(",") if s.strip()}
    scan_mode = bool(a.scan_dir)
    if scan_mode and not a.apply:
        a.dry_run = True
    sources = {}
    if scan_mode:
        sources = resolve_sources(a.scan_dir)
        print("扫描投放目录：%s" % a.scan_dir)
        for cat in ("coarse", "ash_density", "float"):
            p = sources.get(cat)
            print("   %-12s %s" % (cat, p.name + "（%.1f KB，%s）" % (p.stat().st_size / 1024,
                  datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")) if p else "未找到文件"))
        missing = [c for c in ("coarse", "ash_density", "float") if c not in sources]
        if missing:
            print("   注意：缺 %s 的文件，本轮只导找到的那些" % "、".join(missing))

    con = sqlite3.connect(str(a.db or DB), timeout=30)
    cur = con.cursor()
    before = dict(cur.execute("select category, count(*) from coal_records group by category").fetchall())

    plans, failures = {}, {}
    for cat in FILES:
        if scan_mode and cat not in sources:
            continue
        try:
            built, skipped, st, fname = build(cat, ash_systems, sources.get(cat))
        except Exception as exc:              # 单个文件坏掉不该拖垮另外两类
            failures[cat] = "%s: %s" % (sources.get(cat).name if sources.get(cat) else FILES[cat], exc)
            print("[%s] 解析失败，本轮跳过这一类：%s" % (cat, exc))
            continue
        db_ts = [r[0] for r in cur.execute(
            "select ts from coal_records where category=? order by ts", (cat,))]
        exist = set(db_ts)
        new = [b for b in built if b["ts"] not in exist]
        dup = len(built) - len(new)
        pairs, matched = match_overlap(cat, new, db_ts, OVERLAP_WINDOW.get(cat, 0))
        drop_rows = []
        if pairs and a.overlap == "replace":
            drop_rows = sorted({d for _, d, _ in pairs})
        elif pairs and a.overlap == "skip":
            new = [b for i, b in enumerate(new) if i not in matched]
        plans[cat] = dict(new=new, skipped=skipped, st=st, fname=fname, total=len(built),
                          dup=dup, pairs=pairs, drop=drop_rows)
        print("[%s] %s" % (cat, fname))
        print("    文件内可入库 %d 条；库内已存在同时间戳 %d 条；近似重叠 %d 对（--overlap %s）"
              % (len(built), dup, len(pairs), a.overlap))
        if a.overlap == "skip":
            print("    → 实际写入 %d 条；保留库内旧行，不删" % len(new))
        elif a.overlap == "replace":
            print("    → 实际写入 %d 条；同时删除库内被替换的旧行 %d 条" % (len(new), len(drop_rows)))
        else:
            print("    → 实际写入 %d 条；库内旧行保留（两套并存）" % len(new))
        print("    同行统计：" + ("、".join("%s %d" % kv for kv in st.items()) if st else "（无）"))
        if new:
            print("    时间范围：%s → %s" % (new[0]["ts"], new[-1]["ts"]))
        for nt, ot, dm in pairs:
            print("      重叠：新 %s ↔ 库内 %s（差 %s 分钟）" % (nt, ot, dm))
        if skipped:
            kinds = Counter(w for w, _ in skipped)
            print("    明细分类：" + "、".join("%s %d" % kv for kv in kinds.most_common()))
            for why, sample in skipped[:3]:
                print("      · %s | %s" % (why, sample[:78]))

    if a.dry_run:
        print("\n--dry-run：未写入" + ("（--scan-dir 模式需加 --apply 才会写库并归档）" if scan_mode else ""))
        if failures:
            print("解析失败的类别（源文件保持不动，可修正后重跑）：%s" % "、".join(failures))
        con.close()
        return
    if not plans or all(not p["new"] and not p["drop"] for p in plans.values()):
        print("\n没有需要写入的行（都已在库内）")
        con.close()
        if scan_mode:
            _archive_ok(sources, plans)
        return

    cols = ["category", "ts", "system", "source", "ash_content", "coal_amount", "density", "level",
            "raw_ash", "moisture", "belt", "sysA", "sysB", "sys401", "sys402", "desliming473",
            "desliming474", "is_stoppage", "mining_face", "filter_press_running"]
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ins = deleted = 0
    try:
        cur.execute("BEGIN IMMEDIATE")
        for cat, p in plans.items():
            if p["drop"]:
                cur.executemany("delete from coal_records where category=? and ts=?",
                                [(cat, t) for t in p["drop"]])
                deleted += len(p["drop"])
            for b in p["new"]:
                for k in ("sysA", "sysB", "sys401", "sys402", "desliming473", "desliming474", "is_stoppage"):
                    b.setdefault(k, 0)
                vals = [b.get(c) for c in cols] + [now]
                cur.execute("insert into coal_records (%s, created_at) values (%s)"
                            % (",".join(cols), ",".join("?" * (len(cols) + 1))), vals)
                ins += 1
            errs = [{"why": w, "sample": s} for w, s in p["skipped"][:40]]
            if p["drop"]:
                errs.insert(0, {"why": "重叠替换（以新文件为准）",
                                "sample": "删除库内 %d 行：%s" % (len(p["drop"]), "、".join(p["drop"]))})
            elif p["pairs"]:
                errs.insert(0, {"why": "重叠保留（--overlap %s）" % a.overlap,
                                "sample": "；".join("新 %s ↔ 旧 %s" % (n, o) for n, o, _ in p["pairs"])})
            cur.execute("insert into import_logs (ts, category, filename, total, success, failed, skipped,"
                        " status, errors, created_at) values (?,?,?,?,?,?,?,?,?,?)",
                        (now, CAT_NAME[cat], p["fname"], p["total"], len(p["new"]), 0, len(p["skipped"]),
                         "成功", json.dumps(errs, ensure_ascii=False), now))
        for cat, msg in failures.items():     # 解析失败也要留痕（类别/文件名/原因）
            cur.execute("insert into import_logs (ts, category, filename, total, success, failed, skipped,"
                        " status, errors, created_at) values (?,?,?,?,?,?,?,?,?,?)",
                        (now, CAT_NAME[cat], msg.split(": ")[0], 0, 0, 1, 0, "失败",
                         json.dumps([{"why": "解析失败", "sample": msg}], ensure_ascii=False), now))
        con.commit()
    except Exception as exc:
        con.rollback()
        print("\n!! 写入失败已回滚：", exc)
        con.close()
        raise SystemExit(1)

    print("\n写入完成：新增 %d 条、替换删除 %d 条" % (ins, deleted))
    after = dict(cur.execute("select category, count(*) from coal_records group by category").fetchall())
    print("条数变化：", {k: "%s→%s" % (before.get(k, 0), after.get(k, 0)) for k in sorted(set(before) | set(after))})
    print("各表时间范围：", cur.execute(
        "select category, count(*), min(ts), max(ts) from coal_records group by category").fetchall())
    print("本次新增的边界（按类别）：", cur.execute(
        "select category, count(*), min(ts), max(ts) from coal_records where created_at=? group by category",
        (now,)).fetchall())
    print("import_logs 总数：", cur.execute("select count(*) from import_logs").fetchone()[0],
          " 最新 3 条：", cur.execute(
              "select id, ts, category, total, success, skipped from import_logs order by id desc limit 3").fetchall())
    con.close()
    if scan_mode and sources:
        _archive_ok(sources, plans)
    if failures:
        print("解析失败的类别（源文件保持不动，可修正后重跑）：%s" % "、".join(failures))
    if scan_mode:
        print("提示：导入后可在粗精煤泥页看「向前验证」区块，或跑 "
              "python backend/scripts/evaluate_ds_forward.py 复算向前成绩。")


def _archive_ok(sources, plans):
    """只归档**解析成功**的源文件：解析失败的那份留在投放目录里等修正后重跑。"""
    ok = {cat: p for cat, p in sources.items() if cat in plans}
    if not ok:
        return
    moved = archive_sources(ok)
    print("源文件已归档（下轮不会重复导入）：")
    for m in moved or []:
        print("   ", m)


if __name__ == "__main__":
    main()

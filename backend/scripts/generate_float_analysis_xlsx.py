# -*- coding: utf-8 -*-
"""生成 Excel 文档：浮精影响分析-方法说明"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\浮精影响分析-方法说明.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D6E4F0")
SEC_FONT = Font(name="微软雅黑", size=10, bold=True)
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


wb = Workbook()

# ================= Sheet1 分析思路 =================
ws = wb.active
ws.title = "分析思路"
ws.append(["步骤", "内容", "说明"])
rows = [
    [1, "核心思想：控制变量对比",
     "把总精煤灰分拆成两种情况对比：不含浮精的“基准总灰分” vs 含浮精的“总灰分”，两者之差就是浮精对总灰分的贡献（影响值）。"],
    [2, "基准总灰分（不含浮精）",
     "基准 = 重介灰分（动态取灰分密度日志里该时刻最近的灰分值），重介量固定 250 t/h；去掉粗精项，只保留重介作为基准。"],
    [3, "含浮精总灰分",
     "总灰分 = (重介灰分×重介量 + 浮精灰分×浮精量) / (重介量 + 浮精量)；浮精灰分/煤量取表2最新一条（或量数据面板录入值）。"],
    [4, "影响值 = 含浮精 − 基准",
     "固定影响值：用最新一条浮精记录算一个当前影响值；动态波动范围：对全部历史记录逐条算影响值，再求均值与标准差 ±σ。"],
    [5, "展示",
     "4 张卡片（当前浮精灰分/浮精煤量/固定影响值/动态波动范围）+ 联动趋势图 + 影响值分布柱状图 + 明细表 + 详情弹窗。"],
]
for r in rows:
    ws.append(r)
style_sheet(ws, [6, 26, 80])
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

# ================= Sheet2 卡片与公式 =================
ws2 = wb.create_sheet("卡片与公式")
ws2.append(["卡片", "计算方式", "公式/取值", "备注"])
rows2 = [
    ["当前浮精灰分",
     "取时间上最新一条浮精记录的 ash_content 直接读值",
     "浮精灰分 = 最新浮精化验值",
     "数据源：表2 浮精(1).xlsx 或手工补录"],
    ["浮精煤量",
     "量数据统一数据源：录入 或 自动取表2最新 coal_amount",
     "resolveAmount('floatAmount')",
     "与总览“量数据”面板双向同步，可输入"],
    ["固定影响值",
     "含浮精总灰分 − 基准总灰分（最新一条）",
     "影响值 = (重介灰分×250 + 浮精灰分×浮精量)/(250+浮精量) − 重介灰分",
     "重介灰分 = getAshByTime(时间戳) 从表3动态取；重介量固定250"],
    ["动态波动范围",
     "全部记录逐条算影响值 → 均值 x̄ → 标准差 σ",
     "σ = sqrt( Σ(xi − x̄)² / n )，显示 ±σ",
     "表示浮精影响的波动程度，越小越稳定"],
]
for r in rows2:
    ws2.append(r)
style_sheet(ws2, [14, 30, 52, 30])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

# ================= Sheet3 图表与明细 =================
ws3 = wb.create_sheet("图表与明细")
ws3.append(["模块", "内容", "数据口径"])
rows3 = [
    ["联动趋势图",
     "3 条曲线：浮精灰分(%)、总精煤灰分(%)、浮精影响值(%)(右轴)",
     "逐条记录：hAsh=getAshByTime(时间戳)；总灰分=calcTotalAsh(hAsh,250,浮灰,浮量,0,0)；影响值=总灰分−hAsh"],
    ["影响值分布柱状图",
     "最近 12 个时段的影响值柱状图，颜色分级",
     ">0.15 红、0.05~0.15 橙、<0.05 绿"],
    ["明细表",
     "最近 30 条：采样时间/灰分/煤量/压滤机运行/影响值/标注",
     "influence_value 在导入时计算：influence = (灰分−8.5)×煤量/(250+煤量+15)"],
    ["详情弹窗",
     "灰分卡：直接读值说明；煤量卡：数据来源(录入/皮带秤)；影响值卡：完整代入公式过程；范围卡：σ 逐步计算",
     "与卡片同口径"],
    ["时间范围查询",
     "页面上有起止时间输入框与查询按钮",
     "注意：当前 query() 只做整体刷新，未按时间范围过滤数据（UI 已留，逻辑未接）"],
    ["异常工况标注",
     "对记录标注异常类型(压滤机卸料/浮选药剂异常/入浮浓度波动/设备故障/其他)",
     "写入 record.annotation，明细表显示标签"],
]
for r in rows3:
    ws3.append(r)
style_sheet(ws3, [16, 44, 56])
for row in ws3.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER

# ================= Sheet4 数据来源与局限 =================
ws4 = wb.create_sheet("数据来源与局限")
ws4.append(["序号", "类别", "内容"])
rows4 = [
    [1, "数据来源",
     "浮精数据 = 表2《浮精 (1).xlsx》（采样时间/浮精灰分%/浮精煤量t/h/压滤机运行，仅5条）+ 手工补录(float_ash)。"],
    [2, "重介灰分动态获取",
     "getAshByTime(时间戳) 在 calcLogs(表3灰分密度) 中找时间最接近的灰分值，找不到回退 8.50%；可传 system 参数但浮精页当前未传（全局取）。"],
    [3, "局限1：重介量固定 250",
     "影响值与趋势图中的重介精煤量是固定 250 t/h，未与量数据面板的重介精煤量(当前383)同步——口径可进一步统一。"],
    [4, "局限2：去掉粗精项",
     "公式只含重介+浮精，不包含粗精煤泥项（页面注明“去掉粗精项”）；严格来说这是“浮精对(重介+浮精)总灰分的影响”，不是对总精煤灰分的影响。"],
    [5, "局限3：样本少",
     "表2 只有 5 条浮精记录，动态波动范围(σ)的统计意义有限；随真实数据补录会改善。"],
    [6, "局限4：时间范围查询未生效",
     "起止时间输入框与查询按钮存在，但 query() 未按时间过滤（预留待接）。"],
    [7, "与总览的关系",
     "总览卡片实际灰分用的是粗精煤泥前馈预测灰分+三产品加权；浮精页影响值用动态重介灰分+固定重介量，两处口径独立。"],
]
for r in rows4:
    ws4.append(r)
style_sheet(ws4, [6, 20, 100])
for row in ws4.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP if cell.column != 1 else CENTER
        if cell.column == 2 and str(cell.value or "").startswith("局限"):
            cell.fill = WARN_FILL
ws4.auto_filter.ref = f"A1:C{ws4.max_row}"

wb.save(OUT)
print("[gen]", OUT)

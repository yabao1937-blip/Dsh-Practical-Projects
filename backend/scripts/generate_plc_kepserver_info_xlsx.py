# -*- coding: utf-8 -*-
"""生成 Excel：PLC 接入（KEPServer）信息需求清单"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\PLC接入-KEPServer信息需求清单.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
FILL_HEAD = PatternFill("solid", fgColor="D6E4F0")
FILL_BOLD = Font(name="微软雅黑", size=10, bold=True)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def make_table(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER
    return ws


wb = Workbook()
wb.remove(wb.active)

# ---------------- 1 接入架构 ----------------
make_table(wb, "1接入架构",
    ["项", "内容"],
    [
        ["数据流", "PLC ──(驱动 S7/Modbus/...)──▶ KEPServerEX ──(OPC UA / MQTT / REST)──▶ 后端 FastAPI ──▶ coal_records"],
        ["为什么用 KEPServer", "KEPServerEX 自带 150+ 设备驱动，屏蔽 PLC 协议差异；后端只需对接 KEPServer 一个出口，不直接碰 PLC 协议"],
        ["后端读法(推荐)", "① OPC UA（asyncua 客户端，标准/跨平台，推荐）② MQTT（KEPServer IoT Gateway 推送，后端订阅）③ REST 轮询"],
        ["KEPServer 角色", "中间件：向下连 PLC、向上统一提供数据；需在其内建 Channel→Device→Tag 三级配置"],
    ],
    [16, 96])

# ---------------- 2 PLC 信息需求 ----------------
make_table(wb, "2PLC信息需求",
    ["需要确认的信息", "说明", "示例"],
    [
        ["PLC 品牌/型号", "决定 KEPServer 用哪个驱动", "西门子 S7-1200/1500 / 汇川 / 信捷 / 三菱 / AB…"],
        ["通信协议", "决定驱动与参数", "S7（西门子）/ Modbus TCP / EtherNet/IP / …"],
        ["IP 地址", "每台 PLC 的网络地址", "192.168.1.10"],
        ["端口", "协议端口", "S7:102 / Modbus TCP:502"],
        ["站号/单元号(Unit ID)", "Modbus 从站地址；S7 的机架/槽号", "站号 1 / 机架0 槽1"],
        ["PLC 数量与分工", "几台 PLC，各管哪个系统(401/402/501/502)", "2 台：1#管分选、2#管输送"],
        ["是否已有采集", "现场组态软件/上位机是否已在采 PLC？可直接读上位机", "WinCC / 组态王 / 力控…"],
    ],
    [26, 44, 30])

# ---------------- 3 KEPServer 信息需求 ----------------
make_table(wb, "3KEPServer信息需求",
    ["需要确认的信息", "说明", "示例"],
    [
        ["KEPServerEX 版本", "影响驱动/OPC UA/MQTT 能力", "KEPServerEX 6.13"],
        ["KEPServer 安装位置(IP)", "后端要访问的机器", "192.168.1.100"],
        ["是否已建 Channel/Device", "已有配置可复用；没有则新建", "已建：Channel1(S7)→Device1"],
        ["驱动类型", "对应 PLC 协议的驱动", "Siemens TCP/IP Ethernet / Modbus TCP/IP"],
        ["数据暴露方式", "后端用哪种方式读", "OPC UA / OPC DA / MQTT / REST"],
        ["OPC UA 地址 + 认证", "UA 服务器地址、安全策略、匿名/用户名密码/证书", "opc.tcp://192.168.1.100:49320，匿名"],
        ["MQTT broker 地址(若走 MQTT)", "IoT Gateway 推送目标", "192.168.1.100:1883"],
        ["Tag 命名规范", "现有 tag 命名规则，便于映射", "如 DB1.DBD0 → Density_501"],
    ],
    [26, 44, 30])

# ---------------- 4 点位表模板 ----------------
# 预填业务量（PLC 位号留空待填）
BIZ_QUANTITIES = [
    ["悬浮液密度(实测)", "密度", "1s", "Real/Float", "1.30~1.60", "g/cm³"],
    ["501皮带秤(带煤量)", "coal_amount", "1min", "Real/Int", "0~1000", "t/h"],
    ["502皮带秤(带煤量)", "coal_amount", "1min", "Real/Int", "0~1000", "t/h"],
    ["精磁尾液位", "level", "1min", "Real/Int", "0~100", "%"],
    ["原煤灰分(灰分仪)", "raw_ash", "5min", "Real", "0~60", "%"],
    ["全水分(水分仪)", "moisture", "5min", "Real", "0~30", "%"],
    ["A系统启停", "sysA", "1s", "Bool", "0/1", "-"],
    ["B系统启停", "sysB", "1s", "Bool", "0/1", "-"],
    ["401系统启停", "sys401", "1s", "Bool", "0/1", "-"],
    ["402系统启停", "sys402", "1s", "Bool", "0/1", "-"],
    ["473脱粉启停", "desliming473", "1s", "Bool", "0/1", "-"],
    ["474脱粉启停", "desliming474", "1s", "Bool", "0/1", "-"],
    ["给料量", "给料量(协变量)", "1min", "Real", "0~1000", "t/h"],
    ["旋流器压力", "压力(协变量)", "1min", "Real", "0~1.0", "MPa"],
    ["介质煤泥含量", "煤泥含量(协变量)", "10min", "Real", "0~100", "%"],
]
make_table(wb, "4点位表模板",
    ["业务含义", "对应系统字段", "采样周期", "数据类型", "量程", "单位", "PLC 位号/Tag(待填)", "PLC 地址(待填)"],
    [[*row, "", ""] for row in BIZ_QUANTITIES],
    [20, 16, 12, 12, 14, 10, 22, 24])

# ---------------- 5 业务量采集清单 ----------------
make_table(wb, "5业务量采集清单",
    ["业务量", "对应字段", "作用", "优先级", "备注"],
    [
        ["悬浮液密度", "density", "密度建议的核心反馈量", "P0 必采", "建议 1s 采，分钟聚合"],
        ["501/502 皮带秤带煤量", "coal_amount", "总精煤量计算 + 给料量", "P0 必采", "影响总灰分加权"],
        ["精磁尾液位", "level", "粗精煤泥灰分模型特征", "P0 必采", "MLR/PLS 输入"],
        ["原煤灰分", "raw_ash", "粗灰模型 + 密度增益调度", "P0 必采", "灰分仪，可先化验兜底"],
        ["系统/脱粉启停开关", "sysA/B/401/402/473/474", "粗灰模型特征 + 工况识别", "P0 必采", "Bool，秒级"],
        ["全水分", "moisture", "工况/煤质", "P1", "水分仪或化验"],
        ["给料量/旋流器压力", "给料量/压力", "过程辨识协变量", "P1", "密度→灰分建模用"],
        ["介质煤泥含量", "煤泥含量", "K 增益调度(脏介质→K变小)", "P1", "化验或在线粘度"],
    ],
    [20, 18, 34, 12, 32])

# ---------------- 6 网络与运维 ----------------
make_table(wb, "6网络与运维",
    ["项", "要求/需确认"],
    [
        ["网络连通", "后端服务器能 ping 通 KEPServer 机器；KEPServer 能 ping 通 PLC；确认是否同一网段/防火墙端口"],
        ["采集周期", "密度/启停 1s、量/液位 1min、灰分/水分 5min；后端按分钟聚合落库"],
        ["断线重连", "后端读 KEPServer 断线自动重连；KEPServer↔PLC 断线由 KEPServer 重连并置质量位"],
        ["数据质量", "OPC 数据带 Quality（好/坏/不确定），坏点不入库并告警；跳变(>阈值)过滤"],
        ["历史补传", "停采期间的历史是否需补传？(决定是否做缓存/断点续传)"],
        ["安全", "OPC UA 用证书/用户名密码；PLC 网段与办公网隔离"],
    ],
    [16, 88])

wb.save(OUT)
print("[gen]", OUT)

"""生成《项目改造-四项任务清单》Excel"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parents[2] / "docs" / "项目改造-四项任务清单.xlsx"

# 序号 | 任务 | 需求描述 | 关键点/疑问 | 状态
ROWS = [
    (1, "四套系统合并为1个系统",
     "把 401/402/A/B 四套系统合并成一套处理",
     "合并口径待确认：数据层（不再分系统，合并存储）还是显示层（仍分存但界面合一）？历史四系统数据如何归并？",
     "待确认"),
    (2, "浮精量/粗精煤泥量/重介精煤量：录入+输入样式",
     "三个量做成「录入+输入」样式",
     "理解为：既可人工录入、也可由系统按公式计算自动填入，两者可切换/覆盖。待确认",
     "待开始"),
    (3, "总精煤灰分公式与反推",
     "总精煤灰分 = (重介精煤灰分×重介精煤量 + 浮精灰分×浮精量 + 粗精煤泥量×粗精煤泥灰分) ÷ 总煤量；后续获取到总精煤灰分后，用公式反推重介精煤灰分",
     "反推公式：重介精煤灰分 = (总精煤灰分×总煤量 − 浮精灰分×浮精量 − 粗精煤泥量×粗精煤泥灰分) ÷ 重介精煤量（需重介精煤量>0）",
     "待开始"),
    (4, "完善导入功能",
     "等导入模板样式发来后再做",
     "待模板文件",
     "等待模板"),
]


def build():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    status_fills = {
        "待确认": PatternFill("solid", fgColor="FCE4D6"),
        "待开始": PatternFill("solid", fgColor="FFF2CC"),
        "等待模板": PatternFill("solid", fgColor="E2EFDA"),
    }
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")

    ws = wb.active
    ws.title = "四项任务"
    cols = ["序号", "任务", "需求描述", "关键点/疑问", "状态"]
    widths = [6, 30, 48, 52, 10]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for row in ROWS:
        ws.append(list(row))
    for r in ws.iter_rows(min_row=2, max_row=1 + len(ROWS)):
        for c in r:
            c.border = border
            c.alignment = wrap
        r[0].alignment = Alignment(horizontal="center", vertical="center")
        r[4].alignment = Alignment(horizontal="center", vertical="center")
        if r[4].value in status_fills:
            r[4].fill = status_fills[r[4].value]
    ws.freeze_panes = "A2"
    wb.save(OUT_PATH)
    print(f"[gen] 已生成：{OUT_PATH}")


if __name__ == "__main__":
    build()

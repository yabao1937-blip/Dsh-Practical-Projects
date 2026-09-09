"""生成《四系统重构与真实数据接入-执行计划》Excel"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parents[2] / "docs" / "四系统重构与真实数据接入-执行计划.xlsx"

# Sheet1：执行计划（阶段 | 任务 | 具体内容 | 影响/验收）
PLAN = [
    ("阶段1", "系统维度回退（401/402/A/B）",
     "后端：删除 401→A、402→B 归一化；筛选改为四系统；前端下拉框恢复 401/402/A/B；浮精表、重介精煤表同步回退",
     "四系统独立；旧归一化数据清理"),
    ("阶段2", "粗精煤泥表新版导入",
     "解析新版「粗精煤泥、精磁尾 (1).xlsx」：系统列存运行组合字符串、滞后时间列写入元数据、异常值(？)置空；删除旧5月数据后重导",
     "110行新数据入库，系统列显示组合"),
    ("阶段3", "重介精煤表按皮带-系统重构",
     "密度按 401/402/A/B 四系统分行；501皮带量/灰分对应 401/402 组，502 对应 A/B 组；新增字段：磁性物含量、悬浮液煤泥含量",
     "四系统密度齐全，皮带对应关系正确"),
    ("阶段4", "粗精煤泥表新字段与周期修正",
     "新增字段：脱粉筛473/474运行状态、入洗工作面；表头周期标注按真实文件改为 5分钟（灰分仪/密度）、精磁尾液位标注在线(1秒采集)",
     "新字段可见，周期标注准确"),
    ("阶段5", "真实数据导入与模型训练",
     "导入 6.16-7.14 真实数据（原煤灰分/带煤量/运行系统/脱粉筛/精磁尾液位→315灰分）；训练粗精煤泥灰分预测模型替换专家加权占位；预测窗口按滞后1分30秒对齐",
     "模型上线，预测值接近真实化验"),
]

# Sheet2：剩余确认点（序号 | 事项 | 我的建议 | 你的回复（待填））
CONFIRM = [
    (1, "旧5月数据", "删除之前归一化口径导入的5月数据、改用新文件重导", ""),
    (2, "组合系统字符串", "粗精煤泥表系统列直接存组合(如401、A、B)，前端显示原文", ""),
    (3, "训练目标映射", "6.16-7.14 的「315灰分」作为粗精煤泥灰分真实值训练模型；训练后预测模型自动替换专家加权占位模型", ""),
    (4, "煤量异常值", "新版文件「煤量(t/h)」列的「？」置空，「301」按实测值保留", ""),
    (5, "模型训练方式", "先离线训练（脚本），生成新的模型参数JSON，后续采集到更多化验值再定期重训", ""),
]


def build():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    section_fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")

    def style_header(ws):
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border

    def style_body(ws, nrows, section_col=None):
        for r in ws.iter_rows(min_row=2, max_row=1 + nrows):
            for c in r:
                c.border = border
                c.alignment = wrap
            r[0].alignment = Alignment(horizontal="center", vertical="center")
            if section_col is not None:
                r[section_col].alignment = Alignment(horizontal="center", vertical="center")
                r[section_col].fill = section_fill

    # Sheet1
    ws = wb.active
    ws.title = "执行计划"
    cols = ["阶段", "任务", "具体内容", "影响/验收"]
    widths = [8, 24, 62, 30]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws)
    for row in PLAN:
        ws.append(list(row))
    style_body(ws, len(PLAN), section_col=0)
    ws.freeze_panes = "A2"

    # Sheet2
    ws2 = wb.create_sheet("剩余确认点")
    cols2 = ["序号", "事项", "我的建议", "你的回复（待填）"]
    widths2 = [6, 18, 62, 24]
    ws2.append(cols2)
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    style_header(ws2)
    for row in CONFIRM:
        ws2.append(list(row))
    style_body(ws2, len(CONFIRM), section_col=1)
    ws2.freeze_panes = "A2"

    wb.save(OUT_PATH)
    print(f"[gen] 已生成：{OUT_PATH}")


if __name__ == "__main__":
    build()

# -*- coding: utf-8 -*-
"""生成 Excel：MLR/PLS 训练后移（sklearn）——难点与原因 + 我的方案 + GLM/qwen 建议 + 优化结果"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\dense-medium-density-control-system\docs\MLR-PLS训练后移-难点与建议.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FDE9D9")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def make_sheet(wb, title, headers, rows, widths, fill_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if cell.column != 1 else CENTER
    if fill_col:
        for row in ws.iter_rows(min_row=2):
            v = str(row[fill_col - 1].value or "")
            if v.startswith("已") or v.startswith("✅"):
                row[fill_col - 1].fill = OK_FILL
            elif v.startswith("未") or v.startswith("⏸"):
                row[fill_col - 1].fill = WARN_FILL
    return ws


wb = Workbook()
wb.remove(wb.active)

make_sheet(wb, "1难点与原因",
    ["序号", "难点", "原因", "现状"],
    [
        [1, "存量出厂模型无法从当前 seed 复现",
         "raw_ash 均值 35.97→43.96、is_stoppage 由全0变为7行=1，训练数据已漂移；原始 Excel 不在仓库，golden 时代逐行 X 不可恢复",
         "对拍基准改为「Node 现场跑 JS 训练器 dump 全量输出」(train_js.json)，不再对存量常量"],
        [2, "sklearn 与手写 JS 数值对齐",
         "Ridge α 需映射 ×diagMean(=n-1=112)；PLSRegression 1.9 的 coef_ 是 (n_target,n_features) 布局；NIPALS deflation/QR 回代细节不同",
         "纯 Python 逐位移植做 golden；sklearn 只当内层终解，实测差 ~1e-13"],
        [3, "PLS 分量数 A 选优曲线平坦",
         "A=1 Q²=0.4266 vs A=2 0.4283（差 1.7e-3），任何预处理镜像偏差都会翻转 A",
         "离散决策(drop/λ网格点/A/production)零容差断言"],
        [4, "MLR λ 命中网格上边界",
         "LOOCV SSE 沿网格单调下降，选中 f=0.01（λ=1.12）；网格偏窄是建模问题",
         "迁移期禁扩网格，先保 parity；优化另立任务"],
        [5, "_js_round 与 toFixed 不一致",
         "toFixed 对 double 精确十进制值 round-half-away；手写 floor(x·f+0.5) 在负数平局、4.35 类浮点进位处分叉",
         "改用 Fraction(float) 精确镜像 + 边界单测"],
        [6, "持久化字段缺失",
         "history 表缺 q2Time/rmse/mae/A（production 靠 q2Time 选却未存）；同批 mlr+pls 两行无关联；drop 未落库",
         "history 加 detail JSON 快照 + train_run_id；drop 入 metrics JSON"],
    ],
    [6, 34, 56, 46], fill_col=4)

make_sheet(wb, "2我的方案",
    ["步骤", "内容", "验证结果"],
    [
        [1, "Node oracle：dump_train_js.js 在 web(2) 现场跑 App.trainCoarseModel('jun_jul')，dump 全量到 train_js.json",
         "golden：production=pls、mlr.lambda=1.12、pls.A=2、drop=[]"],
        [2, "纯 Python 逐操作移植 trainMlr/trainPls/trainCoarseModel（高斯消元/hat-LOOCV/NIPALS/_timeCvQ2）",
         "与 oracle 逐位一致 diff=0.0（含 q2Time/history/113 行回填）"],
        [3, "sklearn 内层求解器：Ridge(alpha=λ, solver='svd') + PLSRegression(n_components=A, scale=False)",
         "与移植版差 ~1e-13，<1e-9 容差"],
        [4, "API + 持久化 + 测试：POST /api/v1/training/coarse-model",
         "pytest 46 全绿"],
    ],
    [6, 62, 40])

make_sheet(wb, "3GLM建议",
    ["级别", "建议", "处置", "说明"],
    [
        ["高", "验收口径分层：系数/metrics 只做防漂移守护，预测逐值一致才是发布门槛；Q² 选型加平局带宽防「系数都过、选型翻转」", "已采纳",
         "离散决策零容差、连续量 1e-9，见测试"],
        ["高", "确认 golden 模型当初走 walk-forward 还是 LOOCV，后端复现同一条选型路径", "已采纳",
         "_timeCvQ2 + 回退 LOOCV q2 均已移植"],
        ["中", "Ridge α 标度差 n(=diagMean)；λ 选优保留手工 hat-LOOCV，不用 RidgeCV(GCV)", "已采纳",
         "α=λ 网格值=1.12，实测斜率差 7.7e-15"],
        ["中", "PLS scale=False 仍中心化；系数只到 ~1e-6 相对一致，非逐位；A 选优保留手工 LOOCV 且逐 A 重训", "已采纳",
         "coef_ reshape(-1) 兼容布局"],
        ["中", "同步训练 + 单事务 + 进程内锁；加 30s 超时与行数上限", "部分采纳",
         "已同步+单事务+Lock；超时/行数上限未加（数据量小）"],
        ["中", "代码审查：查询在锁外、history 丢字段、train_run_id、drop 未持久化", "已修复",
         "见「5审查优化结果」"],
        ["低", "_js_round 负值 half、bool 穿透数值校验、test fixture 缺失 import 即炸", "已修复",
         "Fraction 镜像 + bool 排除 + skipif"],
    ],
    [6, 56, 10, 50], fill_col=3)

make_sheet(wb, "4qwen建议",
    ["级别", "建议", "处置", "说明"],
    [
        ["P0", "B1 _js_round 非 toFixed 精确镜像（负数平局、4.35）→ Fraction(float) 精确实现", "已修复", "补边界单测"],
        ["P0", "B2 range 无校验 → Literal['jun_jul','30d','all']；B3 tol 无类型强转 → float()", "已修复", "非法 range 422"],
        ["P0", "B4 平局时间戳排序不确定 → order_by(ts, id)", "已修复", "稳定可复现"],
        ["P0", "并发：快照在锁外 → 移入锁内；UniqueConstraint(method) 挡多 worker 双写", "已修复", "两处均改"],
        ["P0", "持久化：history 加 detail JSON；drop 入 metrics；metrics 注释改 camelCase", "已修复", "见 models.py"],
        ["P1", "消除 sklearn 双重训练（抽共享 plan）；PLS LOOCV 前缀复用", "未做", "性能优化，非阻塞，避免改运算序破坏对拍"],
        ["P1", "训练改后台任务 + 409/状态端点", "未做", "单 worker + 数据量小，同步够用"],
        ["P2", "响应加 Pydantic schema；落库 sanitize 非有限值；model_id 去留决策", "未做", "整洁项，可延后"],
    ],
    [6, 56, 10, 52], fill_col=3)

make_sheet(wb, "5审查优化结果",
    ["项", "结论"],
    [
        ["核心验收", "纯 Python 移植 vs JS oracle 逐位一致(diff=0.0)；sklearn vs 移植 ~1e-13；pytest 46 全绿"],
        ["已修复(两模型共识)", "①_js_round Fraction 精确镜像 ②range Literal+tol float() ③快照移锁内+order_by(ts,id) ④coarse_models 加 train_run_id+method 唯一约束、history 加 detail JSON、drop 入 metrics ⑤bool 穿透排除 ⑥test 长度/键断言+skipif ⑦_js_round/q2Time回退/非法range 三新测试"],
        ["未做(非阻塞)", "①sklearn 双重训练消除(性能) ②PLS LOOCV 前缀复用(性能) ③后台任务+409 ④Pydantic response model ⑤sanitize 非有限值"],
        ["下一步", "①接前端训练按钮✅ ②MLR λ 网格扩宽✅(0.01→0.1/1，mlr walk-forward Q² 0.0406→0.1121) ③30d/all 范围验证✅(Python==JS：jun=113/30d=55/all=113)；剩余：接真实生产数据后复核 λ/选型"],
    ],
    [16, 90])

wb.save(OUT)
print("[gen]", OUT)
